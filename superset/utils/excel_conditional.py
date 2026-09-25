# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.
"""Stamp Explore Table / Pivot Table v2 highlights onto an Excel workbook.

Explore paints matching cells in the browser. Chart XLSX download is written
in ``QueryContextProcessor.get_data`` before client post-processing, so this
module is applied there as well as on the reports path.

Rules from ``form_data["conditional_formatting"]`` become native Excel
conditional formatting (CellIs / formula / color scale / data bar). When that
list is empty and ``show_cell_bars`` is on, numeric columns get data bars so
the download matches the Table chart's default gradient.
"""

from __future__ import annotations

import io
from typing import Any, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import (
    CellIsRule,
    ColorScaleRule,
    DataBarRule,
    FormulaRule,
)
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from superset.constants import SHOW_VALUES_AS_PERCENT_MODES
from superset.utils.excel_display import (
    apply_column_display,
    refresh_sheet_bounds,
    styles_from_pivot_form_data,
    styles_from_table_form_data,
)

# Theme tokens used by Table / Pivot Table v2 pickers, plus CSS names.
_NAMED_COLORS = {
    "success": "52C41A",
    "warning": "FAAD14",
    "error": "FF4D4F",
    "red": "FF4D4F",
    "green": "52C41A",
    "blue": "1890FF",
    "yellow": "FAAD14",
    "orange": "FA8C16",
    "purple": "722ED1",
    "cyan": "13C2C2",
    "colorsuccess": "52C41A",
    "colorwarning": "FAAD14",
    "colorerror": "FF4D4F",
    "colorsuccessbg": "F6FFED",
    "colorwarningbg": "FFFBE6",
    "colorerrorbg": "FFF2F0",
}

_CELL_IS_OPERATORS = {
    ">": "greaterThan",
    "<": "lessThan",
    ">=": "greaterThanOrEqual",
    "<=": "lessThanOrEqual",
    "=": "equal",
    "==": "equal",
    "!=": "notEqual",
    "≠": "notEqual",
    "≥": "greaterThanOrEqual",
    "≤": "lessThanOrEqual",
}

_RANGE_OPERATORS = {
    "< x <": (">", "<"),
    "< x ≤": (">", "<="),
    "≤ x <": (">=", "<"),
    "≤ x ≤": (">=", "<="),
}

_DATA_BAR_POSITIVE = "63BE7B"
_SCALE_LOW = "FFFFFF"
_OBJECT_CELL_BAR = "CELL_BAR"
_OBJECT_TEXT = "TEXT_COLOR"


def _hex_rgb(color: Any) -> Optional[str]:
    """Normalize a picker payload to a 6-digit RGB hex string."""
    if isinstance(color, dict):
        hex_value = color.get("hex")
        if isinstance(hex_value, str):
            return _hex_rgb(hex_value)
        red, green, blue = color.get("r"), color.get("g"), color.get("b")
        if None not in (red, green, blue):
            return f"{int(red):02X}{int(green):02X}{int(blue):02X}"
        return None
    if not isinstance(color, str) or not color:
        return None
    token = color.strip().lstrip("#")
    named = _NAMED_COLORS.get(token.lower().replace("_", "").replace("-", ""))
    if named:
        return named
    if len(token) == 3 and all(ch in "0123456789abcdefABCDEF" for ch in token):
        return "".join(ch * 2 for ch in token).upper()
    if len(token) >= 6 and all(ch in "0123456789abcdefABCDEF" for ch in token[:6]):
        return token[:6].upper()
    return None


def _rule_color(rule: dict[str, Any]) -> str:
    return _hex_rgb(rule.get("colorScheme")) or "52C41A"


def _excel_literal(value: Any) -> str:
    """Quote a comparison target so Excel treats it as a constant."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(value)
    text = str(value).replace('"', '""')
    return f'"{text}"'


def _header_matches(sheet_header: str, column: str) -> bool:
    """Match a rule column to a header, including Pivot ``SUM(col)`` titles."""
    left = sheet_header.strip()
    right = column.strip()
    if left == right or left.lower() == right.lower():
        return True
    return left.endswith(f"({right})") or left.endswith(f"({right.lower()})")


def _column_index(sheet: Worksheet, header_row: int, column: str) -> Optional[int]:
    for col_idx in range(1, sheet.max_column + 1):
        header_label = ""
        for row in range(header_row, 0, -1):
            value = sheet.cell(row=row, column=col_idx).value
            if value not in (None, ""):
                header_label = str(value)
                break
        if header_label and _header_matches(header_label, column):
            return col_idx
    return None


def _data_range(sheet: Worksheet, col_idx: int, header_row: int) -> Optional[str]:
    last_row = sheet.max_row
    first_row = header_row + 1
    if last_row < first_row:
        return None
    letter = get_column_letter(col_idx)
    return f"{letter}{first_row}:{letter}{last_row}"


def _solid_fill(rgb: str) -> PatternFill:
    return PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")


def _add_formula_rule(
    sheet: Worksheet,
    cell_range: str,
    formula: str,
    rgb: str,
    *,
    text_color: bool,
) -> None:
    fill = None if text_color else _solid_fill(rgb)
    font = Font(color=rgb) if text_color else None
    sheet.conditional_formatting.add(
        cell_range,
        FormulaRule(formula=[formula], fill=fill, font=font),
    )


def _add_cell_is_rule(
    sheet: Worksheet,
    cell_range: str,
    operator: str,
    formula: list[str],
    rgb: str,
    *,
    text_color: bool,
) -> None:
    fill = None if text_color else _solid_fill(rgb)
    font = Font(color=rgb) if text_color else None
    sheet.conditional_formatting.add(
        cell_range,
        CellIsRule(operator=operator, formula=formula, fill=fill, font=font),
    )


def _add_data_bar(sheet: Worksheet, cell_range: str, rgb: str) -> None:
    sheet.conditional_formatting.add(
        cell_range,
        DataBarRule(
            start_type="min",
            end_type="max",
            color=rgb,
            showValue=True,
            minLength=None,
            maxLength=None,
        ),
    )


def _apply_rule(sheet: Worksheet, header_row: int, rule: dict[str, Any]) -> None:
    column = rule.get("column")
    if not isinstance(column, str) or not column:
        return
    col_idx = _column_index(sheet, header_row, column)
    if col_idx is None:
        return
    cell_range = _data_range(sheet, col_idx, header_row)
    if cell_range is None:
        return

    rgb = _rule_color(rule)
    operator = rule.get("operator")
    object_fmt = rule.get("objectFormatting") or ""
    text_color = object_fmt == _OBJECT_TEXT
    top_left = cell_range.split(":", 1)[0]

    if object_fmt == _OBJECT_CELL_BAR:
        _add_data_bar(sheet, cell_range, rgb)
        return

    if operator in _CELL_IS_OPERATORS:
        _add_cell_is_rule(
            sheet,
            cell_range,
            _CELL_IS_OPERATORS[operator],
            [_excel_literal(rule.get("targetValue"))],
            rgb,
            text_color=text_color,
        )
        return

    if operator in _RANGE_OPERATORS:
        left_op, right_op = _RANGE_OPERATORS[operator]
        left = _excel_literal(rule.get("targetValueLeft"))
        right = _excel_literal(rule.get("targetValueRight"))
        _add_formula_rule(
            sheet,
            cell_range,
            (
                f"AND(NOT(ISBLANK({top_left})),"
                f"{top_left}{left_op}{left},"
                f"{top_left}{right_op}{right})"
            ),
            rgb,
            text_color=text_color,
        )
        return

    if operator in (None, "None", ""):
        if text_color:
            _add_formula_rule(
                sheet,
                cell_range,
                f"NOT(ISBLANK({top_left}))",
                rgb,
                text_color=True,
            )
            return
        sheet.conditional_formatting.add(
            cell_range,
            ColorScaleRule(
                start_type="min",
                start_color=_SCALE_LOW,
                end_type="max",
                end_color=rgb,
            ),
        )


def _is_numeric_header(sheet: Worksheet, col_idx: int, header_row: int) -> bool:
    for row in range(header_row + 1, min(sheet.max_row, header_row + 20) + 1):
        cell = sheet.cell(row=row, column=col_idx)
        if cell.value in (None, ""):
            continue
        if cell.data_type == "n" or isinstance(cell.value, (int, float)):
            return True
        return False
    return False


def _apply_cell_bars(sheet: Worksheet, header_row: int) -> None:
    """Attach Excel data bars on numeric columns (Table ``show_cell_bars``)."""
    for col_idx in range(1, sheet.max_column + 1):
        if not _is_numeric_header(sheet, col_idx, header_row):
            continue
        cell_range = _data_range(sheet, col_idx, header_row)
        if cell_range is None:
            continue
        _add_data_bar(sheet, cell_range, _DATA_BAR_POSITIVE)


def apply_conditional_formatting(
    workbook_bytes: bytes,
    rules: list[dict[str, Any]],
    header_rows: int = 1,
    *,
    show_cell_bars: bool = False,
) -> bytes:
    """Attach native Excel CF for Explore rules and optional Table cell bars."""
    if not rules and not show_cell_bars:
        return workbook_bytes

    workbook: Workbook = load_workbook(io.BytesIO(workbook_bytes))
    sheet = workbook.active
    # xlsxwriter omits a full dimension; without this openpyxl can see only
    # the header row and skip every highlight.
    refresh_sheet_bounds(sheet)
    header_row = max(header_rows, 1)
    for rule in rules:
        if isinstance(rule, dict):
            _apply_rule(sheet, header_row, rule)
    if show_cell_bars and not rules:
        _apply_cell_bars(sheet, header_row)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def polish_explore_xlsx(
    workbook_bytes: bytes,
    df: pd.DataFrame,
    form_data: dict[str, Any],
    include_index: bool = False,
) -> bytes:
    """Apply Explore number formats and conditional formatting to an XLSX."""
    viz_type = form_data.get("viz_type")
    if viz_type not in ("table", "pivot_table_v2"):
        return workbook_bytes

    header_rows = df.columns.nlevels if isinstance(df.columns, pd.MultiIndex) else 1
    if include_index:
        header_rows = max(header_rows, getattr(df.index, "nlevels", 1))

    skip_display = viz_type == "pivot_table_v2" and form_data.get("showValuesAs") in (
        SHOW_VALUES_AS_PERCENT_MODES
    )
    if not skip_display:
        headers = [str(column) for column in df.columns]
        if viz_type == "table":
            styles = styles_from_table_form_data(headers, form_data)
        else:
            styles = styles_from_pivot_form_data(headers, form_data)
        workbook_bytes = apply_column_display(
            workbook_bytes, styles, header_rows=header_rows
        )

    rules = form_data.get("conditionalFormatting") or form_data.get(
        "conditional_formatting"
    )
    if not isinstance(rules, list):
        rules = []
    return apply_conditional_formatting(
        workbook_bytes,
        rules,
        header_rows=header_rows,
        show_cell_bars=bool(form_data.get("show_cell_bars")),
    )

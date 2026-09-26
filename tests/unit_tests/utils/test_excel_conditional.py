# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

import io
from typing import Any

from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.worksheet import Worksheet

from superset.utils.excel_conditional import apply_conditional_formatting


def _cf_rules(sheet: Worksheet) -> list[tuple[str, Any]]:
    """Collect (range, rule) pairs from the first sheet's conditional formatting."""
    rules: list[tuple[str, Any]] = []
    for cf_range, cf_list in sheet.conditional_formatting._cf_rules.items():  # noqa: SLF001
        for rule in cf_list:
            rules.append((str(cf_range), rule))
    return rules


def _workbook_with_sales(*values: float) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "sales"
    sheet["B1"] = "region"
    for offset, value in enumerate(values, start=2):
        sheet.cell(row=offset, column=1, value=value)
    sheet["B2"] = "west"
    raw = io.BytesIO()
    workbook.save(raw)
    return raw.getvalue()


def test_apply_conditional_formatting_writes_cell_is_rule() -> None:
    """A greater-than rule becomes an Excel CellIs rule on the sales column."""
    styled = apply_conditional_formatting(
        _workbook_with_sales(10, 50, 5),
        [
            {
                "column": "sales",
                "operator": ">",
                "targetValue": 8,
                "colorScheme": "#FF0000",
            }
        ],
    )
    result = load_workbook(io.BytesIO(styled)).active
    rules = _cf_rules(result)
    assert len(rules) == 1
    cf_range, rule = rules[0]
    assert "A2" in cf_range
    assert rule.type == "cellIs"
    assert rule.operator == "greaterThan"


def test_apply_cell_bars_when_show_cell_bars_and_no_custom_rules() -> None:
    """Table show_cell_bars adds data bars on numeric columns only."""
    styled = apply_conditional_formatting(
        _workbook_with_sales(10),
        [],
        show_cell_bars=True,
    )
    result = load_workbook(io.BytesIO(styled)).active
    rules = _cf_rules(result)
    assert any(rule.type == "dataBar" for _, rule in rules)
    assert not any("B2" in cf_range for cf_range, _ in rules)


def test_cell_bar_object_writes_data_bar_only_on_matching_cells() -> None:
    """CELL_BAR with a comparator covers matching cells, not the whole column."""
    styled = apply_conditional_formatting(
        _workbook_with_sales(12.5, 5, 20),
        [
            {
                "column": "sales",
                "operator": ">",
                "targetValue": 10,
                "colorScheme": "#5AC189",
                "objectFormatting": "CELL_BAR",
            }
        ],
    )
    result = load_workbook(io.BytesIO(styled)).active
    rules = _cf_rules(result)
    assert any(rule.type == "dataBar" for _, rule in rules)
    cf_range = " ".join(rng for rng, rule in rules if rule.type == "dataBar")
    assert "A2" in cf_range
    assert "A4" in cf_range
    assert "A3" not in cf_range


def test_color_scale_when_operator_is_none() -> None:
    """Operator None becomes a two-color scale, honoring minBound/maxBound."""
    styled = apply_conditional_formatting(
        _workbook_with_sales(10),
        [
            {
                "column": "sales",
                "operator": None,
                "colorScheme": "colorSuccess",
                "minBound": 0,
                "maxBound": 100,
            }
        ],
    )
    result = load_workbook(io.BytesIO(styled)).active
    rules = _cf_rules(result)
    assert any(rule.type == "colorScale" for _, rule in rules)
    scale = next(rule for _, rule in rules if rule.type == "colorScale")
    assert scale.colorScale.cfvo[0].type == "num"
    assert float(scale.colorScale.cfvo[0].val) == 0
    assert scale.colorScale.cfvo[1].type == "num"
    assert float(scale.colorScale.cfvo[1].val) == 100


def test_rule_matches_verbose_column_header() -> None:
    """Rules keyed by the physical column still apply after verbose rename."""
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Revenue"
    sheet["A2"] = 12
    raw = io.BytesIO()
    workbook.save(raw)
    styled = apply_conditional_formatting(
        raw.getvalue(),
        [
            {
                "column": "sales",
                "operator": ">",
                "targetValue": 8,
                "colorScheme": "#FF0000",
            }
        ],
        verbose_map={"sales": "Revenue"},
    )
    rules = _cf_rules(load_workbook(io.BytesIO(styled)).active)
    assert len(rules) == 1
    assert "A2" in rules[0][0]

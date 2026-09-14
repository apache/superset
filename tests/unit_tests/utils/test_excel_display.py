# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

import io
from datetime import datetime

from openpyxl import Workbook, load_workbook

from superset.utils.excel_display import (
    ExcelColumnDisplay,
    apply_column_display,
    d3_number_to_excel,
    d3_time_to_excel,
    styles_from_pivot_form_data,
    styles_from_table_form_data,
)


def test_d3_number_to_excel_common_specifiers() -> None:
    assert d3_number_to_excel(",.2f") == "#,##0.00"
    assert d3_number_to_excel(".0%") == "0%"
    assert d3_number_to_excel(",.1%") == "#,##0.0%"
    assert d3_number_to_excel("d") == "0"
    assert d3_number_to_excel(".3s") is None
    assert d3_number_to_excel("SMART_NUMBER") is None
    assert (
        d3_number_to_excel(",.2f", {"symbol": "EUR", "symbolPosition": "prefix"})
        == '"€"#,##0.00'
    )
    assert (
        d3_number_to_excel(".2f", {"symbol": "USD", "symbolPosition": "suffix"})
        == '0.00"$"'
    )


def test_d3_time_to_excel_common_patterns() -> None:
    assert d3_time_to_excel("%Y-%m-%d") == "yyyy-mm-dd"
    assert d3_time_to_excel("%Y-%m-%d %H:%M:%S") == "yyyy-mm-dd hh:mm:ss"
    assert d3_time_to_excel("%d/%m/%Y") == "dd/mm/yyyy"
    assert d3_time_to_excel("%Q") is None


def test_styles_from_table_form_data_generated_config() -> None:
    styles = styles_from_table_form_data(
        ["revenue", "when"],
        {
            "column_config": {
                "revenue": {
                    "d3NumberFormat": ",.2f",
                    "horizontalAlign": "right",
                    "currencyFormat": {"symbol": "USD", "symbolPosition": "prefix"},
                },
                "when": {"d3TimeFormat": "%Y-%m-%d", "horizontalAlign": "center"},
                "ignored": {"d3NumberFormat": ".2f"},
            }
        },
    )
    assert styles["revenue"].number_format == '"$"#,##0.00'
    assert styles["revenue"].alignment == "right"
    assert styles["when"].number_format == "yyyy-mm-dd"
    assert "ignored" not in styles


def test_styles_from_pivot_form_data_matches_flattened_metric() -> None:
    styles = styles_from_pivot_form_data(
        ["2024 SUM(num)", "state"],
        {"valueFormat": ".1%", "columnFormats": {"SUM(num)": ",.0f"}},
    )
    assert styles["2024 SUM(num)"].number_format == "#,##0"


def test_apply_column_display_on_generated_workbook() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "revenue"
    sheet["B1"] = "label"
    sheet["A2"] = 1234.5
    sheet["B2"] = "west"
    raw = io.BytesIO()
    workbook.save(raw)

    styled = apply_column_display(
        raw.getvalue(),
        {
            "revenue": ExcelColumnDisplay(number_format="#,##0.00", alignment="right"),
            "label": ExcelColumnDisplay(alignment="left"),
        },
    )
    result = load_workbook(io.BytesIO(styled)).active
    assert result["A2"].number_format == "#,##0.00"
    assert result["A2"].alignment.horizontal == "right"
    assert result["B2"].alignment.horizontal == "left"
    assert result["B2"].value == "west"


def test_apply_column_display_formats_datetimes() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "ds"
    sheet["A2"] = datetime(2024, 1, 15, 8, 30)
    raw = io.BytesIO()
    workbook.save(raw)
    styled = apply_column_display(
        raw.getvalue(),
        {"ds": ExcelColumnDisplay(number_format="yyyy-mm-dd hh:mm:ss")},
    )
    cell = load_workbook(io.BytesIO(styled)).active["A2"]
    assert cell.number_format == "yyyy-mm-dd hh:mm:ss"
    assert cell.value == datetime(2024, 1, 15, 8, 30)

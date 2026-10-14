# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.
import io
from datetime import datetime
from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from superset.utils import json, univer_excel
from superset.utils.core import GenericDataType
from superset.utils.pandas_postprocessing import flatten, pivot
from superset.utils.univer_excel import (
    build_grid,
    build_pivot_grid,
    df_to_univer_excel,
    is_univer_chart,
    to_cell_value,
)

DF = pd.DataFrame({"genre": ["Action", "Puzzle"], "Total": [524.999999999985, 2]})
COLTYPES = [GenericDataType.STRING, GenericDataType.NUMERIC]


def saved_state(**overrides: Any) -> str:
    """A saved chart state with live data, like the frontend writes it."""
    state = {
        "version": 1,
        "dataRows": 3,
        "dataCols": 2,
        "stripped": True,
        "headers": ["genre", "Total"],
        "workbook": {
            "sheetOrder": ["data", "notes"],
            "styles": {
                "superset_header": {"bl": 1, "bg": {"rgb": "#F2F2F2"}},
                "red": {"cl": {"rgb": "rgb(255, 0, 0)"}, "ht": 3},
            },
            "sheets": {
                "data": {
                    "name": "Data",
                    "rowCount": 4,
                    "columnCount": 3,
                    "freeze": {"xSplit": 0, "ySplit": 1, "startRow": 1},
                    "columnData": {"0": {"w": 140}},
                    "mergeData": [
                        {"startRow": 3, "endRow": 3, "startColumn": 0, "endColumn": 1}
                    ],
                    "cellData": {
                        "0": {
                            "0": {"v": "Genre", "s": "superset_header"},
                            "1": {"s": "superset_header"},
                        },
                        "1": {"1": {"s": "red"}},
                        "3": {
                            "0": {"v": "=not a formula"},
                            "2": {"f": "=SUM(B2:B3)", "si": "a", "v": 527},
                        },
                        "4": {"2": {"si": "a"}},
                    },
                },
                "notes": {"name": "Notes", "cellData": {"0": {"0": {"v": "hi"}}}},
            },
        },
    }
    state.update(overrides)
    return json.dumps(state)


def export(
    form_data: dict[str, Any], df: pd.DataFrame = DF, data_only: bool = False
) -> Any:
    return load_workbook(
        io.BytesIO(df_to_univer_excel(df, COLTYPES, form_data)), data_only=data_only
    )


def values(ws: Worksheet) -> list[list[Any]]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def test_is_univer_chart() -> None:
    assert is_univer_chart({"viz_type": "univer_sheet"})
    assert not is_univer_chart({"viz_type": "table"})
    assert not is_univer_chart(None)


def test_to_cell_value() -> None:
    assert to_cell_value(datetime(2024, 1, 31), GenericDataType.TEMPORAL) == (
        "2024-01-31"
    )
    assert (
        to_cell_value(pd.Timestamp("2024-01-31 13:05:09"), GenericDataType.TEMPORAL)
        == "2024-01-31 13:05:09"
    )
    assert to_cell_value(0.1 + 0.2) == 0.3
    assert to_cell_value(float("nan")) is None
    assert to_cell_value({"a": 1}) == '{"a": 1}'


def test_build_grid() -> None:
    assert build_grid(DF, COLTYPES) == [
        ["genre", "Total"],
        ["Action", 525],
        ["Puzzle", 2],
    ]


def test_export_without_saved_edits() -> None:
    ws = export({"viz_type": "univer_sheet"}).active
    assert values(ws) == [["genre", "Total"], ["Action", 525], ["Puzzle", 2]]
    assert ws["A1"].font.bold
    assert ws["A1"].fill.fgColor.rgb == "FFF2F2F2"
    assert ws.freeze_panes == "A2"


def test_export_with_saved_edits() -> None:
    book = export(
        {"viz_type": "univer_sheet", "univer_workbook": saved_state()},
        pd.DataFrame({"genre": ["Racing"], "Total": [10]}),
    )
    assert book.sheetnames == ["Data", "Notes"]
    ws = book["Data"]
    # renamed header kept, fresh query values, old rows cleared
    assert ws["A1"].value == "Genre"
    assert ws["B1"].value == "Total"
    assert ws["A2"].value == "Racing"
    assert ws["B2"].value == 10
    assert ws["A3"].value is None
    # formatting, formulas, merges and sizes
    assert ws["B2"].font.color.rgb == "FFFF0000"
    assert ws["B2"].alignment.horizontal == "right"
    assert ws["C4"].value == "=SUM(B2:B3)"
    assert ws["C5"].value == "=SUM(B3:B4)"
    assert ws["A4"].value == "=not a formula"
    assert ws["A4"].data_type == "s"
    assert "A4:B4" in ws.merged_cells
    # 140px wide, about 20 characters
    assert round(ws.column_dimensions["A"].width) == 20
    assert ws.freeze_panes == "A2"
    assert book["Notes"]["A1"].value == "hi"


def test_formulas_keep_their_values() -> None:
    ws = export(
        {"viz_type": "univer_sheet", "univer_workbook": saved_state()},
        data_only=True,
    ).active
    assert ws["C4"].value == 527


def test_renamed_header_dropped_when_the_column_changes() -> None:
    ws = export(
        {"viz_type": "univer_sheet", "univer_workbook": saved_state()},
        pd.DataFrame({"platform": ["PS4"], "Total": [1]}),
    ).active
    assert ws["A1"].value == "platform"


def test_export_without_live_data_uses_saved_values() -> None:
    state = json.loads(saved_state(stripped=False))
    state["workbook"]["sheets"]["data"]["cellData"]["1"] = {"0": {"v": "edited"}}
    ws = export(
        {
            "viz_type": "univer_sheet",
            "univer_live_data": False,
            "univer_workbook": json.dumps(state),
        }
    ).active
    assert ws["A2"].value == "edited"


def test_build_pivot_grid() -> None:
    df = pd.DataFrame(
        {
            "genre": ["Action", "Total"],
            "count, PS4": [1, 1],
            "count, Wii\\, U": [2, 2],
        }
    )
    grid, layout = build_pivot_grid(
        df,
        [GenericDataType.STRING, GenericDataType.NUMERIC, GenericDataType.NUMERIC],
        {"genre": "Genre", "count": "COUNT(*)"},
        1,
    )
    assert grid == [
        ["Genre", "COUNT(*)", "COUNT(*)"],
        [None, "PS4", "Wii, U"],
        ["Action", 1, 2],
        ["Total", 1, 2],
    ]
    assert layout == {
        "headerRows": 2,
        "frozenCols": 1,
        "merges": [
            {"startRow": 0, "endRow": 1, "startColumn": 0, "endColumn": 0},
            {"startRow": 0, "endRow": 0, "startColumn": 1, "endColumn": 2},
        ],
    }


def test_export_pivot() -> None:
    data = pd.DataFrame(
        {
            "genre": ["Action", "Action", "Puzzle"],
            "platform": ["PS4", "Wii", "PS4"],
            "count": [1, 2, 3],
        }
    )
    df = flatten(
        pivot(
            data,
            index=["genre"],
            columns=["platform"],
            aggregates={"count": {"operator": "sum"}},
            marginal_distributions=True,
            marginal_distribution_name="Total",
        )
    )
    contents = df_to_univer_excel(
        df,
        [GenericDataType.STRING] + [GenericDataType.NUMERIC] * 3,
        {
            "viz_type": "univer_sheet",
            "query_mode": "pivot",
            "univer_pivot_rows": ["genre"],
        },
    )
    ws = load_workbook(io.BytesIO(contents)).active
    assert values(ws) == [
        ["genre", "count", None, None],
        [None, "PS4", "Wii", "Total"],
        ["Action", 1, 2, 3],
        ["Puzzle", 3, None, 3],
        ["Total", 4, 2, 6],
    ]
    assert {str(merge) for merge in ws.merged_cells.ranges} == {"A1:A2", "B1:D1"}
    assert ws.freeze_panes == "B3"


def test_export_outline_groups(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(univer_excel, "MANUAL_GROUPS", True)
    state = json.loads(saved_state())
    state["outline"] = {
        "rows": [
            {"start": {"key": "Puzzle"}, "end": {"key": "Racing"}, "collapsed": True}
        ],
        "columns": [{"start": {"index": 1}, "end": {"index": 1}}],
    }
    df = pd.DataFrame(
        {"genre": ["Action", "Puzzle", "Racing", "Sports"], "Total": [1, 2, 3, 4]}
    )
    ws = export(
        {
            "viz_type": "univer_sheet",
            "groupby": ["genre"],
            "univer_workbook": json.dumps(state),
        },
        df,
    ).active
    # Puzzle and Racing are rows 3 and 4, grouped and collapsed
    assert [ws.row_dimensions[r].outline_level for r in (2, 3, 4, 5)] == [
        0,
        1,
        1,
        0,
    ]
    assert ws.row_dimensions[3].hidden
    assert ws.row_dimensions[4].hidden
    assert not ws.row_dimensions[2].hidden
    assert ws.column_dimensions["B"].outline_level == 1

    # while manual groups are turned off, saved ones are ignored
    monkeypatch.setattr(univer_excel, "MANUAL_GROUPS", False)
    ws = export(
        {
            "viz_type": "univer_sheet",
            "groupby": ["genre"],
            "univer_workbook": json.dumps(state),
        },
        df,
    ).active
    assert not ws.row_dimensions[3].outline_level
    assert not ws.row_dimensions[3].hidden


def test_export_grouped_rows() -> None:
    df = pd.DataFrame(
        {
            "genre": ["Action", "Puzzle", "Action"],
            "platform": ["PS4", "PS4", "Wii"],
            "sales": [1, 2, 3.1],
        }
    )
    contents = df_to_univer_excel(
        df,
        [GenericDataType.STRING, GenericDataType.STRING, GenericDataType.NUMERIC],
        {
            "viz_type": "univer_sheet",
            "groupby": ["genre", "platform"],
            "univer_row_grouping": True,
        },
    )
    ws = load_workbook(io.BytesIO(contents)).active
    assert values(ws) == [
        ["genre", "platform", "sales"],
        ["Action", "Subtotals", 4.1],
        [None, "PS4", 1],
        [None, "Wii", 3.1],
        ["Puzzle", "Subtotals", 2],
        [None, "PS4", 2],
        ["Grand Total", None, 6.1],
    ]
    assert {str(merge) for merge in ws.merged_cells.ranges} == {"A2:A4", "A5:A6"}
    assert [ws.row_dimensions[r].outline_level for r in range(2, 8)] == [
        0,
        1,
        1,
        0,
        1,
        0,
    ]
    assert not ws.sheet_properties.outlinePr.summaryBelow
    assert ws["A2"].font.bold
    assert ws["A2"].alignment.vertical == "top"
    assert ws["C2"].font.bold
    assert not ws["C3"].font.bold

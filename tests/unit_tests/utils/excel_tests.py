# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

import io
from datetime import datetime, timedelta, timezone

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from pandas.api.types import is_numeric_dtype

from superset.utils.core import GenericDataType
from superset.utils.excel import (
    apply_column_types,
    df_to_excel,
    NEUTRAL_TIMESTAMP,
    quote_formulas,
)


def test_timezone_conversion() -> None:
    """
    Timezone-aware values are stored as naive Excel datetimes (wall clock).
    """
    df = pd.DataFrame({"dt": [datetime(2023, 1, 1, 0, 0, tzinfo=timezone.utc)]})
    apply_column_types(df, [GenericDataType.TEMPORAL])
    assert df["dt"].tolist() == [datetime(2023, 1, 1, 0, 0)]
    contents = df_to_excel(df)
    exported = pd.read_excel(contents)["dt"][0]
    assert pd.Timestamp(exported) == pd.Timestamp("2023-01-01 00:00:00")


def test_timezone_keeps_wall_clock_not_utc_shift() -> None:
    """
    Offsets are dropped without converting to UTC.

    2023-01-01 00:00 in UTC+3 must remain midnight, not 2022-12-31 21:00.
    """
    plus_three = timezone(timedelta(hours=3))
    df = pd.DataFrame({"dt": [datetime(2023, 1, 1, 0, 0, tzinfo=plus_three)]})
    apply_column_types(df, [GenericDataType.TEMPORAL])
    assert df["dt"].iloc[0] == datetime(2023, 1, 1, 0, 0)


def test_timezone_aware_index_exports() -> None:
    """Pivot-style indexes with tz-aware timestamps still serialize to Excel."""
    index = pd.DatetimeIndex(
        [datetime(2024, 6, 1, 12, 0, tzinfo=timezone.utc)],
        name="ds",
    )
    df = pd.DataFrame({"value": [1.5]}, index=index)
    contents = df_to_excel(df)
    result = pd.read_excel(contents, index_col=0)
    assert pd.Timestamp(result.index[0]) == pd.Timestamp("2024-06-01 12:00:00")


def test_quote_formulas() -> None:
    """
    Test that formulas are quoted in Excel.
    """
    df = pd.DataFrame({"formula": ["=SUM(A1:A2)", "normal", "@SUM(A1:A2)"]})
    contents = df_to_excel(df)
    assert pd.read_excel(contents)["formula"].tolist() == [
        "'=SUM(A1:A2)",
        "normal",
        "'@SUM(A1:A2)",
    ]


def test_quote_formulas_with_leading_whitespace() -> None:
    """
    Formulas are quoted even when a space, tab, or newline comes first.

    Whitespace in front of the prefix stops XlsxWriter recognising a formula at
    all, so a value with whitespace in front of it is written as a plain string
    cell whether or not it is quoted. This is an export-consistency check
    rather than a guard against the workbook evaluating anything: the streaming
    writer (``excel_streaming``) and the CSV writer both strip before reading
    the prefix, so without this the same cell goes out quoted from two of the
    three paths and bare from this one.
    """
    df = pd.DataFrame(
        {"formula": [" =1+1", "\t=1+1", "\n=1+1", "  @SUM(A1)", "normal"]}
    )
    contents = df_to_excel(df)
    assert pd.read_excel(contents)["formula"].tolist() == [
        "' =1+1",
        "'\t=1+1",
        "'\n=1+1",
        "'  @SUM(A1)",
        "normal",
    ]


def test_quote_formulas_leaves_blank_and_whitespace_only_cells_alone() -> None:
    """
    A cell that is empty or holds only whitespace comes back unchanged.

    ``lstrip`` leaves nothing to inspect in either, so testing the stripped
    value before reading its first character is what keeps a whitespace-only
    cell from raising ``IndexError``. The plain ``len(value)`` guard this
    replaced only covered the empty string, because it read the first
    character of the original value rather than of the stripped one.

    The assertion is made on ``quote_formulas`` directly rather than through a
    round-trip, since a blank cell reads back out of a sheet as a missing value
    and would hide the result.
    """
    df = pd.DataFrame({"cell": ["", " \t\n", "   "]})

    assert quote_formulas(df)["cell"].tolist() == ["", " \t\n", "   "]


def test_quote_formulas_in_headers_and_index() -> None:
    """
    Test that formulas in column headers and index labels are quoted too.

    Pivot exports promote data values into the column MultiIndex and row
    index, so hostile warehouse strings can end up there.
    """
    df = pd.DataFrame(
        {"=SUM(A1:A2)": ["normal"]},
        index=pd.Index(['=cmd|" /C calc"!A0'], name="label"),
    )
    contents = df_to_excel(df)
    result = pd.read_excel(contents, index_col=0)
    assert result.columns.tolist() == ["'=SUM(A1:A2)"]
    assert result.index.tolist() == ['\'=cmd|" /C calc"!A0']


def test_quote_formulas_in_axis_names() -> None:
    """
    Test that formula-triggering axis *names* are quoted too, not just axis
    labels/values. Pivot exports can promote a warehouse-controlled column
    name to ``df.index.name`` (or a MultiIndex level name), and pandas
    writes those names into the sheet as header cells. ``rename`` alone
    leaves these untouched since they are a separate attribute from the
    axis labels/values, so this is asserted directly against the
    ``quote_formulas`` output rather than round-tripped through
    ``read_excel``, whose header parsing doesn't reliably preserve the
    columns-axis name.
    """
    df = pd.DataFrame(
        {"value": ["normal"]},
        index=pd.Index(["row"], name="=SUM(A1:A2)"),
    )
    df.columns.name = "+cmd"
    result = quote_formulas(df)
    assert result.index.name == "'=SUM(A1:A2)"
    assert result.columns.name == "'+cmd"

    # exercised end-to-end to confirm it doesn't error when the axis names
    # are written out as sheet header cells
    df_to_excel(df)


def test_document_properties_are_neutral() -> None:
    """
    Test that exported workbooks do not carry identifying document properties.
    """
    df = pd.DataFrame({"a": [1, 2], "b": ["x", "y"]})
    contents = df_to_excel(df, index=False)

    workbook = load_workbook(io.BytesIO(contents))
    properties = workbook.properties

    # Authoring/descriptive fields are cleared.
    for field in (
        "creator",
        "lastModifiedBy",
        "title",
        "subject",
        "description",
        "keywords",
        "category",
    ):
        value = getattr(properties, field)
        assert value in (None, ""), f"{field} should be empty, got {value!r}"

    # Timestamps are pinned to a fixed, neutral value rather than the
    # actual generation time.
    assert properties.created == NEUTRAL_TIMESTAMP
    assert properties.modified == NEUTRAL_TIMESTAMP


def test_column_data_types_with_one_numeric_column():
    df = pd.DataFrame(
        {
            "col0": ["123", "1", "2", "3"],
            "col1": ["456", "5.67", "0", ".45"],
            "col2": [
                datetime(2023, 1, 1, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 2, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 3, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 4, 0, 0, tzinfo=timezone.utc),
            ],
            "col3": ["True", "False", "True", "False"],
        }
    )
    coltypes: list[GenericDataType] = [
        GenericDataType.STRING,
        GenericDataType.NUMERIC,
        GenericDataType.TEMPORAL,
        GenericDataType.BOOLEAN,
    ]

    # only col1 should be converted to numeric, according to coltypes definition
    assert not is_numeric_dtype(df["col1"])
    apply_column_types(df, coltypes)
    assert not is_numeric_dtype(df["col0"])
    assert is_numeric_dtype(df["col1"])
    assert not is_numeric_dtype(df["col2"])
    assert not is_numeric_dtype(df["col3"])


def test_column_data_types_with_failing_conversion():
    df = pd.DataFrame(
        {
            "col0": ["123", "1", "2", "3"],
            "col1": ["456", "non_numeric_value", "0", ".45"],
            "col2": [
                datetime(2023, 1, 1, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 2, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 3, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 4, 0, 0, tzinfo=timezone.utc),
            ],
            "col3": ["True", "False", "True", "False"],
        }
    )
    coltypes: list[GenericDataType] = [
        GenericDataType.STRING,
        GenericDataType.NUMERIC,
        GenericDataType.TEMPORAL,
        GenericDataType.BOOLEAN,
    ]

    # should not fail neither convert
    assert not is_numeric_dtype(df["col1"])
    apply_column_types(df, coltypes)
    assert not is_numeric_dtype(df["col0"])
    assert not is_numeric_dtype(df["col1"])
    assert not is_numeric_dtype(df["col2"])
    assert not is_numeric_dtype(df["col3"])


def test_apply_column_types_with_duplicate_column_labels() -> None:
    """
    Test that duplicate column labels do not break the export.

    The verbose_map rename in QueryContextProcessor.get_data can collapse two
    columns onto the same label, which used to raise
    "'DataFrame' object has no attribute 'dtype'".
    """
    df = pd.DataFrame(
        [
            ["1", datetime(2023, 1, 1, 0, 0, tzinfo=timezone.utc), "2"],
            ["3", datetime(2023, 1, 2, 0, 0, tzinfo=timezone.utc), "4"],
        ],
        columns=["dupe", "dupe", "other"],
    )
    coltypes: list[GenericDataType] = [
        GenericDataType.STRING,
        GenericDataType.TEMPORAL,
        GenericDataType.NUMERIC,
    ]

    apply_column_types(df, coltypes)

    # each position is typed independently, despite sharing a label
    assert not is_numeric_dtype(df.iloc[:, 0])
    assert list(df.iloc[:, 1]) == [
        datetime(2023, 1, 1, 0, 0),
        datetime(2023, 1, 2, 0, 0),
    ]
    assert is_numeric_dtype(df.iloc[:, 2])

    contents = df_to_excel(df, index=False)
    assert pd.read_excel(contents).shape == (2, 3)


def test_quote_formulas_with_duplicate_column_labels() -> None:
    """
    Test that formulas are quoted even when column labels are duplicated.
    """
    df = pd.DataFrame(
        [["=SUM(A1:A2)", "@SUM(A1:A2)", "normal"]],
        columns=["dupe", "dupe", "other"],
    )

    result = quote_formulas(df)

    assert result.iloc[0].tolist() == ["'=SUM(A1:A2)", "'@SUM(A1:A2)", "normal"]


def test_quote_formulas_with_dedicated_string_dtype() -> None:
    """
    Test that formulas are quoted in columns using the dedicated string dtype.

    pandas 3 gives string columns a ``str`` dtype rather than ``object``, so an
    object-only dtype check would skip them and leave formulas unquoted.
    """
    df = pd.DataFrame({"formula": pd.array(["=SUM(A1:A2)", "normal"], dtype="string")})

    result = quote_formulas(df)

    assert result["formula"].tolist() == ["'=SUM(A1:A2)", "normal"]


def test_column_data_types_with_large_numeric_values():
    df = pd.DataFrame(
        {
            "big_number": [
                10**14,
                999999999999999,
                10**15 + 1,
                10**16,
                1100108628127863,
                2**54,
            ],
        }
    )
    apply_column_types(df, [GenericDataType.NUMERIC])
    assert df["big_number"].tolist() == [
        100000000000000,
        999999999999999,
        "1000000000000001",
        "10000000000000000",
        "1100108628127863",
        "18014398509481984",
    ]


def test_numeric_column_with_blanks_preserves_large_integers() -> None:
    """Blank cells must not round integers before Excel's text conversion."""
    large = 2**53 + 1
    df = pd.DataFrame({"value": pd.Series([large, None, 3], dtype=object)})
    apply_column_types(df, [GenericDataType.NUMERIC])
    sheet = load_workbook(io.BytesIO(df_to_excel(df, index=False))).active
    assert sheet["A2"].value == str(large)
    assert sheet["A2"].data_type == "s"
    assert sheet["A3"].value is None
    assert sheet["A4"].value == 3
    assert sheet["A4"].data_type == "n"


@pytest.mark.parametrize(
    "values, expected",
    [
        pytest.param(
            [2**53 + 1, None, 7, -(2**53 + 3)],
            ["9007199254740993", None, 7, "'-9007199254740995"],
            id="python-integers",
        ),
        pytest.param(
            [np.int64(2**53 + 1), None, np.int64(7), np.int64(-(2**63))],
            ["9007199254740993", None, 7, "'-9223372036854775808"],
            id="numpy-integers",
        ),
        pytest.param(
            [2**63 - 1, np.int64(2**63 - 1), None, np.int64(-7)],
            ["9223372036854775807", "9223372036854775807", None, -7],
            id="python-and-numpy-integers",
        ),
    ],
)
def test_apply_column_types_object_big_integers_exact(
    values: list[object], expected: list[object]
) -> None:
    """Nullable integers (an object column) keep every digit in the workbook.

    Values beyond 10**15 become text, smaller values stay numeric cells, nulls
    stay blank, and negative text keeps the formula-quoting prefix.
    """
    df = pd.DataFrame({"value": pd.Series(values, dtype=object)})

    apply_column_types(df, [GenericDataType.NUMERIC])
    sheet = load_workbook(io.BytesIO(df_to_excel(df, index=False))).active

    assert [row[0].value for row in sheet.iter_rows(min_row=2)] == expected

# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

from io import BytesIO, StringIO
from typing import Any
from unittest.mock import MagicMock

import pandas as pd
import pytest
from flask_babel import lazy_gettext as _
from sqlalchemy.orm.session import Session

from superset.charts.client_processing import (
    _collapsed_metric,
    _reduce,
    apply_client_processing,
    apply_pivot_number_formats,
    format_column,
    payload_coltypes,
    pivot_df,
    pivot_table_v2,
    table,
)
from superset.common.chart_data import ChartDataResultFormat
from superset.utils import excel, json
from superset.utils.core import GenericDataType
from superset.utils.csv import chart_data_to_dataframe
from tests.conftest import with_config


def test_pivot_df_no_cols_no_rows_single_metric():
    """
    Pivot table when no cols/rows and 1 metric are selected.
    """
    # when no cols/rows are selected there are no groupbys in the query,
    # and the data has only the metric(s)
    df = pd.DataFrame.from_dict({"SUM(num)": {0: 80679663}})
    assert (
        df.to_markdown()
        == """
|    |   SUM(num) |
|---:|-----------:|
|  0 |   80679663 |
    """.strip()
    )

    pivoted = pivot_df(
        df,
        rows=[],
        columns=[],
        metrics=["SUM(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                  |   ('SUM(num)',) |
|:-----------------|----------------:|
| ('{_("Total")} (Sum)',) |        80679663 |
    """.strip()
    )

    # transpose_pivot and combine_metrics do nothing in this case
    pivoted = pivot_df(
        df,
        rows=[],
        columns=[],
        metrics=["SUM(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  |   ('SUM(num)',) |
|:-----------------|----------------:|
| ('Total (Sum)',) |        80679663 |
    """.strip()
    )

    # apply_metrics_on_rows will pivot the table, moving the metrics
    # to rows
    pivoted = pivot_df(
        df,
        rows=[],
        columns=[],
        metrics=["SUM(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|               |   ('{_("Total")} (Sum)',) |
|:--------------|-------------------:|
| ('SUM(num)',) |           80679663 |
    """.strip()
    )

    # showing totals
    pivoted = pivot_df(
        df,
        rows=[],
        columns=[],
        metrics=["SUM(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=True,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                  |   ('SUM(num)',) |   ('Total (Sum)',) |
|:-----------------|----------------:|-------------------:|
| ('{_("Total")} (Sum)',) |        80679663 |           80679663 |
    """.strip()
    )


def test_pivot_df_no_cols_no_rows_two_metrics():
    """
    Pivot table when no cols/rows and 2 metrics are selected.
    """
    # when no cols/rows are selected there are no groupbys in the query,
    # and the data has only the metrics
    df = pd.DataFrame.from_dict({"SUM(num)": {0: 80679663}, "MAX(num)": {0: 37296}})
    assert (
        df.to_markdown()
        == """
|    |   SUM(num) |   MAX(num) |
|---:|-----------:|-----------:|
|  0 |   80679663 |      37296 |
    """.strip()
    )

    pivoted = pivot_df(
        df,
        rows=[],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                  |   ('SUM(num)',) |   ('MAX(num)',) |
|:-----------------|----------------:|----------------:|
| ('{_("Total")} (Sum)',) |        80679663 |           37296 |
    """.strip()
    )

    # transpose_pivot and combine_metrics do nothing in this case
    pivoted = pivot_df(
        df,
        rows=[],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  |   ('SUM(num)',) |   ('MAX(num)',) |
|:-----------------|----------------:|----------------:|
| ('Total (Sum)',) |        80679663 |           37296 |
    """.strip()
    )

    # apply_metrics_on_rows will pivot the table, moving the metrics
    # to rows
    pivoted = pivot_df(
        df,
        rows=[],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|               |   ('{_("Total")} (Sum)',) |
|:--------------|-------------------:|
| ('SUM(num)',) |           80679663 |
| ('MAX(num)',) |              37296 |
    """.strip()
    )

    # when showing totals we only add a column, since adding a row
    # would be redundant
    pivoted = pivot_df(
        df,
        rows=[],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=True,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                  |   ('SUM(num)',) |   ('MAX(num)',) |   ('{_("Total")} (Sum)',) |
|:-----------------|----------------:|----------------:|-------------------:|
| ('{_("Total")} (Sum)',) |     8.06797e+07 |           37296 |                nan |
    """.strip()
    )


def test_pivot_df_single_row_two_metrics():
    """
    Pivot table when a single column and 2 metrics are selected.
    """
    df = pd.DataFrame.from_dict(
        {
            "gender": {0: "girl", 1: "boy"},
            "SUM(num)": {0: 118065, 1: 47123},
            "MAX(num)": {0: 2588, 1: 1280},
        }
    )
    assert (
        df.to_markdown()
        == """
|    | gender   |   SUM(num) |   MAX(num) |
|---:|:---------|-----------:|-----------:|
|  0 | girl     |     118065 |       2588 |
|  1 | boy      |      47123 |       1280 |
    """.strip()
    )

    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|           |   ('SUM(num)',) |   ('MAX(num)',) |
|:----------|----------------:|----------------:|
| ('boy',)  |           47123 |            1280 |
| ('girl',) |          118065 |            2588 |
    """.strip()
    )

    # transpose_pivot
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                  |   ('SUM(num)', 'boy') |   ('SUM(num)', 'girl') |   ('MAX(num)', 'boy') |   ('MAX(num)', 'girl') |
|:-----------------|----------------------:|-----------------------:|----------------------:|-----------------------:|
| ('{_("Total")} (Sum)',) |                 47123 |                 118065 |                  1280 |                   2588 |
    """.strip()  # noqa: E501
    )

    # combine_metrics does nothing in this case
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|           |   ('SUM(num)',) |   ('MAX(num)',) |
|:----------|----------------:|----------------:|
| ('boy',)  |           47123 |            1280 |
| ('girl',) |          118065 |            2588 |
    """.strip()
    )

    # show totals
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                  |   ('SUM(num)',) |   ('MAX(num)',) |   ('{_("Total")} (Sum)',) |
|:-----------------|----------------:|----------------:|-------------------:|
| ('boy',)         |           47123 |            1280 |                nan |
| ('girl',)        |          118065 |            2588 |                nan |
| ('{_("Total")} (Sum)',) |          165188 |            3868 |                nan |
    """.strip()
    )

    # apply_metrics_on_rows
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                          |   ('{_("Total")} (Sum)',) |
|:-------------------------|-------------------:|
| ('SUM(num)', 'boy')      |              47123 |
| ('SUM(num)', 'girl')     |             118065 |
| ('SUM(num)', '{_("Subtotal")}') |             165188 |
| ('MAX(num)', 'boy')      |               1280 |
| ('MAX(num)', 'girl')     |               2588 |
| ('MAX(num)', '{_("Subtotal")}') |               3868 |
| ('{_("Total")} (Sum)', '')      |                nan |
    """.strip()
    )

    # apply_metrics_on_rows with combine_metrics
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                      |   ('{_("Total")} (Sum)',) |
|:---------------------|-------------------:|
| ('boy', 'SUM(num)')  |              47123 |
| ('boy', 'MAX(num)')  |               1280 |
| ('boy', '{_("Subtotal")}')  |                nan |
| ('girl', 'SUM(num)') |             118065 |
| ('girl', 'MAX(num)') |               2588 |
| ('girl', '{_("Subtotal")}') |                nan |
| ('{_("Total")} (Sum)', '')  |                nan |
    """.strip()
    )


def test_pivot_df_single_row_null_values():
    """
    Pivot table when a single column and 2 metrics are selected.
    """
    df = pd.DataFrame.from_dict(
        {
            "gender": {0: "girl", 1: "boy"},
            "SUM(num)": {0: 118065, 1: None},
            "MAX(num)": {0: 2588, 1: None},
        }
    )
    assert (
        df.to_markdown()
        == """
|    | gender   |   SUM(num) |   MAX(num) |
|---:|:---------|-----------:|-----------:|
|  0 | girl     |     118065 |       2588 |
|  1 | boy      |        nan |        nan |
    """.strip()
    )

    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|           |   ('SUM(num)',) |   ('MAX(num)',) |
|:----------|----------------:|----------------:|
| ('boy',)  |             nan |             nan |
| ('girl',) |          118065 |            2588 |
    """.strip()
    )

    # transpose_pivot
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  |   ('SUM(num)', 'boy') |   ('SUM(num)', 'girl') |   ('MAX(num)', 'boy') |   ('MAX(num)', 'girl') |
|:-----------------|----------------------:|-----------------------:|----------------------:|-----------------------:|
| ('Total (Sum)',) |                   nan |                 118065 |                   nan |                   2588 |
    """.strip()  # noqa: E501
    )

    # combine_metrics does nothing in this case
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|           |   ('SUM(num)',) |   ('MAX(num)',) |
|:----------|----------------:|----------------:|
| ('boy',)  |             nan |             nan |
| ('girl',) |          118065 |            2588 |
    """.strip()
    )

    # show totals
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  |   ('SUM(num)',) |   ('MAX(num)',) |   ('Total (Sum)',) |
|:-----------------|----------------:|----------------:|-------------------:|
| ('boy',)         |             nan |             nan |                nan |
| ('girl',)        |          118065 |            2588 |                nan |
| ('Total (Sum)',) |          118065 |            2588 |                nan |
    """.strip()
    )

    # apply_metrics_on_rows
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=True,
        show_columns_total=False,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                      |   ('{_("Total")} (Sum)',) |
|:---------------------|-------------------:|
| ('SUM(num)', 'boy')  |                nan |
| ('SUM(num)', 'girl') |             118065 |
| ('MAX(num)', 'boy')  |                nan |
| ('MAX(num)', 'girl') |               2588 |
    """.strip()
    )

    # apply_metrics_on_rows with combine_metrics
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                      |   ('{_("Total")} (Sum)',) |
|:---------------------|-------------------:|
| ('boy', 'SUM(num)')  |                nan |
| ('boy', 'MAX(num)')  |                nan |
| ('boy', '{_("Subtotal")}')  |                nan |
| ('girl', 'SUM(num)') |             118065 |
| ('girl', 'MAX(num)') |               2588 |
| ('girl', '{_("Subtotal")}') |                nan |
| ('{_("Total")} (Sum)', '')  |                nan |
    """.strip()
    )


def test_pivot_df_single_row_null_mix_values_strings():
    """
    Pivot table when a single column and 2 metrics are selected.
    """
    df = pd.DataFrame.from_dict(
        {
            "gender": {0: "girl", 1: "boy"},
            "SUM(num)": {0: 118065, 1: "NULL"},
            "MAX(num)": {0: 2588, 1: None},
        }
    )
    assert (
        df.to_markdown()
        == """
|    | gender   | SUM(num)   |   MAX(num) |
|---:|:---------|:-----------|-----------:|
|  0 | girl     | 118065     |       2588 |
|  1 | boy      | NULL       |        nan |
    """.strip()
    )

    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|           | ('SUM(num)',)   |   ('MAX(num)',) |
|:----------|:----------------|----------------:|
| ('boy',)  | NULL            |             nan |
| ('girl',) | 118065          |            2588 |
    """.strip()
    )

    # transpose_pivot
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  | ('SUM(num)', 'boy')   |   ('SUM(num)', 'girl') |   ('MAX(num)', 'boy') |   ('MAX(num)', 'girl') |
|:-----------------|:----------------------|-----------------------:|----------------------:|-----------------------:|
| ('Total (Sum)',) | NULL                  |                 118065 |                   nan |                   2588 |

    """.strip()  # noqa: E501
    )

    # combine_metrics does nothing in this case
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|           | ('SUM(num)',)   |   ('MAX(num)',) |
|:----------|:----------------|----------------:|
| ('boy',)  | NULL            |             nan |
| ('girl',) | 118065          |            2588 |
    """.strip()
    )

    # show totals
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  | ('SUM(num)',)   |   ('MAX(num)',) |   ('Total (Sum)',) |
|:-----------------|:----------------|----------------:|-------------------:|
| ('boy',)         | NULL            |             nan |                nan |
| ('girl',)        | 118065          |            2588 |                nan |
| ('Total (Sum)',) | 118065.0        |            2588 |                nan |
    """.strip()
    )

    # apply_metrics_on_rows with combine_metrics
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=True,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|                      | ('Total (Sum)',)   |
|:---------------------|:-------------------|
| ('boy', 'SUM(num)')  | NULL               |
| ('boy', 'MAX(num)')  | nan                |
| ('boy', 'Subtotal')  | nan                |
| ('girl', 'SUM(num)') | 118065             |
| ('girl', 'MAX(num)') | 2588.0             |
| ('girl', 'Subtotal') | nan                |
| ('Total (Sum)', '')  | nan                |
    """.strip()
    )


def test_pivot_df_single_row_null_mix_values_numbers():
    """
    Pivot table when a single column and 2 metrics are selected.
    """
    df = pd.DataFrame.from_dict(
        {
            "gender": {0: "girl", 1: "boy"},
            "SUM(num)": {0: 118065, 1: 21},
            "MAX(num)": {0: 2588, 1: None},
        }
    )
    assert (
        df.to_markdown()
        == """
|    | gender   |   SUM(num) |   MAX(num) |
|---:|:---------|-----------:|-----------:|
|  0 | girl     |     118065 |       2588 |
|  1 | boy      |         21 |        nan |
    """.strip()
    )

    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|           |   ('SUM(num)',) |   ('MAX(num)',) |
|:----------|----------------:|----------------:|
| ('boy',)  |              21 |             nan |
| ('girl',) |          118065 |            2588 |
    """.strip()
    )

    # transpose_pivot
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  |   ('SUM(num)', 'boy') |   ('SUM(num)', 'girl') |   ('MAX(num)', 'boy') |   ('MAX(num)', 'girl') |
|:-----------------|----------------------:|-----------------------:|----------------------:|-----------------------:|
| ('Total (Sum)',) |                    21 |                 118065 |                   nan |                   2588 |    """.strip()  # noqa: E501
    )

    # combine_metrics does nothing in this case
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|           |   ('SUM(num)',) |   ('MAX(num)',) |
|:----------|----------------:|----------------:|
| ('boy',)  |              21 |             nan |
| ('girl',) |          118065 |            2588 |
    """.strip()
    )

    # show totals
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  |   ('SUM(num)',) |   ('MAX(num)',) |
|:-----------------|----------------:|----------------:|
| ('boy',)         |              21 |             nan |
| ('girl',)        |          118065 |            2588 |
| ('Total (Sum)',) |          118086 |            2588 |
    """.strip()
    )

    # apply_metrics_on_rows
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=True,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|                          |   ('Total (Sum)',) |
|:-------------------------|-------------------:|
| ('SUM(num)', 'boy')      |                 21 |
| ('SUM(num)', 'girl')     |             118065 |
| ('SUM(num)', 'Subtotal') |             118086 |
| ('MAX(num)', 'boy')      |                nan |
| ('MAX(num)', 'girl')     |               2588 |
| ('MAX(num)', 'Subtotal') |               2588 |
| ('Total (Sum)', '')      |                nan |
    """.strip()
    )

    # apply_metrics_on_rows with combine_metrics
    pivoted = pivot_df(
        df,
        rows=["gender"],
        columns=[],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=True,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == f"""
|                      |   ('{_("Total")} (Sum)',) |
|:---------------------|-------------------:|
| ('boy', 'SUM(num)')  |                 21 |
| ('boy', 'MAX(num)')  |                nan |
| ('boy', '{_("Subtotal")}')  |                nan |
| ('girl', 'SUM(num)') |             118065 |
| ('girl', 'MAX(num)') |               2588 |
| ('girl', '{_("Subtotal")}') |                nan |
| ('{_("Total")} (Sum)', '')  |                nan |
    """.strip()
    )


def test_pivot_df_complex():
    """
    Pivot table when a column, rows and 2 metrics are selected.
    """
    df = pd.DataFrame.from_dict(
        {
            "state": {
                0: "CA",
                1: "CA",
                2: "CA",
                3: "FL",
                4: "CA",
                5: "CA",
                6: "FL",
                7: "FL",
                8: "FL",
                9: "CA",
                10: "FL",
                11: "FL",
            },
            "gender": {
                0: "girl",
                1: "boy",
                2: "girl",
                3: "girl",
                4: "girl",
                5: "girl",
                6: "boy",
                7: "girl",
                8: "girl",
                9: "boy",
                10: "boy",
                11: "girl",
            },
            "name": {
                0: "Amy",
                1: "Edward",
                2: "Sophia",
                3: "Amy",
                4: "Cindy",
                5: "Dawn",
                6: "Edward",
                7: "Sophia",
                8: "Dawn",
                9: "Tony",
                10: "Tony",
                11: "Cindy",
            },
            "SUM(num)": {
                0: 45426,
                1: 31290,
                2: 18859,
                3: 14740,
                4: 14149,
                5: 11403,
                6: 9395,
                7: 7181,
                8: 5089,
                9: 3765,
                10: 2673,
                11: 1218,
            },
            "MAX(num)": {
                0: 2227,
                1: 1280,
                2: 2588,
                3: 854,
                4: 842,
                5: 1157,
                6: 389,
                7: 1187,
                8: 461,
                9: 598,
                10: 247,
                11: 217,
            },
        }
    )
    assert (
        df.to_markdown()
        == """
|    | state   | gender   | name   |   SUM(num) |   MAX(num) |
|---:|:--------|:---------|:-------|-----------:|-----------:|
|  0 | CA      | girl     | Amy    |      45426 |       2227 |
|  1 | CA      | boy      | Edward |      31290 |       1280 |
|  2 | CA      | girl     | Sophia |      18859 |       2588 |
|  3 | FL      | girl     | Amy    |      14740 |        854 |
|  4 | CA      | girl     | Cindy  |      14149 |        842 |
|  5 | CA      | girl     | Dawn   |      11403 |       1157 |
|  6 | FL      | boy      | Edward |       9395 |        389 |
|  7 | FL      | girl     | Sophia |       7181 |       1187 |
|  8 | FL      | girl     | Dawn   |       5089 |        461 |
|  9 | CA      | boy      | Tony   |       3765 |        598 |
| 10 | FL      | boy      | Tony   |       2673 |        247 |
| 11 | FL      | girl     | Cindy  |       1218 |        217 |
    """.strip()
    )

    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )

    # Sort the pivoted DataFrame to ensure deterministic output
    pivoted_sorted = pivoted.sort_index()

    assert (
        pivoted_sorted.to_markdown()
        == """
|                    |   ('SUM(num)', 'CA') |   ('SUM(num)', 'FL') |   ('MAX(num)', 'CA') |   ('MAX(num)', 'FL') |
|:-------------------|---------------------:|---------------------:|---------------------:|---------------------:|
| ('boy', 'Edward')  |                31290 |                 9395 |                 1280 |                  389 |
| ('boy', 'Tony')    |                 3765 |                 2673 |                  598 |                  247 |
| ('girl', 'Amy')    |                45426 |                14740 |                 2227 |                  854 |
| ('girl', 'Cindy')  |                14149 |                 1218 |                  842 |                  217 |
| ('girl', 'Dawn')   |                11403 |                 5089 |                 1157 |                  461 |
| ('girl', 'Sophia') |                18859 |                 7181 |                 2588 |                 1187 |
    """.strip()  # noqa: E501
    )

    # transpose_pivot
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|         |   ('SUM(num)', 'boy', 'Edward') |   ('SUM(num)', 'boy', 'Tony') |   ('SUM(num)', 'girl', 'Amy') |   ('SUM(num)', 'girl', 'Cindy') |   ('SUM(num)', 'girl', 'Dawn') |   ('SUM(num)', 'girl', 'Sophia') |   ('MAX(num)', 'boy', 'Edward') |   ('MAX(num)', 'boy', 'Tony') |   ('MAX(num)', 'girl', 'Amy') |   ('MAX(num)', 'girl', 'Cindy') |   ('MAX(num)', 'girl', 'Dawn') |   ('MAX(num)', 'girl', 'Sophia') |
|:--------|--------------------------------:|------------------------------:|------------------------------:|--------------------------------:|-------------------------------:|---------------------------------:|--------------------------------:|------------------------------:|------------------------------:|--------------------------------:|-------------------------------:|---------------------------------:|
| ('CA',) |                           31290 |                          3765 |                         45426 |                           14149 |                          11403 |                            18859 |                            1280 |                           598 |                          2227 |                             842 |                           1157 |                             2588 |
| ('FL',) |                            9395 |                          2673 |                         14740 |                            1218 |                           5089 |                             7181 |                             389 |                           247 |                           854 |                             217 |                            461 |                             1187 |
    """.strip()  # noqa: E501
    )

    # combine_metrics
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                    |   ('CA', 'SUM(num)') |   ('CA', 'MAX(num)') |   ('FL', 'SUM(num)') |   ('FL', 'MAX(num)') |
|:-------------------|---------------------:|---------------------:|---------------------:|---------------------:|
| ('boy', 'Edward')  |                31290 |                 1280 |                 9395 |                  389 |
| ('boy', 'Tony')    |                 3765 |                  598 |                 2673 |                  247 |
| ('girl', 'Amy')    |                45426 |                 2227 |                14740 |                  854 |
| ('girl', 'Cindy')  |                14149 |                  842 |                 1218 |                  217 |
| ('girl', 'Dawn')   |                11403 |                 1157 |                 5089 |                  461 |
| ('girl', 'Sophia') |                18859 |                 2588 |                 7181 |                 1187 |
    """.strip()  # noqa: E501
    )

    # show totals
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                      |   ('SUM(num)', 'CA') |   ('SUM(num)', 'FL') |   ('SUM(num)', 'Subtotal') |   ('MAX(num)', 'CA') |   ('MAX(num)', 'FL') |   ('MAX(num)', 'Subtotal') |   ('Total (Sum)', '') |
|:---------------------|---------------------:|---------------------:|---------------------------:|---------------------:|---------------------:|---------------------------:|----------------------:|
| ('boy', 'Edward')    |                31290 |                 9395 |                      40685 |                 1280 |                  389 |                       1669 |                   nan |
| ('boy', 'Tony')      |                 3765 |                 2673 |                       6438 |                  598 |                  247 |                        845 |                   nan |
| ('boy', 'Subtotal')  |                35055 |                12068 |                      47123 |                 1878 |                  636 |                       2514 |                   nan |
| ('girl', 'Amy')      |                45426 |                14740 |                      60166 |                 2227 |                  854 |                       3081 |                   nan |
| ('girl', 'Cindy')    |                14149 |                 1218 |                      15367 |                  842 |                  217 |                       1059 |                   nan |
| ('girl', 'Dawn')     |                11403 |                 5089 |                      16492 |                 1157 |                  461 |                       1618 |                   nan |
| ('girl', 'Sophia')   |                18859 |                 7181 |                      26040 |                 2588 |                 1187 |                       3775 |                   nan |
| ('girl', 'Subtotal') |                89837 |                28228 |                     118065 |                 6814 |                 2719 |                       9533 |                   nan |
| ('Total (Sum)', '')  |               124892 |                40296 |                     165188 |                 8692 |                 3355 |                      12047 |                   nan |
    """.strip()  # noqa: E501
    )

    # apply_metrics_on_rows
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|                                |   ('CA',) |   ('FL',) |
|:-------------------------------|----------:|----------:|
| ('SUM(num)', 'boy', 'Edward')  |     31290 |      9395 |
| ('SUM(num)', 'boy', 'Tony')    |      3765 |      2673 |
| ('SUM(num)', 'girl', 'Amy')    |     45426 |     14740 |
| ('SUM(num)', 'girl', 'Cindy')  |     14149 |      1218 |
| ('SUM(num)', 'girl', 'Dawn')   |     11403 |      5089 |
| ('SUM(num)', 'girl', 'Sophia') |     18859 |      7181 |
| ('MAX(num)', 'boy', 'Edward')  |      1280 |       389 |
| ('MAX(num)', 'boy', 'Tony')    |       598 |       247 |
| ('MAX(num)', 'girl', 'Amy')    |      2227 |       854 |
| ('MAX(num)', 'girl', 'Cindy')  |       842 |       217 |
| ('MAX(num)', 'girl', 'Dawn')   |      1157 |       461 |
| ('MAX(num)', 'girl', 'Sophia') |      2588 |      1187 |
    """.strip()
    )

    # apply_metrics_on_rows with combine_metrics
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|                                |   ('CA',) |   ('FL',) |
|:-------------------------------|----------:|----------:|
| ('boy', 'Edward', 'SUM(num)')  |     31290 |      9395 |
| ('boy', 'Edward', 'MAX(num)')  |      1280 |       389 |
| ('boy', 'Tony', 'SUM(num)')    |      3765 |      2673 |
| ('boy', 'Tony', 'MAX(num)')    |       598 |       247 |
| ('girl', 'Amy', 'SUM(num)')    |     45426 |     14740 |
| ('girl', 'Amy', 'MAX(num)')    |      2227 |       854 |
| ('girl', 'Cindy', 'SUM(num)')  |     14149 |      1218 |
| ('girl', 'Cindy', 'MAX(num)')  |       842 |       217 |
| ('girl', 'Dawn', 'SUM(num)')   |     11403 |      5089 |
| ('girl', 'Dawn', 'MAX(num)')   |      1157 |       461 |
| ('girl', 'Sophia', 'SUM(num)') |     18859 |      7181 |
| ('girl', 'Sophia', 'MAX(num)') |      2588 |      1187 |
    """.strip()
    )

    # everything
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=True,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|                     |   ('boy', 'Edward') |   ('boy', 'Tony') |   ('boy', 'Subtotal') |   ('girl', 'Amy') |   ('girl', 'Cindy') |   ('girl', 'Dawn') |   ('girl', 'Sophia') |   ('girl', 'Subtotal') |   ('Total (Sum)', '') |
|:--------------------|--------------------:|------------------:|----------------------:|------------------:|--------------------:|-------------------:|---------------------:|-----------------------:|----------------------:|
| ('CA', 'SUM(num)')  |               31290 |              3765 |                 35055 |             45426 |               14149 |              11403 |                18859 |                  89837 |                124892 |
| ('CA', 'MAX(num)')  |                1280 |               598 |                  1878 |              2227 |                 842 |               1157 |                 2588 |                   6814 |                  8692 |
| ('CA', 'Subtotal')  |                 nan |               nan |                   nan |               nan |                 nan |                nan |                  nan |                    nan |                   nan |
| ('FL', 'SUM(num)')  |                9395 |              2673 |                 12068 |             14740 |                1218 |               5089 |                 7181 |                  28228 |                 40296 |
| ('FL', 'MAX(num)')  |                 389 |               247 |                   636 |               854 |                 217 |                461 |                 1187 |                   2719 |                  3355 |
| ('FL', 'Subtotal')  |                 nan |               nan |                   nan |               nan |                 nan |                nan |                  nan |                    nan |                   nan |
| ('Total (Sum)', '') |                 nan |               nan |                   nan |               nan |                 nan |                nan |                  nan |                    nan |                   nan |
    """.strip()  # noqa: E501
    )

    # fraction
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum as Fraction of Columns",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                                            |   ('SUM(num)', 'CA') |   ('SUM(num)', 'FL') |   ('MAX(num)', 'CA') |   ('MAX(num)', 'FL') |
|:-------------------------------------------|---------------------:|---------------------:|---------------------:|---------------------:|
| ('boy', 'Edward')                          |            0.250536  |            0.23315   |            0.147262  |            0.115946  |
| ('boy', 'Tony')                            |            0.030146  |            0.0663341 |            0.0687989 |            0.0736215 |
| ('boy', 'Subtotal')                        |            0.280683  |            0.299484  |            0.216061  |            0.189568  |
| ('girl', 'Amy')                            |            0.363722  |            0.365793  |            0.256213  |            0.254545  |
| ('girl', 'Cindy')                          |            0.11329   |            0.0302263 |            0.0968707 |            0.0646796 |
| ('girl', 'Dawn')                           |            0.0913029 |            0.12629   |            0.133111  |            0.137407  |
| ('girl', 'Sophia')                         |            0.151002  |            0.178206  |            0.297745  |            0.3538    |
| ('girl', 'Subtotal')                       |            0.719317  |            0.700516  |            0.783939  |            0.810432  |
| ('Total (Sum as Fraction of Columns)', '') |            1         |            1         |            1         |            1         |
    """.strip()  # noqa: E501
    )


def test_pivot_df_multi_column():
    """
    Pivot table when 2 columns, no rows and 2 metrics are selected.
    """
    df = pd.DataFrame.from_dict(
        {
            "state": {
                0: "CA",
                1: "CA",
                2: "CA",
                3: "FL",
                4: "CA",
                5: "CA",
                6: "FL",
                7: "FL",
                8: "FL",
                9: "CA",
                10: "FL",
                11: "FL",
            },
            "gender": {
                0: "girl",
                1: "boy",
                2: "girl",
                3: "girl",
                4: "girl",
                5: "girl",
                6: "boy",
                7: "girl",
                8: "girl",
                9: "boy",
                10: "boy",
                11: "girl",
            },
            "SUM(num)": {
                0: 45426,
                1: 31290,
                2: 18859,
                3: 14740,
                4: 14149,
                5: 11403,
                6: 9395,
                7: 7181,
                8: 5089,
                9: 3765,
                10: 2673,
                11: 1218,
            },
            "MAX(num)": {
                0: 2227,
                1: 1280,
                2: 2588,
                3: 854,
                4: 842,
                5: 1157,
                6: 389,
                7: 1187,
                8: 461,
                9: 598,
                10: 247,
                11: 217,
            },
        }
    )

    pivoted = pivot_df(
        df,
        rows=None,
        columns=["state", "gender"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  |   ('SUM(num)', 'boy') |   ('SUM(num)', 'girl') |   ('MAX(num)', 'boy') |   ('MAX(num)', 'girl') |
|:-----------------|----------------------:|-----------------------:|----------------------:|-----------------------:|
| ('CA',)          |                 35055 |                  89837 |                  1878 |                   6814 |
| ('Total (Sum)',) |                 12068 |                  28228 |                   636 |                   2719 |
    """.strip()  # noqa: E501
    )

    # transpose_pivot
    pivoted = pivot_df(
        df,
        rows=None,
        columns=["state", "gender"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                |   ('SUM(num)',) |   ('MAX(num)',) |
|:---------------|----------------:|----------------:|
| ('CA', 'boy')  |           35055 |            1878 |
| ('CA', 'girl') |           89837 |            6814 |
| ('FL', 'boy')  |           12068 |             636 |
| ('FL', 'girl') |           28228 |            2719 |
    """.strip()
    )

    # combine_metrics
    pivoted = pivot_df(
        df,
        rows=None,
        columns=["state", "gender"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  |   ('boy', 'SUM(num)') |   ('boy', 'MAX(num)') |   ('girl', 'SUM(num)') |   ('girl', 'MAX(num)') |
|:-----------------|----------------------:|----------------------:|-----------------------:|-----------------------:|
| ('CA',)          |                 35055 |                  1878 |                  89837 |                   6814 |
| ('Total (Sum)',) |                 12068 |                   636 |                  28228 |                   2719 |
    """.strip()  # noqa: E501
    )

    # show totals
    pivoted = pivot_df(
        df,
        rows=None,
        columns=["state", "gender"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                  |   ('SUM(num)', 'boy') |   ('SUM(num)', 'girl') |   ('SUM(num)', 'Subtotal') |   ('MAX(num)', 'boy') |   ('MAX(num)', 'girl') |   ('MAX(num)', 'Subtotal') |   ('Total (Sum)', '') |
|:-----------------|----------------------:|-----------------------:|---------------------------:|----------------------:|-----------------------:|---------------------------:|----------------------:|
| ('CA',)          |                 35055 |                  89837 |                     124892 |                  1878 |                   6814 |                       8692 |                   nan |
| ('Total (Sum)',) |                 12068 |                  28228 |                      40296 |                   636 |                   2719 |                       3355 |                   nan |
    """.strip()  # noqa: E501
    )

    # apply_metrics_on_rows
    pivoted = pivot_df(
        df,
        rows=None,
        columns=["state", "gender"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|               |   ('CA', 'boy') |   ('CA', 'girl') |   ('FL', 'boy') |   ('FL', 'girl') |
|:--------------|----------------:|-----------------:|----------------:|-----------------:|
| ('SUM(num)',) |           35055 |            89837 |           12068 |            28228 |
| ('MAX(num)',) |            1878 |             6814 |             636 |             2719 |
    """.strip()  # noqa: E501
    )

    # apply_metrics_on_rows with combine_metrics
    pivoted = pivot_df(
        df,
        rows=None,
        columns=["state", "gender"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|               |   ('CA', 'boy') |   ('CA', 'girl') |   ('FL', 'boy') |   ('FL', 'girl') |
|:--------------|----------------:|-----------------:|----------------:|-----------------:|
| ('SUM(num)',) |           35055 |            89837 |           12068 |            28228 |
| ('MAX(num)',) |            1878 |             6814 |             636 |             2719 |
    """.strip()  # noqa: E501
    )

    # everything
    pivoted = pivot_df(
        df,
        rows=None,
        columns=["state", "gender"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=True,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|                      |   ('CA',) |   ('Total (Sum)',) |
|:---------------------|----------:|-------------------:|
| ('boy', 'SUM(num)')  |     35055 |              12068 |
| ('boy', 'MAX(num)')  |      1878 |                636 |
| ('boy', 'Subtotal')  |       nan |                nan |
| ('girl', 'SUM(num)') |     89837 |              28228 |
| ('girl', 'MAX(num)') |      6814 |               2719 |
| ('girl', 'Subtotal') |       nan |                nan |
| ('Total (Sum)', '')  |       nan |                nan |
    """.strip()
    )

    # fraction
    pivoted = pivot_df(
        df,
        rows=None,
        columns=["state", "gender"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum as Fraction of Columns",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                                         |   ('SUM(num)', 'boy') |   ('SUM(num)', 'girl') |   ('MAX(num)', 'boy') |   ('MAX(num)', 'girl') |
|:----------------------------------------|----------------------:|-----------------------:|----------------------:|-----------------------:|
| ('CA',)                                 |              0.743904 |               0.760911 |              0.747017 |                0.71478 |
| ('Total (Sum as Fraction of Columns)',) |              0.256096 |               0.239089 |              0.252983 |                0.28522 |
    """.strip()  # noqa: E501
    )


def test_pivot_df_complex_null_values():
    """
    Pivot table when a column, rows and 2 metrics are selected.
    """
    df = pd.DataFrame.from_dict(
        {
            "state": {
                0: None,
                1: None,
                2: None,
                3: None,
                4: None,
                5: None,
                6: None,
                7: None,
                8: None,
                9: None,
                10: None,
                11: None,
            },
            "gender": {
                0: "girl",
                1: "boy",
                2: "girl",
                3: "girl",
                4: "girl",
                5: "girl",
                6: "boy",
                7: "girl",
                8: "girl",
                9: "boy",
                10: "boy",
                11: "girl",
            },
            "name": {
                0: "Amy",
                1: "Edward",
                2: "Sophia",
                3: "Amy",
                4: "Cindy",
                5: "Dawn",
                6: "Edward",
                7: "Sophia",
                8: "Dawn",
                9: "Tony",
                10: "Tony",
                11: "Cindy",
            },
            "SUM(num)": {
                0: 45426,
                1: 31290,
                2: 18859,
                3: 14740,
                4: 14149,
                5: 11403,
                6: 9395,
                7: 7181,
                8: 5089,
                9: 3765,
                10: 2673,
                11: 1218,
            },
            "MAX(num)": {
                0: 2227,
                1: 1280,
                2: 2588,
                3: 854,
                4: 842,
                5: 1157,
                6: 389,
                7: 1187,
                8: 461,
                9: 598,
                10: 247,
                11: 217,
            },
        }
    )
    assert (
        df.to_markdown()
        == """
|    | state   | gender   | name   |   SUM(num) |   MAX(num) |
|---:|:--------|:---------|:-------|-----------:|-----------:|
|  0 |         | girl     | Amy    |      45426 |       2227 |
|  1 |         | boy      | Edward |      31290 |       1280 |
|  2 |         | girl     | Sophia |      18859 |       2588 |
|  3 |         | girl     | Amy    |      14740 |        854 |
|  4 |         | girl     | Cindy  |      14149 |        842 |
|  5 |         | girl     | Dawn   |      11403 |       1157 |
|  6 |         | boy      | Edward |       9395 |        389 |
|  7 |         | girl     | Sophia |       7181 |       1187 |
|  8 |         | girl     | Dawn   |       5089 |        461 |
|  9 |         | boy      | Tony   |       3765 |        598 |
| 10 |         | boy      | Tony   |       2673 |        247 |
| 11 |         | girl     | Cindy  |       1218 |        217 |
    """.strip()
    )

    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                    |   ('SUM(num)', nan) |   ('MAX(num)', nan) |
|:-------------------|--------------------:|--------------------:|
| ('boy', 'Edward')  |               40685 |                1669 |
| ('boy', 'Tony')    |                6438 |                 845 |
| ('girl', 'Amy')    |               60166 |                3081 |
| ('girl', 'Cindy')  |               15367 |                1059 |
| ('girl', 'Dawn')   |               16492 |                1618 |
| ('girl', 'Sophia') |               26040 |                3775 |

    """.strip()
    )

    # transpose_pivot
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|        |   ('SUM(num)', 'boy', 'Edward') |   ('SUM(num)', 'boy', 'Tony') |   ('SUM(num)', 'girl', 'Amy') |   ('SUM(num)', 'girl', 'Cindy') |   ('SUM(num)', 'girl', 'Dawn') |   ('SUM(num)', 'girl', 'Sophia') |   ('MAX(num)', 'boy', 'Edward') |   ('MAX(num)', 'boy', 'Tony') |   ('MAX(num)', 'girl', 'Amy') |   ('MAX(num)', 'girl', 'Cindy') |   ('MAX(num)', 'girl', 'Dawn') |   ('MAX(num)', 'girl', 'Sophia') |
|:-------|--------------------------------:|------------------------------:|------------------------------:|--------------------------------:|-------------------------------:|---------------------------------:|--------------------------------:|------------------------------:|------------------------------:|--------------------------------:|-------------------------------:|---------------------------------:|
| (nan,) |                           40685 |                          6438 |                         60166 |                           15367 |                          16492 |                            26040 |                            1669 |                           845 |                          3081 |                            1059 |                           1618 |                             3775 |
    """.strip()  # noqa: E501
    )

    # combine_metrics
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                    |   (nan, 'SUM(num)') |   (nan, 'MAX(num)') |
|:-------------------|--------------------:|--------------------:|
| ('boy', 'Edward')  |               40685 |                1669 |
| ('boy', 'Tony')    |                6438 |                 845 |
| ('girl', 'Amy')    |               60166 |                3081 |
| ('girl', 'Cindy')  |               15367 |                1059 |
| ('girl', 'Dawn')   |               16492 |                1618 |
| ('girl', 'Sophia') |               26040 |                3775 |
 """.strip()
    )

    # show totals
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                      |   ('SUM(num)', nan) |   ('SUM(num)', 'Subtotal') |   ('MAX(num)', nan) |   ('MAX(num)', 'Subtotal') |   ('Total (Sum)', '') |
|:---------------------|--------------------:|---------------------------:|--------------------:|---------------------------:|----------------------:|
| ('boy', 'Edward')    |               40685 |                      40685 |                1669 |                       1669 |                   nan |
| ('boy', 'Tony')      |                6438 |                       6438 |                 845 |                        845 |                   nan |
| ('boy', 'Subtotal')  |               47123 |                      47123 |                2514 |                       2514 |                   nan |
| ('girl', 'Amy')      |               60166 |                      60166 |                3081 |                       3081 |                   nan |
| ('girl', 'Cindy')    |               15367 |                      15367 |                1059 |                       1059 |                   nan |
| ('girl', 'Dawn')     |               16492 |                      16492 |                1618 |                       1618 |                   nan |
| ('girl', 'Sophia')   |               26040 |                      26040 |                3775 |                       3775 |                   nan |
| ('girl', 'Subtotal') |              118065 |                     118065 |                9533 |                       9533 |                   nan |
| ('Total (Sum)', '')  |              165188 |                     165188 |               12047 |                      12047 |                   nan |
    """.strip()  # noqa: E501
    )

    # apply_metrics_on_rows
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|                                |   (nan,) |
|:-------------------------------|---------:|
| ('SUM(num)', 'boy', 'Edward')  |    40685 |
| ('SUM(num)', 'boy', 'Tony')    |     6438 |
| ('SUM(num)', 'girl', 'Amy')    |    60166 |
| ('SUM(num)', 'girl', 'Cindy')  |    15367 |
| ('SUM(num)', 'girl', 'Dawn')   |    16492 |
| ('SUM(num)', 'girl', 'Sophia') |    26040 |
| ('MAX(num)', 'boy', 'Edward')  |     1669 |
| ('MAX(num)', 'boy', 'Tony')    |      845 |
| ('MAX(num)', 'girl', 'Amy')    |     3081 |
| ('MAX(num)', 'girl', 'Cindy')  |     1059 |
| ('MAX(num)', 'girl', 'Dawn')   |     1618 |
| ('MAX(num)', 'girl', 'Sophia') |     3775 |
    """.strip()
    )

    # apply_metrics_on_rows with combine_metrics
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=True,
        show_rows_total=False,
        show_columns_total=False,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|                                |   (nan,) |
|:-------------------------------|---------:|
| ('boy', 'Edward', 'SUM(num)')  |    40685 |
| ('boy', 'Edward', 'MAX(num)')  |     1669 |
| ('boy', 'Tony', 'SUM(num)')    |     6438 |
| ('boy', 'Tony', 'MAX(num)')    |      845 |
| ('girl', 'Amy', 'SUM(num)')    |    60166 |
| ('girl', 'Amy', 'MAX(num)')    |     3081 |
| ('girl', 'Cindy', 'SUM(num)')  |    15367 |
| ('girl', 'Cindy', 'MAX(num)')  |     1059 |
| ('girl', 'Dawn', 'SUM(num)')   |    16492 |
| ('girl', 'Dawn', 'MAX(num)')   |     1618 |
| ('girl', 'Sophia', 'SUM(num)') |    26040 |
| ('girl', 'Sophia', 'MAX(num)') |     3775 |
    """.strip()
    )

    # everything
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum",
        transpose_pivot=True,
        combine_metrics=True,
        show_rows_total=True,
        show_columns_total=True,
        apply_metrics_on_rows=True,
    )
    assert (
        pivoted.to_markdown()
        == """
|                     |   ('boy', 'Edward') |   ('boy', 'Tony') |   ('boy', 'Subtotal') |   ('girl', 'Amy') |   ('girl', 'Cindy') |   ('girl', 'Dawn') |   ('girl', 'Sophia') |   ('girl', 'Subtotal') |   ('Total (Sum)', '') |
|:--------------------|--------------------:|------------------:|----------------------:|------------------:|--------------------:|-------------------:|---------------------:|-----------------------:|----------------------:|
| (nan, 'SUM(num)')   |               40685 |              6438 |                 47123 |             60166 |               15367 |              16492 |                26040 |                 118065 |                165188 |
| (nan, 'MAX(num)')   |                1669 |               845 |                  2514 |              3081 |                1059 |               1618 |                 3775 |                   9533 |                 12047 |
| (nan, 'Subtotal')   |                 nan |               nan |                   nan |               nan |                 nan |                nan |                  nan |                    nan |                   nan |
| ('Total (Sum)', '') |                 nan |               nan |                   nan |               nan |                 nan |                nan |                  nan |                    nan |                   nan |
    """.strip()  # noqa: E501
    )

    # fraction
    pivoted = pivot_df(
        df,
        rows=["gender", "name"],
        columns=["state"],
        metrics=["SUM(num)", "MAX(num)"],
        aggfunc="Sum as Fraction of Columns",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )
    assert (
        pivoted.to_markdown()
        == """
|                                            |   ('SUM(num)', nan) |   ('MAX(num)', nan) |
|:-------------------------------------------|--------------------:|--------------------:|
| ('boy', 'Edward')                          |           0.246295  |           0.138541  |
| ('boy', 'Tony')                            |           0.0389738 |           0.0701419 |
| ('boy', 'Subtotal')                        |           0.285269  |           0.208683  |
| ('girl', 'Amy')                            |           0.364227  |           0.255748  |
| ('girl', 'Cindy')                          |           0.0930273 |           0.0879057 |
| ('girl', 'Dawn')                           |           0.0998378 |           0.134307  |
| ('girl', 'Sophia')                         |           0.157639  |           0.313356  |
| ('girl', 'Subtotal')                       |           0.714731  |           0.791317  |
| ('Total (Sum as Fraction of Columns)', '') |           1         |           1         |
    """.strip()  # noqa: E501
    )


# --- `showValuesAs` percent modes (#42809) -----------------------------------
#
# Exports and scheduled reports render server-side, so they have to reproduce
# the client's `fractionOf` aggregator: each cell over its row, column, or grand
# total, computed per metric, with totals dividing by their own rollup rather
# than summing the fractions around them.

SHOW_VALUES_AS_OPTIONS: dict[str, Any] = {
    "rows": ["nation"],
    "columns": ["gender"],
    "metrics": ["SUM(num)"],
    "aggfunc": "Sum",
    "transpose_pivot": False,
    "combine_metrics": False,
    "show_rows_total": True,
    "show_columns_total": True,
    "apply_metrics_on_rows": False,
}


def show_values_as_df() -> pd.DataFrame:
    """A 2x2 pivot: row totals 40/40, column totals 30/50, grand total 80."""
    return pd.DataFrame(
        {
            "nation": ["US", "US", "UK", "UK"],
            "gender": ["boy", "girl", "boy", "girl"],
            "SUM(num)": [10, 30, 20, 20],
        }
    )


def total_label() -> str:
    return f"{_('Total')} (Sum)"


def test_pivot_df_show_values_as_percent_row():
    pivoted = pivot_df(
        show_values_as_df(), **SHOW_VALUES_AS_OPTIONS, show_values_as="percent_row"
    )
    total = total_label()

    assert pivoted.loc[("US",), ("SUM(num)", "boy")] == 0.25
    assert pivoted.loc[("US",), ("SUM(num)", "girl")] == 0.75
    assert pivoted.loc[("UK",), ("SUM(num)", "boy")] == 0.5
    # a row is always 100% of itself
    assert pivoted.loc[("US",), (total, "")] == 1
    # the totals row shows each column's share of the grand total (30/80 and
    # 50/80), not the sum of the fractions above it
    assert pivoted.loc[(total,), ("SUM(num)", "boy")] == 0.375
    assert pivoted.loc[(total,), ("SUM(num)", "girl")] == 0.625
    assert pivoted.loc[(total,), (total, "")] == 1


def test_pivot_df_show_values_as_percent_col():
    pivoted = pivot_df(
        show_values_as_df(), **SHOW_VALUES_AS_OPTIONS, show_values_as="percent_col"
    )
    total = total_label()

    assert pivoted.loc[("US",), ("SUM(num)", "boy")] == pytest.approx(1 / 3)
    assert pivoted.loc[("UK",), ("SUM(num)", "boy")] == pytest.approx(2 / 3)
    assert pivoted.loc[("US",), ("SUM(num)", "girl")] == 0.6
    # a column is always 100% of itself
    assert pivoted.loc[(total,), ("SUM(num)", "boy")] == 1
    # the totals column shows each row's share of the grand total
    assert pivoted.loc[("US",), (total, "")] == 0.5


def test_pivot_df_show_values_as_percent_total():
    pivoted = pivot_df(
        show_values_as_df(), **SHOW_VALUES_AS_OPTIONS, show_values_as="percent_total"
    )
    total = total_label()

    assert pivoted.loc[("US",), ("SUM(num)", "boy")] == 0.125
    assert pivoted.loc[("US",), ("SUM(num)", "girl")] == 0.375
    assert pivoted.loc[("UK",), ("SUM(num)", "boy")] == 0.25
    assert pivoted.loc[(total,), ("SUM(num)", "boy")] == 0.375
    assert pivoted.loc[("US",), (total, "")] == 0.5
    assert pivoted.loc[(total,), (total, "")] == 1


def test_pivot_df_show_values_as_keeps_metrics_separate():
    """One metric's cells are never divided by another metric's total."""
    df = show_values_as_df()
    df["MAX(num)"] = [1, 3, 6, 10]
    pivoted = pivot_df(
        df,
        **{**SHOW_VALUES_AS_OPTIONS, "metrics": ["SUM(num)", "MAX(num)"]},
        show_values_as="percent_row",
    )

    assert pivoted.loc[("US",), ("SUM(num)", "boy")] == 0.25
    assert pivoted.loc[("UK",), ("MAX(num)", "boy")] == 0.375
    assert pivoted.loc[("UK",), ("MAX(num)", "girl")] == 0.625
    # a total collapsing the metric axis mixes SUM(num) and MAX(num), which
    # share no unit, so the renderer blanks it (#44657) instead of picking one
    # arbitrarily (#44725) -- exported values must agree with what it shows
    assert pd.isna(pivoted.loc[("US",), (total_label(), "")])


def test_pivot_df_show_values_as_blanks_mixed_metric_grand_total():
    """https://github.com/apache/superset/issues/44725

    A pivot with 2+ metrics on columns and row/column totals enabled has no
    single meaningful value for the cell where both the row and column totals
    meet (it would mix every configured metric). Before #44657 the renderer
    picked one metric there; the export/report path must match its current,
    blanked behavior rather than showing a specific, plausible-looking number.
    """
    df = show_values_as_df()
    df["MAX(num)"] = [1, 3, 6, 10]
    pivoted = pivot_df(
        df,
        **{**SHOW_VALUES_AS_OPTIONS, "metrics": ["SUM(num)", "MAX(num)"]},
        show_values_as="percent_total",
    )

    assert pd.isna(pivoted.loc[("UK",), (total_label(), "")])
    assert pd.isna(pivoted.loc[("US",), (total_label(), "")])
    assert pd.isna(pivoted.loc[(total_label(),), (total_label(), "")])


@pytest.mark.parametrize("layout", ["COLUMNS", "ROWS"])
@pytest.mark.parametrize("with_rollups", [False, True])
@pytest.mark.parametrize("mode", [None, "percent_row", "percent_col", "percent_total"])
@pytest.mark.parametrize("aggregation", ["Sum", "Average"])
def test_pivot_table_v2_blanks_mixed_metric_totals(
    layout: str, with_rollups: bool, mode: str | None, aggregation: str
):
    """Mixed totals stay blank in both layouts, with or without GROUPING SETS."""
    if with_rollups:
        df = grouping_sets_df()
        df["MAX(num)"] = [1, 6, 6, 1, 6, 6]
        metrics = ["AVG(num)", "MAX(num)"]
    else:
        df = show_values_as_df()
        df["MAX(num)"] = [1, 3, 6, 10]
        metrics = ["SUM(num)", "MAX(num)"]
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": metrics,
        "metricsLayout": layout,
        "rowTotals": True,
        "colTotals": True,
        "showValuesAs": mode,
        "aggregateFunction": aggregation,
    }
    pivoted = pivot_table_v2(df, form_data, apply_number_format=False)
    total = f"{_('Total')} ({aggregation})"
    if layout == "ROWS":
        mixed = pivoted.loc[(total, ""), :]
        leaf = pivoted.loc[(metrics[0], "US"), ("boy",)]
        metric_total = pivoted.loc[(metrics[0], "Subtotal"), (total,)]
    else:
        mixed = pivoted.loc[:, (total, "")]
        leaf = pivoted.loc[("US",), (metrics[0], "boy")]
        metric_total = pivoted.loc[(total,), (metrics[0], "Subtotal")]
    # Includes the grand-total corner: all-blank sums must not revive it as 0.
    assert mixed.isna().all()
    assert pd.notna(leaf)
    assert pd.notna(metric_total)
    if mode is None:
        assert leaf == 10
        assert metric_total == (
            19 if with_rollups else (80 if aggregation == "Sum" else 20)
        )


def test_pivot_df_show_values_as_with_combined_metrics():
    """`combineMetric` moves the metric to the lowest column level."""
    df = show_values_as_df()
    df["MAX(num)"] = [1, 3, 6, 10]
    pivoted = pivot_df(
        df,
        **{
            **SHOW_VALUES_AS_OPTIONS,
            "metrics": ["SUM(num)", "MAX(num)"],
            "combine_metrics": True,
            "show_rows_total": False,
            "show_columns_total": False,
        },
        show_values_as="percent_row",
    )

    assert pivoted.loc[("US",), ("boy", "SUM(num)")] == 0.25
    assert pivoted.loc[("US",), ("girl", "SUM(num)")] == 0.75
    # each metric keeps its own denominator across the combined layout
    assert pivoted.loc[("UK",), ("boy", "MAX(num)")] == 0.375
    assert pivoted.loc[("UK",), ("girl", "MAX(num)")] == 0.625


def test_pivot_table_v2_show_values_as_uses_min_max_rollups():
    """A MIN/MAX metric divides by the row's extreme, not its sum.

    Mirrors `additiveReducerFor` in the plugin's `plugin/utilities.ts`: the
    chart rolls a MAX metric up with max, so a row of [6, 10] reads 60%/100%.
    """
    df = pd.DataFrame(
        {
            "nation": ["US", "US", "UK", "UK"],
            "gender": ["boy", "girl", "boy", "girl"],
            "MAX(num)": [1, 3, 6, 10],
        }
    )
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": [
            {
                "expressionType": "SIMPLE",
                "aggregate": "MAX",
                "column": {"column_name": "num"},
                "label": "MAX(num)",
            }
        ],
        "showValuesAs": "percent_row",
        "rowTotals": True,
    }

    pivoted = pivot_table_v2(df, form_data, apply_number_format=False)

    assert pivoted.loc[("UK",), ("MAX(num)", "boy")] == 0.6
    assert pivoted.loc[("UK",), ("MAX(num)", "girl")] == 1
    # the total divides by itself whatever the reducer
    assert pivoted.loc[("UK",), (total_label(), "")] == 1


def grouping_sets_df() -> pd.DataFrame:
    """A GROUPING SETS result whose rollups differ from any leaf reduction.

    Leaves are 10 and 20, so a leaf-derived total would be 30 (sum) or 15
    (mean). The database rollups are deliberately none of those: 18 down the
    column, 11/21 across the rows, 19 overall -- as a weighted average or a
    distinct count would be.
    """
    return pd.DataFrame(
        {
            "nation": ["US", "UK", None, "US", "UK", None],
            "gender": ["boy", "boy", "boy", None, None, None],
            "nation__superset_grouping": [0, 0, 1, 0, 0, 1],
            "gender__superset_grouping": [0, 0, 0, 1, 1, 1],
            "AVG(num)": [10, 20, 18, 11, 21, 19],
        }
    )


def nested_grouping_sets_df() -> pd.DataFrame:
    """Two row dimensions, so prefix subtotal rows appear beside the rollups.

    Every level holds a value no leaf reduction would produce: the `{region}`
    rollups are 17 and 24, `{region, nation}` are 18 and 25, `{gender}` are 26
    and 24, and the grand total is 21.
    """
    records = [
        # (region, nation, gender, rolled-up markers, value)
        ("EU", "UK", "boy", (0, 0, 0), 10),
        ("EU", "UK", "girl", (0, 0, 0), 30),
        ("NA", "US", "boy", (0, 0, 0), 40),
        ("NA", "US", "girl", (0, 0, 0), 20),
        ("EU", "UK", None, (0, 0, 1), 18),
        ("NA", "US", None, (0, 0, 1), 25),
        ("EU", None, "boy", (0, 1, 0), 10),
        ("EU", None, "girl", (0, 1, 0), 30),
        ("NA", None, "boy", (0, 1, 0), 40),
        ("NA", None, "girl", (0, 1, 0), 20),
        ("EU", None, None, (0, 1, 1), 17),
        ("NA", None, None, (0, 1, 1), 24),
        (None, None, "boy", (1, 1, 0), 26),
        (None, None, "girl", (1, 1, 0), 24),
        (None, None, None, (1, 1, 1), 21),
    ]
    return pd.DataFrame(
        [
            {
                "region": region,
                "nation": nation,
                "gender": gender,
                "region__superset_grouping": markers[0],
                "nation__superset_grouping": markers[1],
                "gender__superset_grouping": markers[2],
                "AVG(num)": value,
            }
            for region, nation, gender, markers, value in records
        ]
    )


NESTED_FORM_DATA = {
    "groupbyRows": ["region", "nation"],
    "groupbyColumns": ["gender"],
    "metrics": ["AVG(num)"],
    "showValuesAs": "percent_row",
    "rowTotals": True,
    "colTotals": True,
}


def test_row_subtotal_keeps_its_total_column_cell():
    """A subtotal row is not a whole-axis total; it must not be looked up as one."""
    pivoted = pivot_table_v2(
        nested_grouping_sets_df(), NESTED_FORM_DATA, apply_number_format=False
    )

    assert not pd.isna(pivoted.loc[("EU", "Subtotal"), (total_label(), "")])


def test_row_subtotal_divides_by_its_own_rollup():
    """An "EU" subtotal divides by the {region} rollup, not the grand total."""
    pivoted = pivot_table_v2(
        nested_grouping_sets_df(), NESTED_FORM_DATA, apply_number_format=False
    )

    assert pivoted.loc[("EU", "Subtotal"), ("AVG(num)", "boy")] == pytest.approx(
        10 / 17
    )
    assert pivoted.loc[("NA", "Subtotal"), ("AVG(num)", "boy")] == pytest.approx(
        40 / 24
    )


def test_nested_leaf_rows_still_divide_by_database_rollups():
    """Regression guard: the leaf and whole-axis behaviour is unchanged."""
    pivoted = pivot_table_v2(
        nested_grouping_sets_df(), NESTED_FORM_DATA, apply_number_format=False
    )

    assert pivoted.loc[("EU", "UK"), ("AVG(num)", "boy")] == pytest.approx(10 / 18)
    assert pivoted.loc[("NA", "US"), ("AVG(num)", "boy")] == pytest.approx(40 / 25)
    assert pivoted.loc[("EU", "UK"), (total_label(), "")] == 1
    assert pivoted.loc[(total_label(), ""), (total_label(), "")] == 1


def test_pivot_table_v2_pivots_only_grouping_sets_leaf_rows():
    """Rollup levels must not be pivoted as if they were leaf rows."""
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "showValuesAs": "percent_col",
    }

    pivoted = pivot_table_v2(grouping_sets_df(), form_data, apply_number_format=False)

    # the rollup rows would otherwise add phantom rows for their NULL dimensions
    assert len(pivoted.index) == 2


def test_pivot_table_v2_metric_default_does_not_raise_key_error():
    """`aggregateFunction: "Metric"` (the result-aggregation control's
    default, meaning "use the metric's own definition") is not a key in
    pivot_v2_aggfunc_map -- this backend path has no result-aggregation
    support yet. It must fall back to the same "Sum" behavior an absent
    field always got, not KeyError.

    Regression test: this frontend default reaching this function used to
    raise ``KeyError('Metric')`` for any Actual Values pivot with rows or
    columns, breaking every scheduled report/alert/CSV/Excel export of a
    pivot_table_v2 chart.
    """
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "aggregateFunction": "Metric",
    }
    without_field = {k: v for k, v in form_data.items() if k != "aggregateFunction"}

    pivoted = pivot_table_v2(grouping_sets_df(), form_data, apply_number_format=False)
    pivoted_default = pivot_table_v2(
        grouping_sets_df(), without_field, apply_number_format=False
    )

    pd.testing.assert_frame_equal(pivoted, pivoted_default)


@pytest.mark.parametrize(
    "mode,expected",
    [
        # each leaf divides by the rollup the database computed, never by a
        # total re-derived from the leaves
        ("percent_col", 20 / 18),
        ("percent_row", 20 / 21),
        ("percent_total", 20 / 19),
    ],
)
def test_pivot_table_v2_divides_by_database_rollups(mode: str, expected: float):
    """For a non-additive metric the DB rollup is the only correct total."""
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "showValuesAs": mode,
    }

    pivoted = pivot_table_v2(grouping_sets_df(), form_data, apply_number_format=False)

    assert pivoted.loc[("UK",), ("AVG(num)", "boy")] == pytest.approx(expected)


def test_pivot_table_v2_actual_values_uses_database_rollups():
    """Actual Values subtotals come from the DB rollups, not leaf reductions.

    Regression test for #44625: in the default (Actual Values) mode the
    report/export path re-reduced the already-aggregated leaf cells, so an
    Average metric's totals were a mean of means instead of the
    database-computed rollups SIP-216 (#41184) added.
    """
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "aggregateFunction": "Average",
        "rowTotals": True,
        "colTotals": True,
    }

    pivoted = pivot_table_v2(grouping_sets_df(), form_data, apply_number_format=False)
    total = f"{_('Total')} (Average)"

    # the database rollups, not the mean of the leaf cells (20 / 15 / 15)
    assert pivoted.loc[("UK",), (total, "")] == 21
    assert pivoted.loc[("US",), (total, "")] == 11
    assert pivoted.loc[(total,), ("AVG(num)", "boy")] == 18
    assert pivoted.loc[(total,), (total, "")] == 19
    # ...while the leaf cells themselves are untouched
    assert pivoted.loc[("UK",), ("AVG(num)", "boy")] == 20
    assert pivoted.loc[("US",), ("AVG(num)", "boy")] == 10


@pytest.mark.parametrize(
    "aggregate_function,expected",
    [
        ("Sample Variance", 50.0),
        ("Sample Standard Deviation", 50.0**0.5),
    ],
)
def test_pivot_table_v2_sample_statistics_aggregations(
    aggregate_function: str, expected: float
):
    """Sample variance/standard deviation must compute for two values at one
    pivot address (10 and 20 give variance 50 and standard deviation ~7.071),
    not raise on a misspelled ``pd.series`` or a non-callable tuple.
    """
    df = pd.DataFrame(
        {
            "nation": ["US", "US"],
            "gender": ["boy", "boy"],
            "SUM(num)": [10, 20],
        }
    )
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["SUM(num)"],
        "aggregateFunction": aggregate_function,
    }

    pivoted = pivot_table_v2(df, form_data, apply_number_format=False)

    assert pivoted.loc[("US",), ("SUM(num)", "boy")] == pytest.approx(expected)


@pytest.mark.parametrize(
    ("aggregate_function", "row_summaries", "column_summaries"),
    [
        # Cells hold the per-cell sample variance (US: 50 and 200, FR: 200 and
        # 450); each summary is the sample variance of those cell values.
        (
            "Sample Variance",
            {"US": 11250.0, "FR": 31250.0},
            {"boy": 11250.0, "girl": 31250.0},
        ),
        # Cells hold the per-cell sample standard deviation (US: ~7.07 and
        # ~14.14, FR: ~14.14 and ~21.21); each summary is the sample standard
        # deviation of those cell values, which is 5 in every scope.
        (
            "Sample Standard Deviation",
            {"US": 5.0, "FR": 5.0},
            {"boy": 5.0, "girl": 5.0},
        ),
    ],
)
def test_pivot_table_v2_sample_statistics_with_summaries(
    aggregate_function: str,
    row_summaries: dict[str, float],
    column_summaries: dict[str, float],
):
    """Row and column summaries invoke the reducer with ``axis``; the sample
    statistics must handle the DataFrame/axis form, not just a Series. Two
    observations per (nation, gender) cell keep the summaries non-zero.
    """
    df = pd.DataFrame(
        {
            "nation": ["US"] * 4 + ["FR"] * 4,
            "gender": ["boy", "boy", "girl", "girl"] * 2,
            "SUM(num)": [10, 20, 30, 50, 40, 60, 70, 100],
        }
    )
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["SUM(num)"],
        "aggregateFunction": aggregate_function,
        "rowTotals": True,
        "colTotals": True,
    }

    pivoted = pivot_table_v2(df, form_data, apply_number_format=False)

    for nation, expected in row_summaries.items():
        assert pivoted.loc[(nation,)].iloc[-1] == pytest.approx(expected)
    for gender, expected in column_summaries.items():
        assert pivoted[("SUM(num)", gender)].iloc[-1] == pytest.approx(expected)


def _normalize_summary(value: Any) -> Any:
    """Make ``List Unique Values`` joins order-insensitive (they come from sets)."""
    if isinstance(value, str):
        return sorted(value.split(", "))
    return value


@pytest.mark.parametrize(
    ("aggregate_function", "row_totals", "col_totals"),
    [
        ("Count", {"US": 2, "FR": 2}, {"boy": 2, "girl": 2}),
        ("Count Unique Values", {"US": 1, "FR": 1}, {"boy": 1, "girl": 1}),
        (
            "List Unique Values",
            {"US": "20, 10", "FR": "50, 30"},
            {"boy": "10, 30", "girl": "50, 20"},
        ),
        ("First", {"US": 10, "FR": 30}, {"boy": 30, "girl": 50}),
        ("Last", {"US": 20, "FR": 50}, {"boy": 10, "girl": 20}),
    ],
)
def test_pivot_table_v2_series_only_reducers_with_summaries(
    aggregate_function: str,
    row_totals: dict[str, Any],
    col_totals: dict[str, Any],
) -> None:
    """Series-only reducers must support each row/column summary axis rather
    than raising ``TypeError`` on the DataFrame/``axis`` invocation, and each
    must select the right element per summary.
    """
    df = pd.DataFrame(
        {
            "nation": ["US", "US", "FR", "FR"],
            "gender": ["boy", "girl", "boy", "girl"],
            "SUM(num)": [10, 20, 30, 50],
        }
    )
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["SUM(num)"],
        "aggregateFunction": aggregate_function,
    }
    total_label = f"Total ({aggregate_function})"

    row_pivot = pivot_table_v2(
        df, {**form_data, "rowTotals": True}, apply_number_format=False
    )
    for nation, expected in row_totals.items():
        assert _normalize_summary(row_pivot.loc[(nation,), (total_label, "")]) == (
            _normalize_summary(expected)
        )

    col_pivot = pivot_table_v2(
        df, {**form_data, "colTotals": True}, apply_number_format=False
    )
    for gender, expected in col_totals.items():
        actual = col_pivot.loc[(total_label,), ("SUM(num)", gender)]
        assert _normalize_summary(actual) == _normalize_summary(expected)


@pytest.mark.parametrize(
    "aggregate_function",
    ["Sample Variance", "Sample Standard Deviation"],
)
def test_pivot_table_v2_sample_statistics_sparse_summaries(
    aggregate_function: str,
) -> None:
    """With one observation per row but two physical columns, the row summary
    is the single-observation 0, not NaN.
    """
    df = pd.DataFrame(
        {
            "nation": ["US", "FR"],
            "gender": ["boy", "girl"],
            "SUM(num)": [10, 20],
        }
    )
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["SUM(num)"],
        "aggregateFunction": aggregate_function,
        "rowTotals": True,
        "colTotals": True,
    }

    pivoted = pivot_table_v2(df, form_data, apply_number_format=False)

    summary_columns = [
        column for column in pivoted.columns if column[1] in ("Subtotal", "")
    ]
    assert summary_columns
    assert (pivoted[summary_columns] == 0).all(axis=None)


def test_sample_dispersion_zero_observations_is_zero_not_nan() -> None:
    """A DataFrame scope with zero non-null observations must render 0, the
    same as the single-observation case just above, not NaN -- the mask only
    covered `count(axis=axis) == 1`, silently leaving an all-null scope (e.g.
    a metric entirely absent for one row of a sparse multi-metric pivot)
    unmasked and NaN instead.
    """
    from superset.charts.client_processing import (  # pylint: disable=import-outside-toplevel
        _sample_dispersion,
    )

    df = pd.DataFrame({"a": [float("nan"), float("nan")], "b": [1.0, 2.0]})

    assert _sample_dispersion(df, "var", axis=0)["a"] == 0
    assert _sample_dispersion(df, "var", axis=0)["b"] == pytest.approx(0.5)


def test_pivot_table_v2_ignores_stale_show_values_as_under_result_aggregation():
    """A result aggregation hides `showValuesAs` in Explore and always wins
    over it on the chart (`resultFactory ?? fractionType` in utilities.ts);
    a persisted `showValuesAs` left over from before `aggregateFunction` was
    set must not silently turn an export/report into a percent table the
    live chart no longer renders as one.
    """
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "aggregateFunction": "Average",
        "showValuesAs": "percent_row",
        "rowTotals": True,
        "colTotals": True,
    }
    without_show_values_as = {k: v for k, v in form_data.items() if k != "showValuesAs"}

    pivoted = pivot_table_v2(grouping_sets_df(), form_data, apply_number_format=False)
    pivoted_without = pivot_table_v2(
        grouping_sets_df(), without_show_values_as, apply_number_format=False
    )

    pd.testing.assert_frame_equal(pivoted, pivoted_without)


def test_pivot_table_v2_actual_values_falls_back_without_rollups():
    """Without GROUPING SETS levels the old leaf-derived totals remain."""
    df = pd.DataFrame(
        {
            "nation": ["US", "UK"],
            "gender": ["boy", "boy"],
            "AVG(num)": [10, 20],
        }
    )
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "aggregateFunction": "Average",
        "rowTotals": True,
        "colTotals": True,
    }

    pivoted = pivot_table_v2(df, form_data, apply_number_format=False)
    total = f"{_('Total')} (Average)"

    # no database rollup to consult: the pandas reduction is kept
    assert pivoted.loc[(total,), ("AVG(num)", "boy")] == 15


def _axis_has_total(axis) -> bool:
    return any(
        total_label() in (label if isinstance(label, tuple) else (label,))
        for label in axis
    )


@pytest.mark.parametrize("layout", ["COLUMNS", "ROWS"])
@pytest.mark.parametrize("toggle", ["rowTotals", "colTotals"])
def test_pivot_table_v2_totals_land_on_the_displayed_axis(layout: str, toggle: str):
    """`rowTotals` is always the right-hand column, whatever the metrics layout.

    Moving metrics to rows transposes the frame on the way out, which would
    otherwise flip the axis each total was inserted on.
    """
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "showValuesAs": "percent_row",
        "metricsLayout": layout,
        toggle: True,
    }

    pivoted = pivot_table_v2(grouping_sets_df(), form_data, apply_number_format=False)

    assert _axis_has_total(pivoted.columns) is (toggle == "rowTotals")
    assert _axis_has_total(pivoted.index) is (toggle == "colTotals")


def test_pivot_table_v2_uses_database_rollups_with_metrics_on_rows():
    """The metrics layout moves the displayed axes but not the rollup levels."""
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "showValuesAs": "percent_col",
        "metricsLayout": "ROWS",
    }

    pivoted = pivot_table_v2(grouping_sets_df(), form_data, apply_number_format=False)

    # 20/18 and 10/18 from the database rollup, not 20/30 from the leaf sum
    assert pivoted.loc[("AVG(num)", "UK"), ("boy",)] == pytest.approx(20 / 18)
    assert pivoted.loc[("AVG(num)", "US"), ("boy",)] == pytest.approx(10 / 18)


def test_pivot_table_v2_keeps_a_null_rollup_blank():
    """A NULL rollup is not a missing one: it blanks the cell, as the chart does."""
    df = grouping_sets_df()
    df.loc[df["nation"].isna() & df["gender"].notna(), "AVG(num)"] = None
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "showValuesAs": "percent_col",
    }

    pivoted = pivot_table_v2(df, form_data, apply_number_format=False)

    # falling back to the leaf-derived total would have produced 10/30 here
    assert pd.isna(pivoted.loc[("US",), ("AVG(num)", "boy")])


def test_pivot_table_v2_rollup_totals_divide_by_themselves():
    """A whole-axis total takes its value from the same level as its denominator."""
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["AVG(num)"],
        "showValuesAs": "percent_row",
        "rowTotals": True,
    }

    pivoted = pivot_table_v2(grouping_sets_df(), form_data, apply_number_format=False)

    # 21/21, not the leaf sum over the database row rollup
    assert pivoted.loc[("UK",), (total_label(), "")] == 1


def test_pivot_df_show_values_as_metrics_on_rows():
    pivoted = pivot_df(
        show_values_as_df(),
        **{**SHOW_VALUES_AS_OPTIONS, "apply_metrics_on_rows": True},
        show_values_as="percent_col",
    )
    total = total_label()

    assert pivoted.loc[("SUM(num)", "US"), ("boy",)] == pytest.approx(1 / 3)
    assert pivoted.loc[("SUM(num)", "US"), ("girl",)] == 0.6
    assert pivoted.loc[("SUM(num)", "US"), (total,)] == 0.5
    assert pivoted.loc[(total, ""), ("boy",)] == 1


def test_pivot_df_show_values_as_zero_denominator_is_blank():
    """A zero total renders blank rather than infinity."""
    df = show_values_as_df()
    df.loc[df["nation"] == "UK", "SUM(num)"] = 0
    pivoted = pivot_df(df, **SHOW_VALUES_AS_OPTIONS, show_values_as="percent_row")

    assert pd.isna(pivoted.loc[("UK",), ("SUM(num)", "boy")])
    assert pivoted.loc[("US",), ("SUM(num)", "boy")] == 0.25


def test_pivot_df_show_values_as_skips_empty_cells_in_denominators():
    """A sparse pivot leaves NaN cells; they must not blank the whole table."""
    df = show_values_as_df()
    # drop UK/girl entirely, so that cell has no rows and pivots to NaN
    sparse = df[~((df["nation"] == "UK") & (df["gender"] == "girl"))]
    options = {**SHOW_VALUES_AS_OPTIONS, "show_rows_total": False}

    pivoted = pivot_df(sparse, **options, show_values_as="percent_total")

    # grand total is 10 + 30 + 20 = 60, the NaN cell contributing nothing
    assert pivoted.loc[("US",), ("SUM(num)", "boy")] == pytest.approx(10 / 60)
    assert pivoted.loc[("UK",), ("SUM(num)", "boy")] == pytest.approx(20 / 60)
    assert pd.isna(pivoted.loc[("UK",), ("SUM(num)", "girl")])


@pytest.mark.parametrize("aggfunc", ["Sum", "Average", "Maximum"])
def test_pivot_df_show_values_as_ignores_legacy_aggregate_function(aggfunc: str):
    """A total always divides by itself, whatever `aggregateFunction` says.

    The chart ignores `aggregateFunction` post-SIP-216, so a percent export must
    too: computing totals with it while the denominators sum leaves the Total
    row/column reading something other than 100%.
    """
    pivoted = pivot_df(
        show_values_as_df(),
        **{**SHOW_VALUES_AS_OPTIONS, "aggfunc": aggfunc},
        show_values_as="percent_row",
    )
    total = total_label()

    assert pivoted.loc[("US",), (total, "")] == 1
    assert pivoted.loc[(total,), (total, "")] == 1
    assert pivoted.loc[("US",), ("SUM(num)", "boy")] == 0.25


def test_pivot_df_show_values_as_supersedes_legacy_fraction_aggfunc():
    """`showValuesAs` wins over a pre-SIP-216 fraction aggregate function."""
    pivoted = pivot_df(
        show_values_as_df(),
        **{**SHOW_VALUES_AS_OPTIONS, "aggfunc": "Sum as Fraction of Total"},
        show_values_as="percent_row",
    )

    assert pivoted.loc[("US",), ("SUM(num)", "boy")] == 0.25


def test_pivot_table_v2_show_values_as_formats_as_percent() -> None:
    """A ratio ignores the configured value format and currency, as the chart does."""
    result = pivot_table_v2(
        show_values_as_df(),
        {
            "groupbyRows": ["nation"],
            "groupbyColumns": ["gender"],
            "metrics": ["SUM(num)"],
            "showValuesAs": "percent_total",
            "valueFormat": "$,.2f",
            "currencyFormat": {"symbol": "USD", "symbolPosition": "prefix"},
        },
    )

    assert result.loc[("US",), ("SUM(num)", "boy")] == "12.5%"
    assert result.loc[("US",), ("SUM(num)", "girl")] == "37.5%"


def test_pivot_table_v2_unrecognized_aggregate_function_keeps_show_values_as() -> None:
    """An unrecognized `aggregateFunction` isn't a result aggregation to the
    frontend's own `getResultAggregation` lookup -- `RESULT_AGGREGATIONS`
    doesn't contain it, so it falls back to the metric definition and still
    renders the stored percent mode. A stale or otherwise unrecognized value
    (e.g. a legacy spelling this map doesn't know, or an empty string) must
    not suppress `showValuesAs` here either, or the export disagrees with
    what the chart actually shows.
    """
    form_data = {
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["SUM(num)"],
        "showValuesAs": "percent_total",
        "aggregateFunction": "Integer Sum",
        "valueFormat": "$,.2f",
        "currencyFormat": {"symbol": "USD", "symbolPosition": "prefix"},
    }

    result = pivot_table_v2(show_values_as_df(), form_data)

    assert result.loc[("US",), ("SUM(num)", "boy")] == "12.5%"
    assert result.loc[("US",), ("SUM(num)", "girl")] == "37.5%"


def test_pivot_table_v2_show_values_as_actual_keeps_raw_values() -> None:
    result = pivot_table_v2(
        show_values_as_df(),
        {
            "groupbyRows": ["nation"],
            "groupbyColumns": ["gender"],
            "metrics": ["SUM(num)"],
            "showValuesAs": "actual",
            "valueFormat": ",d",
        },
    )

    assert result.loc[("US",), ("SUM(num)", "boy")] == "10"


def test_apply_client_processing_xlsx_formats_percentages_as_cells():
    """The workbook shows percentages without giving up numeric cells.

    A scheduled XLSX report and the browser's pivoted-Excel download have to
    agree; writing formatted text would match the download but stop Excel from
    summing the column.
    """
    import openpyxl  # noqa: PLC0415

    source = pd.DataFrame(
        {"nation": ["US", "US"], "gender": ["boy", "girl"], "SUM(num)": [1, 99]}
    )
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.XLSX,
                "data": excel.df_to_excel(source, index=False),
            }
        ]
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["SUM(num)"],
        "showValuesAs": "percent_row",
        "metricsLayout": "COLUMNS",
    }

    processed = apply_client_processing(result, form_data)

    workbook = openpyxl.load_workbook(BytesIO(processed["queries"][0]["data"]))
    sheet = workbook.active
    assert sheet.cell(row=2, column=2).value == pytest.approx(0.01)
    assert sheet.cell(row=2, column=2).number_format == "0.0%"


def test_apply_client_processing_xlsx_no_percent_format_under_result_aggregation():
    """A stale percent `showValuesAs` under a result aggregation is ignored
    when pivoting (the chart no longer renders it, see
    `effective_show_values_as`), so the workbook must not apply the percent
    number format either: an actual value of 10 would otherwise display as
    1000.0%.
    """
    import openpyxl  # noqa: PLC0415

    source = pd.DataFrame(
        {"nation": ["US", "US"], "gender": ["boy", "girl"], "SUM(num)": [1, 99]}
    )
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.XLSX,
                "data": excel.df_to_excel(source, index=False),
            }
        ]
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["SUM(num)"],
        "aggregateFunction": "Average",
        "showValuesAs": "percent_row",
        "metricsLayout": "COLUMNS",
    }

    processed = apply_client_processing(result, form_data)

    workbook = openpyxl.load_workbook(BytesIO(processed["queries"][0]["data"]))
    sheet = workbook.active
    assert sheet.cell(row=2, column=2).value == 1
    assert sheet.cell(row=2, column=2).number_format != "0.0%"


def test_apply_client_processing_csv_format_show_values_as():
    """CSV exports carry the percentages the chart displays."""
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": (
                    "nation,gender,SUM(num)\n"
                    "US,boy,10\nUS,girl,30\nUK,boy,20\nUK,girl,20\n"
                ),
            }
        ]
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["nation"],
        "groupbyColumns": ["gender"],
        "metrics": ["SUM(num)"],
        "showValuesAs": "percent_row",
        "metricsLayout": "COLUMNS",
    }

    processed = apply_client_processing(result, form_data)

    assert processed["queries"][0]["data"] == (
        ",SUM(num) boy,SUM(num) girl\nUK,0.5,0.5\nUS,0.25,0.75\n"
    )


def test_table():
    """
    Test that the table reports honor `d3NumberFormat`.
    """
    df = pd.DataFrame.from_dict({"count": {0: 80679663}})
    form_data = {
        "adhoc_filters": [
            {
                "clause": "WHERE",
                "comparator": "NULL",
                "expressionType": "SIMPLE",
                "filterOptionName": "filter_ameaka2efjv_rfv1et5nwng",
                "isExtra": False,
                "isNew": False,
                "operator": "!=",
                "sqlExpression": None,
                "subject": "lang_at_home",
            }
        ],
        "all_columns": [],
        "color_pn": True,
        "column_config": {"count": {"d3NumberFormat": ",d"}},
        "conditional_formatting": [],
        "datasource": "8__table",
        "extra_form_data": {},
        "granularity_sqla": "time_start",
        "groupby": ["lang_at_home"],
        "metrics": ["count"],
        "order_by_cols": [],
        "order_desc": True,
        "percent_metrics": [],
        "query_mode": "aggregate",
        "row_limit": "15",
        "server_page_length": 10,
        "show_cell_bars": True,
        "table_timestamp_format": "smart_date",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "url_params": {},
        "viz_type": "table",
    }
    formatted = table(df, form_data)
    assert (
        formatted.to_markdown()
        == """
|    |      count |
|---:|-----------:|
|  0 | 80,679,663 |
    """.strip()
    )


def test_table_applies_currency_format() -> None:
    """
    Table reports honor a column's `currencyFormat`.
    """
    df = pd.DataFrame.from_dict({"amount": {0: 1234.5}})
    form_data = {
        "viz_type": "table",
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "currencyFormat": {"symbol": "USD", "symbolPosition": "prefix"},
            }
        },
    }
    formatted = table(df, form_data)
    assert formatted["amount"].tolist() == ["$ 1,234.50"]


def test_table_applies_time_format_alongside_number_format() -> None:
    """
    Table reports honor an explicit ``d3TimeFormat`` on a temporal column,
    mirroring the Table plugin's first formatter branch, while still applying
    number formats on the same call. Without the time branch the timestamp was
    left raw even though the browser formats it.
    """
    df = pd.DataFrame.from_dict(
        {
            "ds": {
                0: pd.Timestamp("2024-01-15 13:45:00"),
                1: pd.Timestamp("2024-06-30 08:00:00"),
            },
            "SUM(num)": {0: 1234567.0, 1: 42.0},
        }
    )
    form_data = {
        "viz_type": "table",
        "metrics": ["SUM(num)"],
        "column_config": {
            "ds": {"d3TimeFormat": "%Y"},
            "SUM(num)": {"d3NumberFormat": ",.2f"},
        },
    }
    formatted = table(df, form_data)
    assert formatted["ds"].tolist() == ["2024", "2024"]
    assert formatted["SUM(num)"].tolist() == ["1,234,567.00", "42.00"]


@pytest.mark.parametrize(
    ("time_format", "expected"),
    [
        ("%d", "01"),
        ("%G", "2020"),
        ("%H", "13"),
        ("%I", "01"),
        ("%j", "001"),
        ("%m", "01"),
        ("%M", "05"),
        ("%S", "09"),
        ("%u", "5"),
        ("%U", "00"),
        ("%V", "53"),
        ("%w", "5"),
        ("%W", "00"),
        ("%y", "21"),
        ("%Y", "2021"),
        ("%%", "%"),
        ("%G-W%V-%u", "2020-W53-5"),
    ],
)
def test_table_time_format_matches_d3_for_each_shared_directive(
    time_format: str, expected: str
) -> None:
    """
    Every directive in ``SHARED_TIME_DIRECTIVES`` is held to d3's own output.

    The expected strings are what ``utcFormat`` from ``d3-time-format`` returns
    for this timestamp, so taking a directive out of the allowlist fails the
    case that names it rather than passing unnoticed. 1 January 2021 is a
    Friday in ISO week 53 of 2020, which is what separates ``%G`` from ``%Y``
    and ``%V`` from ``%U`` and ``%W``; 13:05:09 separates ``%I`` from ``%H``.
    """
    df = pd.DataFrame.from_dict({"ds": {0: pd.Timestamp("2021-01-01 13:05:09")}})
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": time_format}},
    }

    formatted = table(df, form_data)

    assert formatted["ds"].tolist() == [expected]


@pytest.mark.parametrize(
    "time_format",
    [
        "%Y%",
        "%d/%m/%Y%",
        "%%%",
        "%Y%%%",
    ],
)
def test_table_skips_a_time_format_with_an_unconsumed_percent(
    time_format: str,
) -> None:
    """
    A ``%`` that no directive consumes is dropped by d3 and copied by
    ``strftime``.

    ``%Y%`` is ``2021`` in the chart and would be ``2021%`` in the report, and
    ``%%%`` is ``%`` against ``%%``. Every directive in these is shared, so the
    allowlist alone accepts them and the column has to be left to the browser
    on the stray ``%`` instead.
    """
    df = pd.DataFrame.from_dict({"ds": {0: pd.Timestamp("2021-01-01 13:05:09")}})
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": time_format}},
    }

    formatted = table(df, form_data)

    assert formatted["ds"].tolist() == [pd.Timestamp("2021-01-01 13:05:09")]


def test_table_leaves_adaptive_time_format_alone() -> None:
    """
    ``smart_date`` ("Adaptive formatting") is an alias the browser resolves to a
    formatter, not a format string. It carries no ``%`` directive, so the rule
    that a format has to have one is what keeps it away from ``strftime``,
    which would hand the alias back unchanged and print it in every cell.
    """
    df = pd.DataFrame.from_dict(
        {
            "ds": {
                0: pd.Timestamp("2024-01-15 13:45:00"),
                1: pd.Timestamp("2024-06-30 08:00:00"),
            },
        }
    )
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": "smart_date"}},
    }
    formatted = table(df, form_data)
    assert formatted["ds"].tolist() == [
        pd.Timestamp("2024-01-15 13:45:00"),
        pd.Timestamp("2024-06-30 08:00:00"),
    ]


@pytest.mark.parametrize(
    "time_format",
    [
        "yyyy-mm-dd",
        "Q1",
        "Date",
        "100%",
    ],
)
def test_table_skips_a_time_format_with_no_directive(time_format: str) -> None:
    """
    A format carrying no ``%`` directive is not a format string.

    ``strftime`` has nothing to substitute and returns the string it was given,
    so accepting one of these would put ``yyyy-mm-dd`` in every cell of the
    column instead of a date. This is the rule that also rejects the adaptive
    alias.
    """
    df = pd.DataFrame.from_dict({"ds": {0: pd.Timestamp("2024-04-02 13:45:00")}})
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": time_format}},
    }

    formatted = table(df, form_data)

    assert formatted["ds"].tolist() == [pd.Timestamp("2024-04-02 13:45:00")]


@pytest.mark.parametrize(
    ("values", "expected"),
    [
        pytest.param(
            {0: "2024-01-15", 1: "2024-06-30"},
            ["2024-01-15", "2024-06-30"],
            id="string column",
        ),
        pytest.param({0: 2024, 1: 2025}, [2024, 2025], id="integer column"),
    ],
)
def test_table_skips_a_time_format_on_a_non_datetime_column(
    values: dict[int, Any], expected: list[Any]
) -> None:
    """
    ``d3TimeFormat`` on a column that is not datetime leaves the column alone.

    A string or a number has no ``strftime`` to call, so formatting one would
    raise ``AttributeError`` in the report worker. The browser reaches these
    through a d3-time-format port, which parses the value first, and that is
    what the follow-up needs rather than something ``strftime`` can stand in
    for.
    """
    df = pd.DataFrame.from_dict({"ds": values})
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": "%Y"}},
    }

    formatted = table(df, form_data)

    assert formatted["ds"].tolist() == expected


def test_table_time_format_keeps_nulls() -> None:
    """
    A missing timestamp stays missing, and does not fail the report.

    ``strftime`` on a ``NaT`` does not produce the string ``"NaT"``; it raises
    ``ValueError: NaTType does not support strftime``, so without the
    ``pd.notna`` guard one null would fail the whole report rather than print
    anything.
    """
    df = pd.DataFrame.from_dict(
        {"ds": {0: pd.Timestamp("2024-01-15"), 1: pd.NaT}},
    )
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": "%Y"}},
    }

    formatted = table(df, form_data)

    assert formatted["ds"].tolist()[0] == "2024"
    assert pd.isna(formatted["ds"].tolist()[1])


@pytest.mark.parametrize(
    "time_format",
    ["  %Y-%m-%d  ", " %Y-%m-%d", "%Y-%m-%d ", "\t%Y-%m-%d\n", "\r%Y-%m-%d\x0b"],
)
def test_table_time_format_strips_surrounding_whitespace(time_format: str) -> None:
    """
    A padded format renders like the chart, which trims before formatting.

    ``TimeFormatterRegistry.get`` trims the format string before building the
    formatter, while d3 itself does not: ``utcFormat(" %Y ")`` returns
    ``" 2024 "``. Without the strip the report kept the padding the chart had
    already dropped. Every character here is whitespace to both ``str.strip``
    and ``String.prototype.trim``; the handful they disagree on is in
    ``format_time_column``.
    """
    df = pd.DataFrame.from_dict({"ds": {0: pd.Timestamp("2024-01-15")}})
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": time_format}},
    }

    formatted = table(df, form_data)

    assert formatted["ds"].tolist() == ["2024-01-15"]


@pytest.mark.parametrize(
    "time_format",
    [
        "%Y Q%q",
        "%H:%M:%S.%L",
        "%Q",
        "%-d %b",
        "%x",
        "local!%Y-%m-%d",
        "%Y local!",
        "%H:%M:%S.%f",
    ],
)
def test_table_skips_a_time_format_strftime_cannot_express(time_format: str) -> None:
    """
    ``strftime`` copies a directive it does not recognize into its output, so a
    format carrying a d3-only one would reach the report as literal text.

    So would the ``local!`` marker, wherever it sits. ``TimeFormatterRegistry``
    reads local time off ``startsWith`` but takes the marker out of the format
    with ``replace``, which drops it from anywhere, so a trailing one is not
    part of the chart's output either and a format carrying one is left to the
    browser rather than matched on its prefix.

    ``%f`` is recognized by both but means different things: d3 scales the
    milliseconds it has, so its microseconds always end in ``000``. The column
    is left raw instead, the way an unsupported number format is.
    """
    df = pd.DataFrame.from_dict({"ds": {0: pd.Timestamp("2024-04-02 13:45:00")}})
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": time_format}},
    }

    formatted = table(df, form_data)

    assert formatted["ds"].tolist() == [pd.Timestamp("2024-04-02 13:45:00")]


@pytest.mark.parametrize(
    "time_format",
    [
        "%B %Y",
        "%b %d",
        "%A",
        "%a",
        "%I:%M %p",
    ],
)
def test_table_skips_a_time_format_naming_a_month_weekday_or_period(
    time_format: str,
) -> None:
    """
    The browser reads weekday, month and period names from the d3 locale
    definition, which ``D3_TIME_FORMAT`` overrides key by key, while ``strftime``
    reads them from the worker's ``LC_TIME``. With French ``months`` configured
    the chart renders ``janvier 2024`` where ``strftime`` on an English worker
    renders ``January 2024``, so the column is left raw rather than formatted
    with names the chart does not use.
    """
    df = pd.DataFrame.from_dict({"ds": {0: pd.Timestamp("2024-01-15 14:30:00")}})
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": time_format}},
    }

    formatted = table(df, form_data)

    assert formatted["ds"].tolist() == [pd.Timestamp("2024-01-15 14:30:00")]


def test_table_time_format_reads_an_aware_value_in_utc() -> None:
    """
    ``createD3TimeFormatter`` formats with ``utcFormat`` unless the format asks
    for local time, so an aware value has to be read in UTC rather than in its
    own offset, which is a different instant.
    """
    df = pd.DataFrame(
        {"ds": pd.to_datetime(["2024-01-01 00:30"]).tz_localize("Asia/Kolkata")}
    )
    form_data = {
        "viz_type": "table",
        "column_config": {"ds": {"d3TimeFormat": "%Y-%m-%d %H:%M"}},
    }

    formatted = table(df, form_data)

    assert formatted["ds"].tolist() == ["2023-12-31 19:00"]


def test_table_report_time_format_survives_the_notification_table() -> None:
    """
    ``chart_data_to_dataframe`` casts every column the payload types as
    ``TEMPORAL`` back to a datetime, so the payload has to report a formatted
    temporal column as text or the notification table re-parses the strings and
    the configured format never reaches the recipient.
    """
    datasource = MagicMock()
    datasource.data = {"verbose_map": {}, "column_formats": {}}
    datasource.columns = [{"column_name": "ds", "is_dttm": True}]
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.JSON,
                "data": [
                    {"ds": pd.Timestamp("2024-01-02"), "cnt": 1},
                    {"ds": pd.Timestamp("2024-01-15"), "cnt": 2},
                ],
            }
        ],
    }
    form_data = {
        "viz_type": "table",
        "metrics": ["cnt"],
        "column_config": {"ds": {"d3TimeFormat": "%d/%m/%Y"}},
    }

    processed = apply_client_processing(result, form_data, datasource)
    query = processed["queries"][0]

    assert list(query["data"]["ds"].values()) == ["02/01/2024", "15/01/2024"]
    assert query["coltypes"][0] == GenericDataType.STRING

    payload = json.dumps(
        {
            "result": [
                {
                    "data": query["data"],
                    "coltypes": [int(coltype) for coltype in query["coltypes"]],
                    "colnames": query["colnames"],
                    "indexnames": query["indexnames"],
                }
            ]
        }
    ).encode("utf-8")

    notification_table = chart_data_to_dataframe(payload)

    assert notification_table is not None
    assert notification_table[("ds",)].tolist() == ["02/01/2024", "15/01/2024"]


def test_table_report_time_format_survives_a_null_in_the_column() -> None:
    """
    A null alongside the formatted strings does not cost the column its type.

    ``format_time_column`` leaves nulls in place, and ``infer_dtype`` reads a
    ``NaT`` as a value of its own type rather than one ``skipna`` drops, so it
    calls the column ``"mixed"`` rather than text. Typing it ``TEMPORAL`` on
    that basis sends it back through the datetime cast in
    ``chart_data_to_dataframe``, which reads ``02/01/2024`` month-first and
    puts 1 February in the report, so any null in the column would undo the
    format.
    """
    datasource = MagicMock()
    datasource.data = {"verbose_map": {}, "column_formats": {}}
    datasource.columns = [{"column_name": "ds", "is_dttm": True}]
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.JSON,
                "data": [
                    {"ds": pd.Timestamp("2024-01-02"), "cnt": 1},
                    {"ds": None, "cnt": 2},
                ],
            }
        ],
    }
    form_data = {
        "viz_type": "table",
        "metrics": ["cnt"],
        "column_config": {"ds": {"d3TimeFormat": "%d/%m/%Y"}},
    }

    query = apply_client_processing(result, form_data, datasource)["queries"][0]

    assert query["coltypes"][0] == GenericDataType.STRING
    values = list(query["data"]["ds"].values())
    assert values[0] == "02/01/2024"
    assert pd.isna(values[1])


def test_payload_coltypes_reads_repeated_labels_positionally() -> None:
    """
    Two columns sharing a label are typed one at a time, not together.

    ``extract_dataframe_dtypes`` walks the frame by position, so the types it
    returns are read back the same way. Addressing the frame by label instead
    would hand a repeated label every column that carries it, type them all
    from the first one's values, and cost a real timestamp column its
    ``TEMPORAL`` in the payload.
    """
    datasource = MagicMock()
    datasource.data = {"verbose_map": {}, "column_formats": {}}
    datasource.columns = [{"column_name": "ds", "is_dttm": True}]
    df = pd.DataFrame(
        [["02/01/2024", pd.Timestamp("2024-02-01")]], columns=["ds", "ds"]
    )

    assert payload_coltypes(df, datasource) == [
        GenericDataType.STRING,
        GenericDataType.TEMPORAL,
    ]


def test_report_payload_keeps_an_unformatted_temporal_column_temporal() -> None:
    """
    A temporal column nobody rendered as text keeps its ``TEMPORAL`` type.

    The payload type is what ``chart_data_to_dataframe`` reads to rebuild dates
    for the notification table, and what the XLSX branch hands to
    ``excel.apply_column_types`` to get date cells. Reporting a datetime column
    as text would cost it both.
    """
    datasource = MagicMock()
    datasource.data = {"verbose_map": {}, "column_formats": {}}
    datasource.columns = [{"column_name": "ds", "is_dttm": True}]
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.JSON,
                "data": [
                    {"ds": pd.Timestamp("2024-01-02"), "cnt": 1},
                    {"ds": pd.Timestamp("2024-01-15"), "cnt": 2},
                ],
            }
        ],
    }
    form_data = {"viz_type": "table", "metrics": ["cnt"]}

    query = apply_client_processing(result, form_data, datasource)["queries"][0]

    assert query["coltypes"][0] == GenericDataType.TEMPORAL
    assert list(query["data"]["ds"].values()) == [
        pd.Timestamp("2024-01-02"),
        pd.Timestamp("2024-01-15"),
    ]


def test_report_payload_types_a_csv_read_temporal_column_as_text() -> None:
    """
    The type follows the values, not whether this module wrote them.

    Number and time formatting only runs for the JSON result format, so a CSV
    payload's temporal column arrives as text from ``read_csv`` and is reported
    that way. Keying the override on "did a post-processor format this column"
    instead of on the values would report text as ``TEMPORAL`` here.
    """
    datasource = MagicMock()
    datasource.data = {"verbose_map": {}, "column_formats": {}}
    datasource.columns = [{"column_name": "ds", "is_dttm": True}]
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": "ds,cnt\n2024-01-02,1\n2024-01-15,2\n",
            }
        ],
    }
    form_data = {
        "viz_type": "table",
        "metrics": ["cnt"],
        "column_config": {"ds": {"d3TimeFormat": "%d/%m/%Y"}},
    }

    query = apply_client_processing(result, form_data, datasource)["queries"][0]

    assert query["coltypes"] == [GenericDataType.STRING, GenericDataType.NUMERIC]
    # The CSV text itself is handed back untouched, formats included.
    assert query["data"] == "ds,cnt\n2024-01-02,1\n2024-01-15,2\n"


def test_table_applies_si_number_format() -> None:
    """
    Table reports honor d3 formats that Python's str.format cannot express.
    """
    df = pd.DataFrame.from_dict({"amount": {0: 1234.0}})
    form_data = {
        "viz_type": "table",
        "column_config": {"amount": {"d3NumberFormat": ".3s"}},
    }
    formatted = table(df, form_data)
    assert formatted["amount"].tolist() == ["1.23k"]


def test_table_applies_small_number_format() -> None:
    """
    A cell below one in absolute value takes the column's small-number format.

    ``formatColumnValue`` in the Table plugin picks the formatter per value, so
    the report has to switch per cell rather than per column.

    The switch point itself is covered on both sides and in both signs. The
    browser's test is ``Math.abs(value) < 1``, so exactly one is a large cell
    and a large negative stays large: widening the comparison to ``<= 1`` or
    comparing the signed value would render one of these as a percentage.
    """
    df = pd.DataFrame.from_dict(
        {
            "ratio": {
                0: 0.0523,
                1: 12.5,
                2: -0.5,
                3: 0.0,
                4: 1.0,
                5: -1.0,
                6: -12.5,
            }
        }
    )
    form_data = {
        "viz_type": "table",
        "metrics": ["ratio"],
        "column_config": {
            "ratio": {"d3NumberFormat": ",.2f", "d3SmallNumberFormat": ".1%"}
        },
    }
    formatted = table(df, form_data)
    assert formatted["ratio"].tolist() == [
        "5.2%",
        "12.50",
        "-50.0%",
        "0.0%",
        "1.00",
        "-1.00",
        "-12.50",
    ]


def test_table_ignores_blank_small_number_format() -> None:
    """
    A cleared Select leaves an empty string behind, which
    ``getSmallNumberFormatter`` treats as unset.
    """
    df = pd.DataFrame.from_dict({"ratio": {0: 0.0523, 1: 12.5}})
    form_data = {
        "viz_type": "table",
        "metrics": ["ratio"],
        "column_config": {
            "ratio": {"d3NumberFormat": ",.2f", "d3SmallNumberFormat": "  "}
        },
    }
    formatted = table(df, form_data)
    assert formatted["ratio"].tolist() == ["0.05", "12.50"]


def test_table_small_number_format_without_a_normal_format() -> None:
    """
    A small-number format on its own is enough to format the small cells.

    The other cells follow the browser: a metric falls back to SMART_NUMBER,
    while a numeric dimension gets no formatter at all and is printed as is.
    """
    df = pd.DataFrame.from_dict({"ratio": {0: 0.0523, 1: 1200.0}})
    form_data = {
        "viz_type": "table",
        "metrics": ["ratio"],
        "column_config": {"ratio": {"d3SmallNumberFormat": ".1%"}},
    }
    assert table(df, form_data)["ratio"].tolist() == ["5.2%", "1.2k"]

    # 1200.0 separates the two: SMART_NUMBER renders it ``1.2k`` while an
    # unformatted cell prints ``1200``. A value such as 12.5 reads the same
    # either way and would not show the difference.
    df = pd.DataFrame.from_dict({"ratio": {0: 0.0523, 1: 1200.0}})
    form_data = {
        "viz_type": "table",
        "metrics": [],
        "column_config": {"ratio": {"d3SmallNumberFormat": ".1%"}},
    }
    assert table(df, form_data)["ratio"].tolist() == ["5.2%", "1200"]


def test_table_small_number_format_uses_the_column_config_currency() -> None:
    """
    ``getSmallNumberFormatter`` pairs the small format with the column config's
    own currency, and the currency formatter drops percent signs.
    """
    df = pd.DataFrame.from_dict({"amount": {0: 0.5, 1: 12.5}})
    form_data = {
        "viz_type": "table",
        "metrics": ["amount"],
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "d3SmallNumberFormat": ".1%",
                "currencyFormat": {"symbol": "USD", "symbolPosition": "prefix"},
            }
        },
    }
    formatted = table(df, form_data)
    assert formatted["amount"].tolist() == ["$ 50.0", "$ 12.50"]


def test_table_small_number_format_drops_a_dollar_for_a_cleared_currency() -> None:
    """
    ``getSmallNumberFormatter`` builds a ``CurrencyFormatter`` from any
    ``currencyFormat`` object, and that constructor strips ``$`` from the format
    before it checks the symbol, so a cleared currency still drops the ``$``.
    """
    df = pd.DataFrame.from_dict({"amount": {0: 0.5, 1: 12.5}})
    form_data = {
        "viz_type": "table",
        "metrics": ["amount"],
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "d3SmallNumberFormat": "$,.2f",
                "currencyFormat": {"symbolPosition": "prefix"},
            }
        },
    }

    formatted = table(df, form_data)

    assert formatted["amount"].tolist() == ["0.50", "12.50"]


def test_table_small_number_format_of_only_a_dollar_falls_back() -> None:
    """
    A small-number format that is nothing but ``$`` lands on the default.

    ``CurrencyFormatter.format`` calls ``getNumberFormatter`` with
    ``getNormalizedD3Format()``, which strips every ``$``, and the registry
    answers an empty id with its ``defaultKey``, SMART_NUMBER. Leaving the
    format empty instead would skip the small-number branch and give the cell
    the column's normal format.
    """
    df = pd.DataFrame.from_dict({"amount": {0: 0.5, 1: 12.5}})
    form_data = {
        "viz_type": "table",
        "metrics": ["amount"],
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "d3SmallNumberFormat": "$",
                "currencyFormat": {"symbol": "", "symbolPosition": "prefix"},
            }
        },
    }

    # SMART_NUMBER gives ``0.5``; the column's own ``,.2f`` would give ``0.50``.
    assert table(df, form_data)["amount"].tolist() == ["0.5", "12.50"]


def test_table_small_number_format_keeps_a_dollar_without_a_currency_config() -> None:
    """
    With no ``currencyFormat`` object the browser builds a plain number
    formatter instead, and d3 renders the ``$`` in the format itself.
    """
    df = pd.DataFrame.from_dict({"amount": {0: 0.5, 1: 12.5}})
    form_data = {
        "viz_type": "table",
        "metrics": ["amount"],
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "d3SmallNumberFormat": "$,.2f",
            }
        },
    }

    formatted = table(df, form_data)

    assert formatted["amount"].tolist() == ["$0.50", "12.50"]


def test_table_small_number_format_drops_a_dollar_with_a_currency_context() -> None:
    """
    The per-row branch builds the same formatter, so a cleared currency drops
    the ``$`` there too, and no row's currency reaches a symbol-less config.
    """
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    df = pd.DataFrame(
        {
            "amount": [0.5, 12.5],
            "currency": ["USD", "USD"],
        }
    )
    form_data = {
        "viz_type": "table",
        "metrics": ["amount"],
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "d3SmallNumberFormat": "$,.2f",
                "currencyFormat": {"symbolPosition": "prefix"},
            }
        },
    }

    formatted = table(df, form_data, datasource, detected_currency="USD")

    assert formatted["amount"].tolist() == ["0.50", "12.50"]


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        pytest.param(None, "", id="null"),
        pytest.param("n/a", "n/a", id="text"),
        pytest.param(float("nan"), "", id="nan"),
    ],
)
def test_table_small_number_format_on_a_non_numeric_percent_metric(
    value: Any, expected: str
) -> None:
    """
    A percent metric cell that is not a number is carried through.

    The percent metric branch is the one that does not test
    ``is_numeric_dtype`` first, so these values reach the small-number
    predicate. Handing them to ``float`` raises ``TypeError`` for a null and
    ``ValueError`` for text, which would be a failed report rather than a
    misformatted cell.
    """
    # Built as object dtype on purpose: ``from_dict`` would turn the null into
    # a ``nan`` in a float column, and the value has to arrive as it is for the
    # type check to be the thing under test.
    df = pd.DataFrame({"%share": pd.Series([value, 0.5], dtype=object)})
    form_data = {
        "viz_type": "table",
        "percent_metrics": ["share"],
        "column_config": {"%share": {"d3SmallNumberFormat": ".1%"}},
    }

    assert table(df, form_data)["%share"].tolist() == [expected, "50.0%"]


def test_table_small_number_format_on_a_percent_metric() -> None:
    """
    Percent metric columns are numeric too, so they switch formatter the same
    way; their other cells keep the PERCENT_3_POINT default.
    """
    df = pd.DataFrame.from_dict({"%cnt": {0: 0.004, 1: 2.5}})
    form_data = {
        "viz_type": "table",
        "metrics": [],
        "percent_metrics": ["cnt"],
        "column_config": {"%cnt": {"d3SmallNumberFormat": ".2%"}},
    }
    formatted = table(df, form_data)
    assert formatted["%cnt"].tolist() == ["0.40%", "250.000%"]


def test_table_small_number_format_on_a_percent_metric_keeps_the_row_currency() -> None:
    """
    A percent metric's own formatter carries no currency, but the small-number
    one is built from the column config, so its cells resolve AUTO per row.
    """
    form_data = {
        "viz_type": "table",
        "metrics": [],
        "percent_metrics": ["cnt"],
        "column_config": {
            "%cnt": {
                "d3SmallNumberFormat": ",.3f",
                "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
            }
        },
    }
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    df = pd.DataFrame(
        {
            "%cnt": [0.5, 0.25, 2.5],
            "currency": ["USD", None, "EUR"],
        }
    )

    formatted = table(df, form_data, datasource, detected_currency="GBP")

    assert formatted["%cnt"].tolist() == ["$ 0.500", "0.250", "250.000%"]


def test_table_small_number_format_ignores_a_saved_currency() -> None:
    """
    The small format is paired with the column config's own currency, so a
    currency the dataset saved for the metric reaches only the normal format.
    """
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "metrics": [
            {
                "metric_name": "amount",
                "currency": {"symbol": "USD", "symbolPosition": "prefix"},
            }
        ],
    }
    df = pd.DataFrame.from_dict({"amount": {0: 0.5, 1: 12.5}})
    form_data = {
        "viz_type": "table",
        "metrics": ["amount"],
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "d3SmallNumberFormat": ".1%",
            }
        },
    }

    formatted = table(df, form_data, datasource)

    assert formatted["amount"].tolist() == ["50.0%", "$ 12.50"]


def test_ag_grid_table_ignores_a_small_number_format() -> None:
    """
    The AG Grid table paints its cells through ``valueFormatter``, which never
    builds the small-number formatter, so a report on one keeps the column's
    normal format for every cell.
    """
    df = pd.DataFrame.from_dict({"amount": {0: 0.5, 1: 12.5}})
    column_config = {
        "amount": {"d3NumberFormat": ",.2f", "d3SmallNumberFormat": ".1%"},
    }

    ag_grid = table(
        df.copy(),
        {
            "viz_type": "ag-grid-table",
            "metrics": ["amount"],
            "column_config": column_config,
        },
    )
    plain = table(
        df.copy(),
        {"viz_type": "table", "metrics": ["amount"], "column_config": column_config},
    )

    assert ag_grid["amount"].tolist() == ["0.50", "12.50"]
    assert plain["amount"].tolist() == ["50.0%", "12.50"]


def test_table_small_number_format_leaves_non_numeric_columns_alone() -> None:
    """
    The plugin only switches formatter for a numeric column, so a string column
    configured with a small-number format is untouched.
    """
    df = pd.DataFrame.from_dict({"name": {0: "a", 1: "b"}})
    form_data = {
        "viz_type": "table",
        "metrics": [],
        "column_config": {"name": {"d3SmallNumberFormat": ".1%"}},
    }
    formatted = table(df, form_data)
    assert formatted["name"].tolist() == ["a", "b"]


def test_table_applies_smart_number_default_to_unconfigured_metric() -> None:
    """
    A metric with no saved d3 format still renders like Explore. The Table
    plugin gives every metric column a formatter, and ``getNumberFormatter``
    defaults to SMART_NUMBER, so the report must not leave the value raw.
    """
    df = pd.DataFrame.from_dict({"count": {0: 1234567}})
    form_data = {"viz_type": "table", "metrics": ["count"], "percent_metrics": []}
    formatted = table(df, form_data)
    assert formatted["count"].tolist() == ["1.23M"]


def test_table_leaves_unconfigured_numeric_dimension_untouched() -> None:
    """
    A numeric non-metric column with no configured format is left raw, matching
    the browser, which only formats numeric dimensions when a format or currency
    is set.
    """
    df = pd.DataFrame.from_dict({"year": {0: 2024}, "count": {0: 1234567}})
    form_data = {"viz_type": "table", "metrics": ["count"], "percent_metrics": []}
    formatted = table(df, form_data)
    assert formatted["year"].tolist() == [2024]
    assert formatted["count"].tolist() == ["1.23M"]


def test_table_applies_percent_3_point_default_to_percent_metric() -> None:
    """
    Percent metric columns default to PERCENT_3_POINT in the Table plugin.
    """
    df = pd.DataFrame.from_dict({"%count": {0: 0.1234}})
    form_data = {"viz_type": "table", "metrics": [], "percent_metrics": ["count"]}
    formatted = table(df, form_data)
    assert formatted["%count"].tolist() == ["12.340%"]


def test_table_applies_datasource_saved_metric_format_without_chart_override() -> None:
    df = pd.DataFrame.from_dict({"amount": {0: 1234.5}})
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {"amount": ",.2f"},
        "verbose_map": {},
    }

    formatted = table(df, {"viz_type": "table"}, datasource)

    assert formatted["amount"].tolist() == ["1,234.50"]


def test_table_chart_format_overrides_datasource_saved_metric_format() -> None:
    df = pd.DataFrame.from_dict({"amount": {0: 1234.5}})
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {"amount": ",.2f"},
        "verbose_map": {},
    }
    form_data = {
        "viz_type": "table",
        "column_config": {"amount": {"d3NumberFormat": ",.1f"}},
    }

    formatted = table(df, form_data, datasource)

    assert formatted["amount"].tolist() == ["1,234.5"]


def test_pivot_table_v2_applies_value_format() -> None:
    """
    Pivot table reports honor `valueFormat` and per-metric `columnFormats`.
    """
    df = pd.DataFrame(
        {"region": ["A", "B"], "sales": [1234.5, 6789.0], "qty": [10.0, 20.0]}
    )
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales", "qty"],
        "aggregateFunction": "Sum",
        "metricsLayout": "COLUMNS",
        "valueFormat": ",.2f",
        "columnFormats": {"qty": ",d"},
    }
    formatted = pivot_table_v2(df, form_data)
    assert formatted[("sales",)].tolist() == ["1,234.50", "6,789.00"]
    assert formatted[("qty",)].tolist() == ["10", "20"]


def test_pivot_table_v2_applies_smart_number_default_without_value_format() -> None:
    """
    Pivot value cells with no ``valueFormat`` and no per-metric format default to
    SMART_NUMBER, mirroring the frontend's ``getNumberFormatter`` default.
    """
    df = pd.DataFrame({"region": ["A", "B"], "sales": [1234567.0, 6789.0]})
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "metricsLayout": "COLUMNS",
    }
    formatted = pivot_table_v2(df, form_data)
    assert formatted[("sales",)].tolist() == ["1.23M", "6.79k"]


def test_pivot_table_v2_applies_datasource_saved_metric_format_without_override() -> (
    None
):
    df = pd.DataFrame({"region": ["A", "B"], "sales": [1234.5, 6789.0]})
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {"sales": ",d"},
        "verbose_map": {},
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "valueFormat": ",.2f",
    }

    formatted = pivot_table_v2(df, form_data, datasource)

    assert formatted[("sales",)].tolist() == ["1,235", "6,789"]


def test_pivot_table_v2_chart_format_overrides_datasource_saved_metric_format() -> None:
    df = pd.DataFrame({"region": ["A"], "sales": [1234.5]})
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {"sales": ",d"},
        "verbose_map": {},
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "valueFormat": ",.2f",
        "columnFormats": {"sales": ",.1f"},
    }

    formatted = pivot_table_v2(df, form_data, datasource)

    assert formatted[("sales",)].tolist() == ["1,234.5"]


def test_pivot_table_v2_applies_per_metric_format_when_metrics_combined() -> None:
    """
    Per-metric formats apply when `combineMetric` moves the metric to the last
    column level.
    """
    df = pd.DataFrame(
        {
            "dept": ["A", "B"],
            "region": ["x", "x"],
            "sales": [100.0, 200.0],
            "qty": [1111.0, 2222.0],
        }
    )
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["dept"],
        "groupbyColumns": ["region"],
        "metrics": ["sales", "qty"],
        "aggregateFunction": "Sum",
        "metricsLayout": "COLUMNS",
        "combineMetric": True,
        "valueFormat": ",.2f",
        "columnFormats": {"qty": ",d"},
    }
    formatted = pivot_table_v2(df, form_data)
    assert formatted[("x", "qty")].tolist() == ["1,111", "2,222"]
    assert formatted[("x", "sales")].tolist() == ["100.00", "200.00"]


def test_table_auto_currency_uses_detected_currency() -> None:
    """
    AUTO currency resolves to the payload's `detected_currency`, or falls back
    to the plain number when detection found mixed currencies.
    """
    form_data = {
        "viz_type": "table",
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
            }
        },
    }
    df = pd.DataFrame.from_dict({"amount": {0: 1234.5}})
    formatted = table(df, form_data, detected_currency="USD")
    assert formatted["amount"].tolist() == ["$ 1,234.50"]

    df = pd.DataFrame.from_dict({"amount": {0: 1234.5}})
    formatted = table(df, form_data, detected_currency=None)
    assert formatted["amount"].tolist() == ["1,234.50"]


def test_table_auto_currency_uses_per_row_currency_context() -> None:
    form_data = {
        "viz_type": "table",
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
            }
        },
    }
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    df = pd.DataFrame(
        {
            "amount": [100.0, 200.0, 300.0, 400.0],
            "currency": ["USD", " eur ", None, "invalid"],
        }
    )

    formatted = table(df, form_data, datasource, detected_currency="GBP")

    assert formatted["amount"].tolist() == [
        "$ 100.00",
        "€ 200.00",
        "300.00",
        "400.00",
    ]


def test_table_small_number_format_resolves_auto_currency_per_row() -> None:
    """
    The browser hands the row's data to the small-number formatter as well, so a
    value below one keeps the row's own currency symbol.
    """
    form_data = {
        "viz_type": "table",
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "d3SmallNumberFormat": ",.3f",
                "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
            }
        },
    }
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    df = pd.DataFrame(
        {
            "amount": [100.0, 0.5, 0.25, 0.75],
            "currency": ["USD", " eur ", None, "invalid"],
        }
    )

    formatted = table(df, form_data, datasource, detected_currency="GBP")

    assert formatted["amount"].tolist() == [
        "$ 100.00",
        "€ 0.500",
        "0.250",
        "0.750",
    ]


def test_table_saved_auto_currency_uses_per_row_currency_context() -> None:
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {"amount": ",.2f"},
        "verbose_map": {},
        "currency_code_column": "currency",
        "metrics": [
            {
                "metric_name": "amount",
                "currency": {"symbol": "AUTO", "symbolPosition": "prefix"},
            }
        ],
    }
    df = pd.DataFrame(
        {
            "amount": [100.0, 200.0],
            "currency": ["USD", "EUR"],
        }
    )

    formatted = table(df, {"viz_type": "table"}, datasource)

    assert formatted["amount"].tolist() == ["$ 100.00", "€ 200.00"]


def test_pivot_table_v2_auto_currency_uses_detected_currency() -> None:
    """
    AUTO currency in pivot tables resolves to the payload's `detected_currency`.
    """
    df = pd.DataFrame({"region": ["A", "B"], "sales": [1234.5, 6789.0]})
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "valueFormat": ",.2f",
        "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
    }
    formatted = pivot_table_v2(df, form_data, detected_currency="EUR")
    assert formatted[("sales",)].tolist() == ["€ 1,234.50", "€ 6,789.00"]


def test_pivot_table_v2_auto_currency_uses_per_cell_currency_context() -> None:
    df = pd.DataFrame(
        {
            "region": ["USD cell", "USD cell", "EUR cell", "Mixed", "Mixed", "Empty"],
            "sales": [100.0, 50.0, 200.0, 300.0, 400.0, 500.0],
            "currency": ["USD", " usd ", "EUR", "USD", "EUR", None],
        }
    )
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "valueFormat": ",.2f",
        "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
    }

    formatted = pivot_table_v2(df, form_data, datasource, detected_currency="GBP")

    assert formatted[("sales",)].to_dict() == {
        ("EUR cell",): "€ 200.00",
        ("Empty",): "£ 500.00",
        ("Mixed",): "700.00",
        ("USD cell",): "$ 150.00",
    }


def test_pivot_table_v2_auto_currency_reads_stored_form_data_key() -> None:
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    df = pd.DataFrame(
        {
            "region": ["US", "EU"],
            "sales": [100.0, 200.0],
            "currency": ["USD", "EUR"],
        }
    )
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "valueFormat": ",.2f",
        "currency_format": {"symbol": "AUTO", "symbolPosition": "prefix"},
    }

    formatted = pivot_table_v2(df, form_data, datasource)

    assert formatted[("sales",)].to_dict() == {
        ("EU",): "€ 200.00",
        ("US",): "$ 100.00",
    }


def test_pivot_table_v2_auto_currency_handles_sparse_2d_pivot() -> None:
    """
    A pivot with both rows and columns has empty cross-product cells. Pandas
    fills those with scalar ``NaN`` rather than an empty currency tuple, which
    must not crash AUTO currency resolution for the whole report.
    """
    df = pd.DataFrame(
        {
            "region": ["EU", "EU", "US"],
            "product": ["a", "b", "a"],
            "sales": [10.0, 20.0, 30.0],
            "currency": ["EUR", "EUR", "USD"],
        }
    )
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": ["product"],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "valueFormat": ",.2f",
        "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
    }

    formatted = pivot_table_v2(df, form_data, datasource, detected_currency="GBP")

    assert formatted[("sales", "a")].to_dict() == {
        ("EU",): "€ 10.00",
        ("US",): "$ 30.00",
    }
    # The missing (US, b) combination stays empty; the present EUR cell formats.
    assert formatted[("sales", "b")].to_dict() == {
        ("EU",): "€ 20.00",
        ("US",): "",
    }


def test_pivot_table_v2_saved_auto_currency_uses_per_cell_context() -> None:
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {"sales": ",.2f"},
        "verbose_map": {},
        "currency_code_column": "currency",
        "metrics": [
            {
                "metric_name": "sales",
                "currency": {"symbol": "AUTO", "symbolPosition": "prefix"},
            }
        ],
    }
    df = pd.DataFrame(
        {
            "region": ["US", "EU"],
            "sales": [100.0, 200.0],
            "currency": ["USD", "EUR"],
        }
    )
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "valueFormat": ",.1f",
    }

    formatted = pivot_table_v2(df, form_data, datasource)

    assert formatted[("sales",)].to_dict() == {
        ("EU",): "€ 200.00",
        ("US",): "$ 100.00",
    }


def test_pivot_table_v2_count_auto_currency_uses_detected_fallback() -> None:
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    df = pd.DataFrame(
        {
            "region": ["US", "EU"],
            "sales": [100.0, 200.0],
            "currency": ["USD", "EUR"],
        }
    )
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales"],
        "aggregateFunction": "Count",
        "valueFormat": ",d",
        "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
    }

    formatted = pivot_table_v2(df, form_data, datasource, detected_currency="GBP")

    assert formatted[("sales",)].to_dict() == {
        ("EU",): "£ 1",
        ("US",): "£ 1",
    }


def test_pivot_table_v2_auto_currency_tracks_mixed_total_context() -> None:
    df = pd.DataFrame(
        {
            "region": ["US", "EU"],
            "sales": [100.0, 200.0],
            "currency": ["USD", "EUR"],
        }
    )
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": [],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "valueFormat": ",.2f",
        "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
        "colTotals": True,
    }

    formatted = pivot_table_v2(df, form_data, datasource)

    assert formatted[("sales",)].to_dict() == {
        ("EU",): "€ 200.00",
        ("US",): "$ 100.00",
        ("Total (Sum)",): "300.00",
    }


def test_pivot_table_v2_auto_currency_tracks_subtotal_context() -> None:
    df = pd.DataFrame(
        {
            "region": ["Mixed", "Mixed", "USD", "USD"],
            "quarter": ["Q1", "Q2", "Q1", "Q2"],
            "sales": [100.0, 200.0, 300.0, 400.0],
            "currency": ["USD", "EUR", "USD", "USD"],
        }
    )
    datasource = MagicMock()
    datasource.data = {
        "column_formats": {},
        "verbose_map": {},
        "currency_code_column": "currency",
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["region"],
        "groupbyColumns": ["quarter"],
        "metrics": ["sales"],
        "aggregateFunction": "Sum",
        "valueFormat": ",.2f",
        "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
        "rowTotals": True,
    }

    formatted = pivot_table_v2(df, form_data, datasource)

    assert formatted.loc[("Mixed",), ("sales", "Q1")] == "$ 100.00"
    assert formatted.loc[("Mixed",), ("sales", "Q2")] == "€ 200.00"
    assert formatted.loc[("Mixed",), ("sales", "Subtotal")] == "300.00"
    assert formatted.loc[("USD",), ("sales", "Subtotal")] == "$ 700.00"
    assert formatted.loc[("Mixed",), ("Total (Sum)", "")] == "300.00"


def test_apply_client_processing_passes_detected_currency() -> None:
    """
    The query payload's `detected_currency` reaches the number formatters.
    """
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.JSON,
                "detected_currency": "USD",
                "data": [{"amount": 1234.5}],
            }
        ]
    }
    form_data = {
        "viz_type": "table",
        "column_config": {
            "amount": {
                "d3NumberFormat": ",.2f",
                "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
            }
        },
    }
    processed = apply_client_processing(result, form_data)
    assert processed["queries"][0]["data"] == {"amount": {0: "$ 1,234.50"}}


def test_pivot_table_v2_applies_per_metric_format_when_metrics_on_rows() -> None:
    """
    Per-metric formats apply when `metricsLayout` is "ROWS" and the metric is
    on the index instead of the columns.
    """
    df = pd.DataFrame(
        {
            "dept": ["A", "B"],
            "region": ["x", "x"],
            "sales": [100.0, 200.0],
            "qty": [1111.0, 2222.0],
        }
    )
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["dept"],
        "groupbyColumns": ["region"],
        "metrics": ["sales", "qty"],
        "aggregateFunction": "Sum",
        "metricsLayout": "ROWS",
        "valueFormat": ",.2f",
        "columnFormats": {"qty": ",d"},
        "currencyFormats": {"sales": {"symbol": "USD", "symbolPosition": "prefix"}},
    }
    formatted = pivot_table_v2(df, form_data)
    assert formatted[("x",)].tolist() == ["$ 100.00", "$ 200.00", "1,111", "2,222"]


def test_pivot_table_v2_applies_per_metric_format_when_metrics_on_rows_combined() -> (
    None
):
    """
    Per-metric formats apply when `metricsLayout` is "ROWS" and `combineMetric`
    moves the metric to the last index level.
    """
    df = pd.DataFrame(
        {
            "dept": ["A", "B"],
            "region": ["x", "x"],
            "sales": [100.0, 200.0],
            "qty": [1111.0, 2222.0],
        }
    )
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyRows": ["dept"],
        "groupbyColumns": ["region"],
        "metrics": ["sales", "qty"],
        "aggregateFunction": "Sum",
        "metricsLayout": "ROWS",
        "combineMetric": True,
        "valueFormat": ",.2f",
        "columnFormats": {"qty": ",d"},
    }
    formatted = pivot_table_v2(df, form_data)
    assert formatted[("x",)].tolist() == ["100.00", "1,111", "200.00", "2,222"]


def test_format_column_applies_d3_and_currency() -> None:
    df = pd.DataFrame({"amount": [1234.5, 6789.0]})
    format_column(df, "amount", ",.2f", {})
    assert df["amount"].tolist() == ["1,234.50", "6,789.00"]

    df = pd.DataFrame({"amount": [1234.5]})
    format_column(df, "amount", ",.2f", {"symbol": "USD", "symbolPosition": "prefix"})
    assert df["amount"].tolist() == ["$ 1,234.50"]


def test_format_column_is_noop_without_format() -> None:
    df = pd.DataFrame({"amount": [1234.5]})
    format_column(df, "amount", None, {})
    assert df["amount"].tolist() == [1234.5]


def test_format_column_preserves_numeric_format_when_currency_is_invalid() -> None:
    df = pd.DataFrame({"amount": [1234.5]})
    format_column(
        df,
        "amount",
        ",.2f",
        {"symbol": {"invalid": True}, "symbolPosition": "prefix"},
    )
    assert df["amount"].tolist() == ["1,234.50"]


def test_format_column_preserves_raw_value_for_invalid_number_format() -> None:
    df = pd.DataFrame({"amount": [1234.5]})
    format_column(df, "amount", "not-a-format", {})
    assert df["amount"].tolist() == ["1234.5"]


def test_apply_pivot_number_formats_resolves_metric_level() -> None:
    df = pd.DataFrame({("sales",): [1234.5], ("qty",): [10.0]})
    df.columns = pd.MultiIndex.from_tuples([("sales",), ("qty",)])
    apply_pivot_number_formats(
        df, {"valueFormat": ",.2f", "columnFormats": {"qty": ",d"}}
    )
    assert df[("sales",)].tolist() == ["1,234.50"]
    assert df[("qty",)].tolist() == ["10"]


def test_apply_pivot_number_formats_metric_at_last_level_when_combined() -> None:
    df = pd.DataFrame({("x", "sales"): [100.0], ("x", "qty"): [1111.0]})
    df.columns = pd.MultiIndex.from_tuples([("x", "sales"), ("x", "qty")])
    apply_pivot_number_formats(
        df,
        {"combineMetric": True, "valueFormat": ",.2f", "columnFormats": {"qty": ",d"}},
    )
    assert df[("x", "sales")].tolist() == ["100.00"]
    assert df[("x", "qty")].tolist() == ["1,111"]


def test_apply_pivot_number_formats_falls_back_to_global_format() -> None:
    df = pd.DataFrame({("sales",): [1234.5]})
    df.columns = pd.MultiIndex.from_tuples([("sales",)])
    apply_pivot_number_formats(
        df, {"valueFormat": ",.2f", "columnFormats": {"sales": ""}}
    )
    assert df[("sales",)].tolist() == ["1,234.50"]


def test_apply_pivot_number_formats_preserves_raw_value_on_format_error() -> None:
    df = pd.DataFrame({("sales",): [1234.5]})
    df.columns = pd.MultiIndex.from_tuples([("sales",)])
    apply_pivot_number_formats(df, {"valueFormat": "not-a-format"})
    assert df[("sales",)].tolist() == ["1234.5"]


def test_apply_pivot_number_formats_auto_currency_without_detection() -> None:
    df = pd.DataFrame({("sales",): [1234.5]})
    df.columns = pd.MultiIndex.from_tuples([("sales",)])
    apply_pivot_number_formats(
        df,
        {
            "valueFormat": ",.2f",
            "currencyFormat": {"symbol": "AUTO", "symbolPosition": "prefix"},
        },
    )
    assert df[("sales",)].tolist() == ["1,234.50"]


def test_apply_client_processing_no_form_invalid_viz_type():
    """
    Test with invalid viz type. It should just return the result
    """
    result = {"foo": "bar"}
    form_data = {"viz_type": "baz"}
    assert apply_client_processing(result, form_data) == result


def test_apply_client_processing_without_result_format():
    """
    A query without result_format should raise an exception
    """
    result = {"queries": [{"result_format": "foo"}]}
    form_data = {"viz_type": "pivot_table_v2"}

    with pytest.raises(Exception) as ex:  # noqa: PT011
        apply_client_processing(result, form_data)

    assert ex.match("Result format foo not supported") is True  # noqa: E712


def test_apply_client_processing_json_format():
    """
    It should be able to process json results
    """

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.JSON,
                "data": {
                    "result": [
                        {
                            "data": [{"COUNT(is_software_dev)": 4725}],
                            "colnames": ["COUNT(is_software_dev)"],
                            "coltypes": [0],
                        }
                    ]
                },
            }
        ]
    }
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [
            {
                "aggregate": "COUNT",
                "column": {
                    "column_name": "is_software_dev",
                    "description": None,
                    "expression": None,
                    "filterable": True,
                    "groupby": True,
                    "id": 1463,
                    "is_dttm": False,
                    "python_date_format": None,
                    "type": "DOUBLE PRECISION",
                    "verbose_name": None,
                },
                "expressionType": "SIMPLE",
                "hasCustomLabel": False,
                "isNew": False,
                "label": "COUNT(is_software_dev)",
                "optionName": "metric_9i1kctig9yr_sizo6ihd2o",
                "sqlExpression": None,
            }
        ],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [
            {
                "clause": "WHERE",
                "comparator": "Currently A Developer",
                "expressionType": "SIMPLE",
                "filterOptionName": "filter_fvi0jg9aii_2lekqrhy7qk",
                "isExtra": False,
                "isNew": False,
                "operator": "==",
                "sqlExpression": None,
                "subject": "developer_type",
            }
        ],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "json",
        "result_type": "results",
    }

    assert apply_client_processing(result, form_data) == {
        "queries": [
            {
                "result_format": ChartDataResultFormat.JSON,
                "data": {
                    "result": {
                        "Total (Sum)": {
                            "data": [{"COUNT(is_software_dev)": 4725}],
                            "colnames": ["COUNT(is_software_dev)"],
                            "coltypes": [0],
                        }
                    }
                },
                "colnames": [("result",)],
                "indexnames": [("Total (Sum)",)],
                "coltypes": [GenericDataType.STRING],
                "rowcount": 1,
            }
        ]
    }


def test_apply_client_processing_csv_format():
    """
    It should be able to process csv results
    """

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": """
COUNT(is_software_dev)
4725
""",
            }
        ]
    }
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [
            {
                "aggregate": "COUNT",
                "column": {
                    "column_name": "is_software_dev",
                    "description": None,
                    "expression": None,
                    "filterable": True,
                    "groupby": True,
                    "id": 1463,
                    "is_dttm": False,
                    "python_date_format": None,
                    "type": "DOUBLE PRECISION",
                    "verbose_name": None,
                },
                "expressionType": "SIMPLE",
                "hasCustomLabel": False,
                "isNew": False,
                "label": "COUNT(is_software_dev)",
                "optionName": "metric_9i1kctig9yr_sizo6ihd2o",
                "sqlExpression": None,
            }
        ],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [
            {
                "clause": "WHERE",
                "comparator": "Currently A Developer",
                "expressionType": "SIMPLE",
                "filterOptionName": "filter_fvi0jg9aii_2lekqrhy7qk",
                "isExtra": False,
                "isNew": False,
                "operator": "==",
                "sqlExpression": None,
                "subject": "developer_type",
            }
        ],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "json",
        "result_type": "results",
    }

    assert apply_client_processing(result, form_data) == {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": ",COUNT(is_software_dev)\nTotal (Sum),4725\n",
                "colnames": [("COUNT(is_software_dev)",)],
                "indexnames": [("Total (Sum)",)],
                "coltypes": [GenericDataType.NUMERIC],
                "rowcount": 1,
            }
        ]
    }


def test_apply_client_processing_csv_format_simple_table():
    """
    It should be able to process csv results
    And not show a default column
    """

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": """
COUNT(is_software_dev)
4725
""",
            }
        ]
    }
    form_data = {
        "datasource": "19__table",
        "viz_type": "table",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [
            {
                "aggregate": "COUNT",
                "column": {
                    "column_name": "is_software_dev",
                    "description": None,
                    "expression": None,
                    "filterable": True,
                    "groupby": True,
                    "id": 1463,
                    "is_dttm": False,
                    "python_date_format": None,
                    "type": "DOUBLE PRECISION",
                    "verbose_name": None,
                },
                "expressionType": "SIMPLE",
                "hasCustomLabel": False,
                "isNew": False,
                "label": "COUNT(is_software_dev)",
                "optionName": "metric_9i1kctig9yr_sizo6ihd2o",
                "sqlExpression": None,
            }
        ],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [
            {
                "clause": "WHERE",
                "comparator": "Currently A Developer",
                "expressionType": "SIMPLE",
                "filterOptionName": "filter_fvi0jg9aii_2lekqrhy7qk",
                "isExtra": False,
                "isNew": False,
                "operator": "==",
                "sqlExpression": None,
                "subject": "developer_type",
            }
        ],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "json",
        "result_type": "results",
    }

    assert apply_client_processing(result, form_data) == {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": "COUNT(is_software_dev)\n4725\n",
                "colnames": ["COUNT(is_software_dev)"],
                "indexnames": [0],
                "coltypes": [GenericDataType.NUMERIC],
                "rowcount": 1,
            }
        ]
    }


def test_apply_client_processing_csv_format_escapes_formula_values():
    """
    A value starting with a formula trigger should be escaped in the CSV
    output, consistent with the other CSV export paths.
    """

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": "is_software_dev\n=SUM(1+1)\n",
            }
        ]
    }
    form_data = {
        "datasource": "19__table",
        "viz_type": "table",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "csv",
        "result_type": "results",
    }

    processed = apply_client_processing(result, form_data)
    # the leading "=" is neutralized with a single-quote prefix
    assert "'=SUM(1+1)" in processed["queries"][0]["data"]
    assert "\n=SUM(1+1)" not in processed["queries"][0]["data"]


def test_apply_client_processing_csv_format_bytes_data():
    """
    Regression for #32370: "Export to pivoted .csv" fails with a 505 error.

    For a real ``resultType=post_processed``/``resultFormat=csv`` request,
    the query's ``data`` is CSV *bytes* by the time it reaches
    ``apply_client_processing`` --
    ``QueryContextProcessor.get_data`` encodes it to bytes for any CSV
    result_format (``superset/common/query_context_processor.py``), and
    that's the exact payload ``ChartDataRestApi._send_chart_response`` hands
    to ``apply_client_processing`` for a POST_PROCESSED result. The CSV
    branch here passes that straight into ``StringIO(data)``, which raises
    ``TypeError: initial_value must be str or None, not bytes`` -- a crash
    that reproduces for every pivoted CSV export, not only the reporter's
    special-character metric labels (this test uses one anyway, to also
    pin the originally reported symptom).
    """

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": b"name,% of total\nA,1\nB,2\n",
            }
        ]
    }
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "groupbyColumns": [],
        "groupbyRows": ["name"],
        "metrics": ["% of total"],
        "metricsLayout": "COLUMNS",
        "aggregateFunction": "Sum",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "result_format": "csv",
        "result_type": "post_processed",
    }

    processed = apply_client_processing(result, form_data)
    assert "% of total" in processed["queries"][0]["data"]


@with_config({"CSV_EXPORT": {"encoding": "latin-1"}})
def test_apply_client_processing_csv_format_bytes_data_non_default_encoding():
    """
    Regression for #32370: the CSV bytes payload must be decoded with the
    configured ``CSV_EXPORT.encoding``, not a hardcoded ``utf-8``, since
    ``QueryContextProcessor.get_data`` encodes with that same config value.
    A UTF-8 decode of latin-1 bytes containing e.g. "é" (0xE9) would raise
    ``UnicodeDecodeError`` instead of the intended CSV text.
    """

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": "name,city\nA,Montr\xe9al\n".encode("latin-1"),
            }
        ]
    }
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "groupbyColumns": [],
        "groupbyRows": ["name"],
        "metrics": ["city"],
        "metricsLayout": "COLUMNS",
        "aggregateFunction": "Sum",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "result_format": "csv",
        "result_type": "post_processed",
    }

    processed = apply_client_processing(result, form_data)
    assert "Montréal" in processed["queries"][0]["data"]


def test_apply_client_processing_csv_format_empty_string():
    """
    It should be able to process csv results with no data
    """

    result = {"queries": [{"result_format": ChartDataResultFormat.CSV, "data": ""}]}
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [
            {
                "aggregate": "COUNT",
                "column": {
                    "column_name": "is_software_dev",
                    "description": None,
                    "expression": None,
                    "filterable": True,
                    "groupby": True,
                    "id": 1463,
                    "is_dttm": False,
                    "python_date_format": None,
                    "type": "DOUBLE PRECISION",
                    "verbose_name": None,
                },
                "expressionType": "SIMPLE",
                "hasCustomLabel": False,
                "isNew": False,
                "label": "COUNT(is_software_dev)",
                "optionName": "metric_9i1kctig9yr_sizo6ihd2o",
                "sqlExpression": None,
            }
        ],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [
            {
                "clause": "WHERE",
                "comparator": "Currently A Developer",
                "expressionType": "SIMPLE",
                "filterOptionName": "filter_fvi0jg9aii_2lekqrhy7qk",
                "isExtra": False,
                "isNew": False,
                "operator": "==",
                "sqlExpression": None,
                "subject": "developer_type",
            }
        ],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "json",
        "result_type": "results",
    }

    assert apply_client_processing(result, form_data) == {
        "queries": [{"result_format": ChartDataResultFormat.CSV, "data": ""}]
    }


@pytest.mark.parametrize("data", [None, "", "\n", b"", b"\n"])
def test_apply_client_processing_csv_format_no_data(data):
    """
    It should be able to process csv results with no data, including the
    bytes forms ``QueryContextProcessor.get_data`` actually produces for a
    columnless frame (e.g. a bare newline), which must be decoded before
    the empty-data check runs or they'd reach ``pd.read_csv`` and raise
    ``EmptyDataError``.
    """

    result = {"queries": [{"result_format": ChartDataResultFormat.CSV, "data": data}]}
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [
            {
                "aggregate": "COUNT",
                "column": {
                    "column_name": "is_software_dev",
                    "description": None,
                    "expression": None,
                    "filterable": True,
                    "groupby": True,
                    "id": 1463,
                    "is_dttm": False,
                    "python_date_format": None,
                    "type": "DOUBLE PRECISION",
                    "verbose_name": None,
                },
                "expressionType": "SIMPLE",
                "hasCustomLabel": False,
                "isNew": False,
                "label": "COUNT(is_software_dev)",
                "optionName": "metric_9i1kctig9yr_sizo6ihd2o",
                "sqlExpression": None,
            }
        ],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [
            {
                "clause": "WHERE",
                "comparator": "Currently A Developer",
                "expressionType": "SIMPLE",
                "filterOptionName": "filter_fvi0jg9aii_2lekqrhy7qk",
                "isExtra": False,
                "isNew": False,
                "operator": "==",
                "sqlExpression": None,
                "subject": "developer_type",
            }
        ],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "json",
        "result_type": "results",
    }

    assert apply_client_processing(result, form_data) == {
        "queries": [{"result_format": ChartDataResultFormat.CSV, "data": data}]
    }


def test_apply_client_processing_csv_format_no_data_multiple_queries():
    """
    It should be able to process csv results multiple queries if one query has no data
    """

    result = {
        "queries": [
            {"result_format": ChartDataResultFormat.CSV, "data": ""},
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": """
COUNT(is_software_dev)
4725
""",
            },
        ]
    }
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [
            {
                "aggregate": "COUNT",
                "column": {
                    "column_name": "is_software_dev",
                    "description": None,
                    "expression": None,
                    "filterable": True,
                    "groupby": True,
                    "id": 1463,
                    "is_dttm": False,
                    "python_date_format": None,
                    "type": "DOUBLE PRECISION",
                    "verbose_name": None,
                },
                "expressionType": "SIMPLE",
                "hasCustomLabel": False,
                "isNew": False,
                "label": "COUNT(is_software_dev)",
                "optionName": "metric_9i1kctig9yr_sizo6ihd2o",
                "sqlExpression": None,
            }
        ],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [
            {
                "clause": "WHERE",
                "comparator": "Currently A Developer",
                "expressionType": "SIMPLE",
                "filterOptionName": "filter_fvi0jg9aii_2lekqrhy7qk",
                "isExtra": False,
                "isNew": False,
                "operator": "==",
                "sqlExpression": None,
                "subject": "developer_type",
            }
        ],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "json",
        "result_type": "results",
    }

    assert apply_client_processing(result, form_data) == {
        "queries": [
            {"result_format": ChartDataResultFormat.CSV, "data": ""},
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": ",COUNT(is_software_dev)\nTotal (Sum),4725\n",
                "colnames": [("COUNT(is_software_dev)",)],
                "indexnames": [("Total (Sum)",)],
                "coltypes": [GenericDataType.NUMERIC],
                "rowcount": 1,
            },
        ]
    }


def test_apply_client_processing_json_format_empty_string():
    """
    It should be able to process json results with no data
    """

    result = {"queries": [{"result_format": ChartDataResultFormat.JSON, "data": ""}]}
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [
            {
                "aggregate": "COUNT",
                "column": {
                    "column_name": "is_software_dev",
                    "description": None,
                    "expression": None,
                    "filterable": True,
                    "groupby": True,
                    "id": 1463,
                    "is_dttm": False,
                    "python_date_format": None,
                    "type": "DOUBLE PRECISION",
                    "verbose_name": None,
                },
                "expressionType": "SIMPLE",
                "hasCustomLabel": False,
                "isNew": False,
                "label": "COUNT(is_software_dev)",
                "optionName": "metric_9i1kctig9yr_sizo6ihd2o",
                "sqlExpression": None,
            }
        ],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [
            {
                "clause": "WHERE",
                "comparator": "Currently A Developer",
                "expressionType": "SIMPLE",
                "filterOptionName": "filter_fvi0jg9aii_2lekqrhy7qk",
                "isExtra": False,
                "isNew": False,
                "operator": "==",
                "sqlExpression": None,
                "subject": "developer_type",
            }
        ],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "json",
        "result_type": "results",
    }

    assert apply_client_processing(result, form_data) == {
        "queries": [{"result_format": ChartDataResultFormat.JSON, "data": ""}]
    }


def test_apply_client_processing_json_format_data_is_none():
    """
    It should be able to process json results with no data
    """

    result = {"queries": [{"result_format": ChartDataResultFormat.JSON, "data": None}]}
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [
            {
                "aggregate": "COUNT",
                "column": {
                    "column_name": "is_software_dev",
                    "description": None,
                    "expression": None,
                    "filterable": True,
                    "groupby": True,
                    "id": 1463,
                    "is_dttm": False,
                    "python_date_format": None,
                    "type": "DOUBLE PRECISION",
                    "verbose_name": None,
                },
                "expressionType": "SIMPLE",
                "hasCustomLabel": False,
                "isNew": False,
                "label": "COUNT(is_software_dev)",
                "optionName": "metric_9i1kctig9yr_sizo6ihd2o",
                "sqlExpression": None,
            }
        ],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [
            {
                "clause": "WHERE",
                "comparator": "Currently A Developer",
                "expressionType": "SIMPLE",
                "filterOptionName": "filter_fvi0jg9aii_2lekqrhy7qk",
                "isExtra": False,
                "isNew": False,
                "operator": "==",
                "sqlExpression": None,
                "subject": "developer_type",
            }
        ],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "json",
        "result_type": "results",
    }

    assert apply_client_processing(result, form_data) == {
        "queries": [{"result_format": ChartDataResultFormat.JSON, "data": None}]
    }


def test_apply_client_processing_verbose_map(session: Session):
    from superset import db
    from superset.connectors.sqla.models import SqlaTable, SqlMetric
    from superset.models.core import Database

    engine = db.session.get_bind()
    SqlaTable.metadata.create_all(engine)  # pylint: disable=no-member
    database = Database(database_name="my_database", sqlalchemy_uri="sqlite://")
    sqla_table = SqlaTable(
        table_name="my_sqla_table",
        columns=[],
        metrics=[
            SqlMetric(
                metric_name="count",
                verbose_name="COUNT(*)",
                metric_type="count",
                expression="COUNT(*)",
            )
        ],
        database=database,
    )

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.JSON,
                "data": [{"count": 4725}],
            }
        ]
    }
    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": ["COUNT(*)"],
        "metricsLayout": "COLUMNS",
        "row_limit": 10000,
        "order_desc": True,
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "json",
        "result_type": "results",
    }

    assert apply_client_processing(result, form_data, datasource=sqla_table) == {
        "queries": [
            {
                "result_format": ChartDataResultFormat.JSON,
                "data": {"COUNT(*)": {"Total (Sum)": "4.73k"}},
                "colnames": [("COUNT(*)",)],
                "indexnames": [("Total (Sum)",)],
                "coltypes": [GenericDataType.STRING],
                "rowcount": 1,
            }
        ]
    }


def test_pivot_multi_level_index():
    """
    Pivot table with multi-level indexing.
    """
    arrays = [
        ["Region1", "Region1", "Region1", "Region2", "Region2", "Region2"],
        ["State1", "State1", "State2", "State3", "State3", "State4"],
        ["City1", "City2", "City3", "City4", "City5", "City6"],
    ]
    index = pd.MultiIndex.from_tuples(
        list(zip(*arrays, strict=False)),
        names=["Region", "State", "City"],
    )

    data = {
        "Metric1": [10, 20, 30, 40, 50, 60],
        "Metric2": [5, 10, 15, 20, 25, 30],
        "Metric3": [None, None, None, None, None, None],
    }
    df = pd.DataFrame(data, index=index)

    pivoted = pivot_df(
        df,
        rows=["Region", "State", "City"],
        columns=[],
        metrics=["Metric1", "Metric2", "Metric3"],
        aggfunc="Sum",
        transpose_pivot=False,
        combine_metrics=False,
        show_rows_total=False,
        show_columns_total=True,
        apply_metrics_on_rows=False,
    )

    # Sort the pivoted DataFrame to ensure deterministic output
    pivoted_sorted = pivoted.sort_index()

    assert (
        pivoted_sorted.to_markdown()
        == """
|                                   |   ('Metric1',) |   ('Metric2',) |   ('Metric3',) |
|:----------------------------------|---------------:|---------------:|---------------:|
| ('Region1', 'State1', 'City1')    |             10 |              5 |            nan |
| ('Region1', 'State1', 'City2')    |             20 |             10 |            nan |
| ('Region1', 'State1', 'Subtotal') |             30 |             15 |              0 |
| ('Region1', 'State2', 'City3')    |             30 |             15 |            nan |
| ('Region1', 'State2', 'Subtotal') |             30 |             15 |              0 |
| ('Region1', 'Subtotal', '')       |             60 |             30 |              0 |
| ('Region2', 'State3', 'City4')    |             40 |             20 |            nan |
| ('Region2', 'State3', 'City5')    |             50 |             25 |            nan |
| ('Region2', 'State3', 'Subtotal') |             90 |             45 |              0 |
| ('Region2', 'State4', 'City6')    |             60 |             30 |            nan |
| ('Region2', 'State4', 'Subtotal') |             60 |             30 |              0 |
| ('Region2', 'Subtotal', '')       |            150 |             75 |              0 |
| ('Total (Sum)', '', '')           |            210 |            105 |              0 |
    """.strip()
    )


@with_config({"REPORTS_CSV_NA_NAMES": []})
def test_apply_client_processing_csv_format_preserves_na_strings():
    """
    Test that apply_client_processing preserves "NA" when REPORTS_CSV_NA_NAMES is [].

    This ensures that scheduled reports can be configured to
    preserve strings like "NA" as literal values.
    """

    # CSV data with "NA" string that should be preserved
    csv_data = "first_name,last_name\nJeff,Smith\nAlice,NA"

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": csv_data,
            }
        ]
    }

    form_data = {
        "datasource": "1__table",
        "viz_type": "table",
        "slice_id": 1,
        "url_params": {},
        "metrics": [],
        "groupby": [],
        "columns": ["first_name", "last_name"],
        "extra_form_data": {},
        "force": False,
        "result_format": "csv",
        "result_type": "results",
    }

    # Test with REPORTS_CSV_NA_NAMES set to empty list (disable NA conversion)

    processed_result = apply_client_processing(result, form_data)

    # Verify the CSV data still contains "NA" as string, not converted to null
    output_data = processed_result["queries"][0]["data"]
    assert "NA" in output_data
    # The "NA" should be preserved in the output CSV
    lines = output_data.strip().split("\n")
    assert "Alice,NA" in lines[2]  # Second data row should preserve "NA"


@with_config({"REPORTS_CSV_NA_NAMES": ["MISSING"]})
def test_apply_client_processing_csv_format_custom_na_values():
    """
    Test that apply_client_processing respects custom NA values configuration.
    """

    csv_data = "name,status\nJeff,MISSING\nAlice,OK"

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": csv_data,
            }
        ]
    }

    form_data = {
        "datasource": "1__table",
        "viz_type": "table",
        "slice_id": 1,
        "url_params": {},
        "metrics": [],
        "groupby": [],
        "columns": ["name", "status"],
        "extra_form_data": {},
        "force": False,
        "result_format": "csv",
        "result_type": "results",
    }

    # Test with custom NA values - only "MISSING" should be treated as NA
    processed_result = apply_client_processing(result, form_data)

    output_data = processed_result["queries"][0]["data"]
    lines = output_data.strip().split("\n")
    assert len(lines) >= 3  # header + 2 data rows
    assert "Jeff," in lines[1]  # First data row should have empty status after "Jeff,"
    assert "Alice,OK" in lines[2]  # Second data row should preserve "OK"


@with_config({"REPORTS_CSV_NA_NAMES": []})
def test_apply_client_processing_csv_format_default_na_behavior():
    """
    Test that apply_client_processing uses default pandas NA behavior
    when REPORTS_CSV_NA_NAMES is not configured.
    This ensures backwards compatibility.
    """

    # CSV data with "NA" string that should be converted to null in default behavior
    csv_data = "first_name,last_name\nJeff,Smith\nAlice,NA"

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": csv_data,
            }
        ]
    }

    form_data = {
        "datasource": "1__table",
        "viz_type": "table",
        "slice_id": 1,
        "url_params": {},
        "metrics": [],
        "groupby": [],
        "columns": ["first_name", "last_name"],
        "extra_form_data": {},
        "force": False,
        "result_format": "csv",
        "result_type": "results",
    }

    processed_result = apply_client_processing(result, form_data)

    # Verify the CSV data has "NA" converted to empty (default pandas behavior)
    output_data = processed_result["queries"][0]["data"]
    lines = output_data.strip().split("\n")
    assert len(lines) >= 3  # header + 2 data rows
    # The "NA" should be converted to empty by default pandas behavior
    assert (
        "Alice," in lines[2]
    )  # Second data row should have empty last_name (NA converted to null)


def _assert_xlsx_client_processing(index: bool) -> None:
    """Assert XLSX post-processing preserves columns and configured index."""
    source_df = pd.DataFrame(
        {
            "city": ["Paris", "London"],
            "value": [10, 20],
        },
        index=pd.Index(["row-1", "row-2"], name="row"),
    )
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.XLSX,
                "data": excel.df_to_excel(source_df, index=index),
            }
        ]
    }
    form_data = {
        "viz_type": "table",
        "columns": ["city", "value"],
        "metrics": [],
    }

    processed_result = apply_client_processing(result, form_data)
    query = processed_result["queries"][0]
    output_df = pd.read_excel(
        BytesIO(query["data"]),
        index_col=0 if index else None,
    )
    expected_df = source_df if index else source_df.reset_index(drop=True)

    pd.testing.assert_frame_equal(output_df, expected_df, check_names=False)
    assert query["colnames"] == ["city", "value"]
    if index:
        assert query["indexnames"] == ["row-1", "row-2"]
    else:
        assert query["indexnames"] == [0, 1]
    assert query["rowcount"] == 2


@with_config({"EXCEL_EXPORT": {"index": True}})
def test_apply_client_processing_xlsx_format_with_index() -> None:
    """XLSX post-processing should preserve an exported index."""
    _assert_xlsx_client_processing(index=True)


@with_config({"EXCEL_EXPORT": {"index": False}})
def test_apply_client_processing_xlsx_format_without_index() -> None:
    """XLSX post-processing should not shift columns when index is omitted."""
    _assert_xlsx_client_processing(index=False)


@with_config({"EXCEL_EXPORT": {}})
def test_apply_client_processing_xlsx_format_without_index_default_config() -> None:
    """XLSX post-processing derives omitted index from the payload."""
    _assert_xlsx_client_processing(index=False)


@with_config({"EXCEL_EXPORT": {"index": False}})
def test_apply_client_processing_xlsx_format_pivot_table_groupby_columns() -> None:
    """XLSX post-processing should preserve pivot groupby columns."""
    source_df = pd.DataFrame(
        {
            "city": ["Paris", "Paris", "London"],
            "segment": ["Consumer", "Corporate", "Consumer"],
            "value": [10, 20, 30],
        },
    )
    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.XLSX,
                "data": excel.df_to_excel(source_df, index=False),
            }
        ]
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyColumns": ["segment"],
        "groupbyRows": ["city"],
        "metrics": ["value"],
    }

    processed_result = apply_client_processing(result, form_data)
    query = processed_result["queries"][0]
    output_df = pd.read_excel(BytesIO(query["data"]), index_col=0)

    assert query["rowcount"] == 2
    assert query["indexnames"] == [("London",), ("Paris",)]
    assert set(output_df.index) == {"London", "Paris"}
    assert "value Consumer" in output_df.columns
    assert "value Corporate" in output_df.columns


@pytest.mark.parametrize(
    "result_format", [ChartDataResultFormat.XLSX, ChartDataResultFormat.CSV]
)
@pytest.mark.parametrize("selection", ["columns", "all_columns", "groupby"])
@pytest.mark.parametrize("verbose", [False, True])
@pytest.mark.parametrize("values", [["Books", "Games"], [0, 1]])
def test_table_export_preserves_selected_dimension(
    result_format: ChartDataResultFormat,
    selection: str,
    verbose: bool,
    values: list[str] | list[int],
) -> None:
    """Selected dimensions remain columns even with labels or index-like values."""
    labels: dict[str, str] = (
        {"category": "Category", "total_amount": "Total amount"} if verbose else {}
    )
    source_df: pd.DataFrame = pd.DataFrame(
        {"category": values, "total_amount": [15, 20]}
    ).rename(columns=labels)
    datasource: MagicMock = MagicMock()
    datasource.data = {"verbose_map": labels}
    form_data: dict[str, Any] = {
        "viz_type": "table",
        selection: ["category"],
        "metrics": ["total_amount"],
    }
    result: dict[str, Any] = {
        "queries": [
            {
                "result_format": result_format,
                "data": excel.df_to_excel(source_df, index=False)
                if result_format == ChartDataResultFormat.XLSX
                else source_df.to_csv(index=False),
            }
        ]
    }
    query: dict[str, Any] = apply_client_processing(result, form_data, datasource)[
        "queries"
    ][0]
    output_df: pd.DataFrame = (
        pd.read_excel(BytesIO(query["data"]))
        if result_format == ChartDataResultFormat.XLSX
        else pd.read_csv(StringIO(query["data"]))
    )
    pd.testing.assert_frame_equal(output_df, source_df)
    assert query["colnames"] == source_df.columns.tolist()


@pytest.mark.parametrize(
    "result_format", [ChartDataResultFormat.XLSX, ChartDataResultFormat.CSV]
)
def test_serialized_export_does_not_rename_display_label_twice(
    result_format: ChartDataResultFormat,
) -> None:
    """A display label can equal another column's raw name without remapping."""
    # Avoid app-init regression: QueryContextProcessor imports SQL dataset models.
    from superset.common.query_context_processor import QueryContextProcessor

    labels: dict[str, str] = {
        "category": "total_amount",
        "total_amount": "Total amount",
    }
    source: pd.DataFrame = pd.DataFrame(
        {"category": ["Books", "Games"], "total_amount": [15, 20]}
    )
    expected: pd.DataFrame = source.rename(columns=labels)
    datasource: MagicMock = MagicMock(data={"verbose_map": labels})
    processor: QueryContextProcessor = QueryContextProcessor(
        MagicMock(datasource=datasource, result_format=result_format)
    )
    result: dict[str, Any] = {
        "queries": [
            {
                "result_format": result_format,
                "data": processor.get_data(
                    source, [GenericDataType.STRING, GenericDataType.NUMERIC]
                ),
            }
        ]
    }
    query: dict[str, Any] = apply_client_processing(
        result,
        {"viz_type": "table", "groupby": ["category"], "metrics": ["total_amount"]},
        datasource,
    )["queries"][0]
    actual: pd.DataFrame = (
        pd.read_excel(BytesIO(query["data"]))
        if result_format == ChartDataResultFormat.XLSX
        else pd.read_csv(StringIO(query["data"]))
    )
    pd.testing.assert_frame_equal(actual, expected)
    assert query["colnames"] == ["total_amount", "Total amount"]


@pytest.mark.parametrize(
    ("column", "values", "columns"),
    [("Category", ["Books", "Games"], []), (2024, [0, 1], ["2024"])],
)
def test_table_xlsx_preserves_unlisted_and_numeric_headers(
    column: str | int, values: list[str] | list[int], columns: list[str]
) -> None:
    """Retain fallback index names and recognize numeric-looking selections."""
    source: pd.DataFrame = pd.DataFrame({column: values, "Total": [15, 20]})
    result: dict[str, Any] = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.XLSX,
                "data": excel.df_to_excel(source, index=False),
            }
        ]
    }
    query: dict[str, Any] = apply_client_processing(
        result, {"viz_type": "table", "groupby": columns, "metrics": ["Total"]}
    )["queries"][0]
    actual: pd.DataFrame = pd.read_excel(BytesIO(query["data"]))
    pd.testing.assert_frame_equal(actual, source)


def test_table_xlsx_preserves_percent_metric_header() -> None:
    """A percent-only selection is a data column, not an unnamed row index."""
    source_df: pd.DataFrame = pd.DataFrame({"%count": [0, 1]})
    result: dict[str, Any] = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.XLSX,
                "data": excel.df_to_excel(source_df, index=False),
            }
        ]
    }
    query: dict[str, Any] = apply_client_processing(
        result, {"viz_type": "table", "percent_metrics": ["count"]}
    )["queries"][0]
    output_df: pd.DataFrame = pd.read_excel(BytesIO(query["data"]))
    pd.testing.assert_frame_equal(output_df, source_df)


@pytest.mark.parametrize("viz_type", ["table", "pivot_table_v2"])
@pytest.mark.parametrize(
    "result_format", [ChartDataResultFormat.XLSX, ChartDataResultFormat.CSV]
)
def test_sql_dataset_export_headers(
    session: Session, viz_type: str, result_format: ChartDataResultFormat
) -> None:
    """SQL Table and multi-row Pivot exports preserve their header contract."""
    # avoid app-init regression: SQL dataset models require initialized field factories.
    from superset.connectors.sqla.models import SqlaTable, SqlMetric, TableColumn
    from superset.models.core import Database

    datasource: SqlaTable = SqlaTable(
        table_name="orders",
        columns=[
            TableColumn(column_name="category", type="VARCHAR"),
            TableColumn(column_name="region", type="VARCHAR"),
        ],
        metrics=[SqlMetric(metric_name="total_amount", expression="SUM(amount)")],
        database=Database(database_name="export_test", sqlalchemy_uri="sqlite://"),
    )
    source_df: pd.DataFrame = pd.DataFrame(
        {"category": ["Books"], "region": ["West"], "total_amount": [15]}
    )
    form_data: dict[str, Any] = {
        "viz_type": viz_type,
        "columns": ["category", "region"],
        "groupbyRows": ["category", "region"],
        "groupbyColumns": [],
        "metrics": ["total_amount"],
    }
    result: dict[str, Any] = {
        "queries": [
            {
                "result_format": result_format,
                "data": excel.df_to_excel(source_df, index=False)
                if result_format == ChartDataResultFormat.XLSX
                else source_df.to_csv(index=False),
            }
        ]
    }
    query: dict[str, Any] = apply_client_processing(result, form_data, datasource)[
        "queries"
    ][0]
    output_df: pd.DataFrame = (
        pd.read_excel(BytesIO(query["data"]))
        if result_format == ChartDataResultFormat.XLSX
        else pd.read_csv(StringIO(query["data"]))
    )
    pd.testing.assert_frame_equal(output_df, source_df)


@with_config({"CSV_EXPORT": {"sep": ";", "decimal": ","}})
def test_apply_client_processing_csv_format_custom_delimiter():
    """
    Test that apply_client_processing respects CSV_EXPORT sep and decimal config.
    Without the fix, pd.read_csv() uses default comma separator and fails to parse
    semicolon-delimited CSV correctly, causing HTTP 500 in email reports.
    """
    csv_data = "name;value\nfoo;1,5\nbar;2,0"

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": csv_data,
            }
        ]
    }

    form_data = {
        "datasource": "1__table",
        "viz_type": "table",
        "slice_id": 1,
        "url_params": {},
        "metrics": [],
        "groupby": [],
        "columns": ["name", "value"],
        "extra_form_data": {},
        "force": False,
        "result_format": "csv",
        "result_type": "results",
    }

    processed_result = apply_client_processing(result, form_data)

    output_data = processed_result["queries"][0]["data"]
    lines = output_data.strip().split("\n")
    # Should have header + 2 data rows, with correct column parsing
    assert len(lines) == 3
    # name and value should be separate columns, not merged into one
    assert processed_result["queries"][0]["colnames"] == ["name", "value"]
    # Output CSV must also use the configured separator and decimal
    assert lines[0] == "name;value", f"Expected semicolon header, got: {lines[0]}"
    assert "1,5" in lines[1], f"Expected comma decimal in row 1, got: {lines[1]}"
    assert "2,0" in lines[2], f"Expected comma decimal in row 2, got: {lines[2]}"


@with_config({"CSV_EXPORT": {"sep": ";", "decimal": ","}})
def test_apply_client_processing_pivot_table_v2_custom_sep_decimal():
    """
    Test that pivot_table_v2 respects CSV_EXPORT sep and decimal config.
    pivot_table_v2 performs DataFrame manipulations before writing to CSV, so we
    verify that the final to_csv() call correctly uses the configured sep and decimal.
    """
    csv_data = "COUNT(is_software_dev)\n4725"

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": csv_data,
            }
        ]
    }

    form_data = {
        "datasource": "19__table",
        "viz_type": "pivot_table_v2",
        "slice_id": 69,
        "url_params": {},
        "granularity_sqla": "time_start",
        "time_grain_sqla": "P1D",
        "time_range": "No filter",
        "groupbyColumns": [],
        "groupbyRows": [],
        "metrics": [
            {
                "aggregate": "COUNT",
                "column": {
                    "column_name": "is_software_dev",
                    "description": None,
                    "expression": None,
                    "filterable": True,
                    "groupby": True,
                    "id": 1463,
                    "is_dttm": False,
                    "python_date_format": None,
                    "type": "DOUBLE PRECISION",
                    "verbose_name": None,
                },
                "expressionType": "SIMPLE",
                "hasCustomLabel": False,
                "isNew": False,
                "label": "COUNT(is_software_dev)",
                "optionName": "metric_9i1kctig9yr_sizo6ihd2o",
                "sqlExpression": None,
            }
        ],
        "metricsLayout": "COLUMNS",
        "adhoc_filters": [],
        "row_limit": 10000,
        "order_desc": True,
        "aggregateFunction": "Sum",
        "valueFormat": "SMART_NUMBER",
        "date_format": "smart_date",
        "rowOrder": "key_a_to_z",
        "colOrder": "key_a_to_z",
        "extra_form_data": {},
        "force": False,
        "result_format": "csv",
        "result_type": "results",
    }

    processed_result = apply_client_processing(result, form_data)

    output_data = processed_result["queries"][0]["data"]
    lines = output_data.strip().split("\n")
    # pivot_table_v2 adds a row index (Total (Sum)) and produces output
    # with the configured sep
    assert len(lines) == 2
    # Output must use the configured separator
    assert ";" in lines[0], f"Expected semicolon separator in header, got: {lines[0]}"
    assert ";" in lines[1], f"Expected semicolon separator in data row, got: {lines[1]}"
    # The Total (Sum) label should appear in the index column
    assert "Total (Sum)" in lines[1]


@with_config(
    {
        "REPORTS_CSV_NA_NAMES": ["MISSING"],
        "CSV_EXPORT": {"sep": ";", "decimal": ","},
    }
)
def test_apply_client_processing_csv_format_na_values_and_sep_decimal_combined():
    """
    Test that apply_client_processing correctly handles both REPORTS_CSV_NA_NAMES and
    CSV_EXPORT sep/decimal config at the same time.
    """
    csv_data = "name;status\nJeff;MISSING\nAlice;OK"

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": csv_data,
            }
        ]
    }

    form_data = {
        "datasource": "1__table",
        "viz_type": "table",
        "slice_id": 1,
        "url_params": {},
        "metrics": [],
        "groupby": [],
        "columns": ["name", "status"],
        "extra_form_data": {},
        "force": False,
        "result_format": "csv",
        "result_type": "results",
    }

    processed_result = apply_client_processing(result, form_data)

    output_data = processed_result["queries"][0]["data"]
    lines = output_data.strip().split("\n")
    assert len(lines) == 3  # header + 2 data rows
    # Output must use configured separator
    assert ";" in lines[0], f"Expected semicolon separator in header, got: {lines[0]}"
    # "MISSING" should be treated as NA and rendered as empty in output
    assert lines[1].endswith(";"), f"Expected empty status for Jeff, got: {lines[1]}"
    # "OK" should be preserved as-is
    assert lines[2] == "Alice;OK", f"Expected Alice;OK, got: {lines[2]}"


@with_config({"CSV_EXPORT": {"decimal": ","}})
def test_apply_client_processing_csv_format_partial_config_decimal_only():
    """
    Test that apply_client_processing handles a partial CSV_EXPORT config where
    only decimal is set (sep falls back to the default comma).
    """
    # Default sep="," is used since no sep key is present in CSV_EXPORT
    csv_data = "name,value\nfoo,5\nbar,10"

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": csv_data,
            }
        ]
    }

    form_data = {
        "datasource": "1__table",
        "viz_type": "table",
        "slice_id": 1,
        "url_params": {},
        "metrics": [],
        "groupby": [],
        "columns": ["name", "value"],
        "extra_form_data": {},
        "force": False,
        "result_format": "csv",
        "result_type": "results",
    }

    processed_result = apply_client_processing(result, form_data)

    output_data = processed_result["queries"][0]["data"]
    lines = output_data.strip().split("\n")
    # Should parse and output correctly using default sep="," and decimal=","
    assert len(lines) == 3  # header + 2 data rows
    assert processed_result["queries"][0]["colnames"] == ["name", "value"]
    # Output uses default sep="," (no sep in partial config)
    assert lines[0] == "name,value", f"Expected comma-separated header, got: {lines[0]}"
    assert "foo" in lines[1]
    assert "bar" in lines[2]


def test_apply_client_processing_csv_format_pivot_table_multiple_rows():
    """
    When multiple "Rows" fields are selected in a pivot table, exporting to
    pivoted CSV should keep each field in its own column instead of merging
    them into a single column.

    See: https://github.com/apache/superset/issues/32369
    """
    csv_data = (
        "city,segment,value\n"
        "Paris,Consumer,10\n"
        "Paris,Corporate,20\n"
        "London,Consumer,30\n"
        "London,Corporate,40\n"
    )

    result = {
        "queries": [
            {
                "result_format": ChartDataResultFormat.CSV,
                "data": csv_data,
            }
        ]
    }
    form_data = {
        "viz_type": "pivot_table_v2",
        "groupbyColumns": [],
        "groupbyRows": ["city", "segment"],
        "metrics": ["value"],
        "metricsLayout": "COLUMNS",
    }

    processed_result = apply_client_processing(result, form_data)
    query = processed_result["queries"][0]

    output_data = query["data"]
    lines = output_data.strip().split("\n")

    # the header should have a separate column for each "Rows" field, instead
    # of a single merged column
    assert lines[0] == "city,segment,value"

    # each data row should have the city and segment in their own columns
    output_df = pd.read_csv(StringIO(output_data))
    assert list(output_df.columns) == ["city", "segment", "value"]
    assert set(zip(output_df["city"], output_df["segment"], strict=False)) == {
        ("Paris", "Consumer"),
        ("Paris", "Corporate"),
        ("London", "Consumer"),
        ("London", "Corporate"),
    }


@pytest.mark.parametrize("reducer", ["sum", "min", "max"])
def test_pivot_rollup_reducer_preserves_blank_and_zero(reducer: str):
    assert pd.isna(_reduce(pd.Series([float("nan"), float("nan")]), reducer))
    assert _reduce(pd.Series([0, float("nan")]), reducer) == 0


@pytest.mark.parametrize("layout", ["COLUMNS", "ROWS"])
@pytest.mark.parametrize("combine", [False, True])
@pytest.mark.parametrize("transpose", [False, True])
def test_pivot_mixed_totals_preserve_integer_csv(
    layout: str, combine: bool, transpose: bool
):
    """Blank totals must not round integer leaves or change their CSV spelling."""
    large = 2**53 + 1
    frame = pd.DataFrame(
        {
            "nation": ["US", "UK"],
            "gender": ["boy", "boy"],
            "SUM(a)": [large, 2],
            "SUM(b)": [3, 4],
        }
    )
    result = pivot_table_v2(
        frame,
        {
            "groupbyRows": ["nation"],
            "groupbyColumns": ["gender"],
            "metrics": ["SUM(a)", "SUM(b)"],
            "metricsLayout": layout,
            "combineMetric": combine,
            "transposePivot": transpose,
            "rowTotals": True,
            "colTotals": True,
        },
        apply_number_format=False,
    )
    csv = result.to_csv()
    assert str(large) in csv
    assert str(large - 1) not in csv
    assert ".0" not in csv


@pytest.mark.parametrize("mode", [None, "percent_row", "percent_col", "percent_total"])
@pytest.mark.parametrize("transpose", [False, True])
@pytest.mark.parametrize("single_metric", [False, True])
@pytest.mark.parametrize("layout", ["COLUMNS", "ROWS"])
def test_combined_metric_subtotals(
    mode: str | None, transpose: bool, single_metric: bool, layout: str
):
    frame = show_values_as_df()
    frame["MAX(num)"] = [1, 3, 6, 10]
    metrics = ["SUM(num)"] if single_metric else ["SUM(num)", "MAX(num)"]
    result = pivot_table_v2(
        frame,
        {
            "groupbyRows": ["nation"],
            "groupbyColumns": ["gender"],
            "metrics": metrics,
            "metricsLayout": layout,
            "combineMetric": True,
            "transposePivot": transpose,
            "rowTotals": True,
            "colTotals": True,
            "showValuesAs": mode,
        },
        apply_number_format=False,
    )
    # Transpose swaps the dimension axes before constructing the pivot.
    metrics_on_rows = layout == "ROWS"
    if metrics_on_rows:
        result = result.T
    groups = ["US", "UK"] if transpose != metrics_on_rows else ["boy", "girl"]
    for group in groups:
        subtotal = result.loc[:, (group, "Subtotal")]
        assert subtotal.notna().all() if single_metric else subtotal.isna().all()


@pytest.mark.parametrize("metrics", [set(), {"SUM(num)"}, {"SUM(num)", "MAX(num)"}])
def test_collapsed_metric_preserves_shared_metric_set(metrics: set[str]):
    original = metrics.copy()
    expected = "SUM(num)" if len(metrics) == 1 else None
    assert _collapsed_metric(metrics) == expected
    assert _collapsed_metric(metrics) == expected
    assert metrics == original


@pytest.mark.parametrize("combine", [False, True])
@pytest.mark.parametrize("transpose", [False, True])
def test_pivot_mixed_totals_preserve_integer_xlsx(combine: bool, transpose: bool):
    """The full XLSX path keeps large integers as exact text and small ones numeric."""
    from openpyxl import load_workbook

    large = 2**53 + 1
    source = pd.DataFrame(
        {
            "nation": ["US", "UK"],
            "gender": ["boy", "boy"],
            "SUM(a)": [str(large), "2"],
            "SUM(b)": [3, 4],
        }
    )
    result = apply_client_processing(
        {
            "queries": [
                {
                    "result_format": ChartDataResultFormat.XLSX,
                    "data": excel.df_to_excel(source, index=False),
                }
            ]
        },
        {
            "viz_type": "pivot_table_v2",
            "groupbyRows": ["nation"],
            "groupbyColumns": ["gender"],
            "metrics": ["SUM(a)", "SUM(b)"],
            "metricsLayout": "ROWS",
            "combineMetric": combine,
            "transposePivot": transpose,
            "rowTotals": True,
            "colTotals": True,
        },
    )
    sheet = load_workbook(BytesIO(result["queries"][0]["data"])).active
    cells = [cell for row in sheet for cell in row]
    assert any(cell.value == str(large) and cell.data_type == "s" for cell in cells)
    assert any(cell.value == 3 and cell.data_type == "n" for cell in cells)
    assert not any(str(cell.value).startswith(str(large - 1)) for cell in cells)


@pytest.mark.parametrize("apply_number_format", [False, True])
def test_pivot_mixed_str_int_big_integer_dimension(apply_number_format: bool) -> None:
    """A dimension that mixes exact big-integer strings, numbers and nulls (the
    shape the chart-data API uses for integers beyond 2**53 - 1) pivots without
    error and keeps every label exact."""
    wire = [
        {"big": "4611686018427387905", "metric": 1},
        {"big": None, "metric": 2},
        {"big": 7, "metric": 3},
        {"big": "4611686018427387905", "metric": 4},
        {"big": "-4611686018427387905", "metric": 5},
    ]
    df = pd.DataFrame.from_dict(wire)
    form_data = {
        "groupbyRows": ["big"],
        "groupbyColumns": [],
        "metrics": ["metric"],
        "aggregateFunction": "Sum",
        "rowTotals": True,
        "colTotals": True,
    }

    pivoted = pivot_table_v2(
        df.copy(), form_data, apply_number_format=apply_number_format
    )
    unformatted = table(df.copy(), {"all_columns": ["big", "metric"]}, None, True)

    sums = {
        label: str(value)
        for label, value, *_ in pivoted.reset_index().itertuples(index=False)
        if isinstance(label, str)
    }
    assert sums == {
        "7": "3",
        "-4611686018427387905": "5",
        "4611686018427387905": "5",
        "Total (Sum)": "15",
    }
    assert unformatted["big"].tolist() == df["big"].tolist()

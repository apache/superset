# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.
"""
Chart data over a real database: nullable large integers stay exact in plain
queries and keep today's float behaviour wherever Python-side processing runs.
"""

from __future__ import annotations

import copy
import io
import sqlite3
from collections.abc import Iterator
from datetime import datetime
from pathlib import Path
from typing import Any
from unittest.mock import patch

import pandas as pd
import pyarrow as pa
import pytest
from flask_caching import Cache
from openpyxl import load_workbook
from sqlalchemy.orm.session import Session

from superset.app import SupersetApp
from superset.charts.client_processing import apply_client_processing
from superset.charts.schemas import ChartDataQueryContextSchema
from superset.common.chart_data import ChartDataResultFormat, ChartDataResultType
from superset.common.query_context import QueryContext
from superset.common.query_context_factory import QueryContextFactory
from superset.common.query_serialization import load_serialized_query, serialize_query
from superset.connectors.sqla.models import SqlaTable, SqlMetric, TableColumn
from superset.constants import CacheRegion
from superset.models.core import Database
from superset.utils import json

LARGE = 333466480986292423


@pytest.fixture
def table(session: Session, tmp_path: Path) -> SqlaTable:
    """A table whose ``big`` column is a nullable BIGINT holding a large value."""
    path = tmp_path / "chart_data.db"
    connection = sqlite3.connect(path)
    connection.execute(
        "CREATE TABLE t (id INTEGER, grp TEXT, ts DATETIME, big BIGINT, val INTEGER)"
    )
    connection.executemany(
        "INSERT INTO t VALUES (?, ?, ?, ?, ?)",
        [
            (1, "a", "2024-01-01 00:00:00", LARGE, 10),
            (2, "a", "2024-01-02 00:00:00", None, 20),
            (3, "b", "2024-01-03 00:00:00", 1000, 30),
            (4, "b", "2024-01-04 00:00:00", LARGE, 40),
            (5, "c", "2024-01-05 00:00:00", None, 50),
        ],
    )
    connection.commit()
    connection.close()

    SqlaTable.metadata.create_all(session.get_bind())  # pylint: disable=no-member
    database = Database(database_name="db", sqlalchemy_uri=f"sqlite:///{path}")
    sqla_table = SqlaTable(
        table_name="t",
        database=database,
        columns=[
            TableColumn(column_name="id", type="INTEGER"),
            TableColumn(column_name="grp", type="TEXT", groupby=True),
            TableColumn(column_name="ts", type="DATETIME", is_dttm=True),
            TableColumn(column_name="big", type="BIGINT", verbose_name="Big Value"),
            TableColumn(column_name="val", type="INTEGER"),
        ],
        metrics=[
            SqlMetric(
                metric_name="sum_val", expression="SUM(val)", verbose_name="Sum Val"
            ),
            SqlMetric(
                metric_name="max_big", expression="MAX(big)", verbose_name="Max Big"
            ),
        ],
    )
    session.add(database)
    session.add(sqla_table)
    session.flush()
    # ``changed_on`` is part of every cache key; settle it so keys are stable
    sqla_table.changed_on = datetime(2024, 1, 1)
    session.flush()
    return sqla_table


def create_context(
    table: SqlaTable,
    queries: list[dict[str, Any]],
    result_type: ChartDataResultType = ChartDataResultType.FULL,
    result_format: ChartDataResultFormat = ChartDataResultFormat.JSON,
    form_data: dict[str, Any] | None = None,
    force: bool = True,
) -> QueryContext:
    return QueryContextFactory().create(
        datasource={"type": "table", "id": table.id},
        queries=queries,
        form_data=form_data or {},
        result_type=result_type,
        result_format=result_format,
        force=force,
    )


def run(table: SqlaTable, queries: list[dict[str, Any]], **kwargs: Any) -> Any:
    """Return the first query's payload for a freshly built context."""
    payload = create_context(table, queries, **kwargs).get_payload()
    return payload["queries"][0]


def rows(payload: dict[str, Any]) -> list[dict[str, Any]]:
    return sorted(payload["data"], key=lambda row: json.dumps(row, default=str))


def cells(data: bytes) -> list[list[Any]]:
    sheet = load_workbook(io.BytesIO(data)).active
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def schema_context(
    table: SqlaTable,
    queries: list[dict[str, Any]],
    result_type: ChartDataResultType,
    form_data: dict[str, Any] | None = None,
    result_format: ChartDataResultFormat = ChartDataResultFormat.JSON,
) -> QueryContext:
    """Build a context the way the chart-data API does for a saved chart."""
    body = {
        "datasource": {"id": table.id, "type": "table"},
        "queries": copy.deepcopy(queries),
        "form_data": form_data or {},
        "result_format": result_format.value,
        "result_type": result_type.value,
    }
    schema = ChartDataQueryContextSchema()
    schema.query_context_factory = QueryContextFactory()
    return schema.load(body)


def client_processed(
    table: SqlaTable,
    queries: list[dict[str, Any]],
    form_data: dict[str, Any],
    result_format: ChartDataResultFormat = ChartDataResultFormat.JSON,
) -> Any:
    """What a report receives: the client-processed result of a saved chart."""
    query_context = schema_context(
        table, queries, ChartDataResultType.POST_PROCESSED, form_data, result_format
    )
    processed = apply_client_processing(
        query_context.get_payload(), form_data, query_context.datasource
    )
    data = processed["queries"][0]["data"]
    return cells(data) if result_format == ChartDataResultFormat.XLSX else data


QUERY = {"row_limit": 100}


def test_plain_dimension_queries_keep_nullable_big_integers_exact(
    app_context: None, table: SqlaTable
) -> None:
    raw = rows(run(table, [{**QUERY, "columns": ["big"]}]))
    assert raw == [
        {"big": "333466480986292423"},
        {"big": "333466480986292423"},
        {"big": 1000},
        {"big": None},
        {"big": None},
    ]

    aggregate = rows(
        run(table, [{**QUERY, "columns": ["big"], "metrics": ["sum_val", "max_big"]}])
    )
    # the dimension is exact; the saved metrics (one of them nullable) are
    # normalized as before, whatever their verbose names
    assert aggregate == [
        {"big": "333466480986292423", "sum_val": 50, "max_big": 3.334664809862924e17},
        {"big": 1000, "sum_val": 30, "max_big": 1000.0},
        {"big": None, "sum_val": 70, "max_big": None},
    ]


@pytest.mark.parametrize(
    "result_type, queries",
    [
        pytest.param(
            ChartDataResultType.SAMPLES, [{**QUERY, "columns": ["big"]}], id="samples"
        ),
        pytest.param(
            ChartDataResultType.DRILL_DETAIL, [{**QUERY}], id="drill-to-detail"
        ),
    ],
)
def test_sample_queries_keep_nullable_big_integers_exact(
    app_context: None,
    table: SqlaTable,
    result_type: ChartDataResultType,
    queries: list[dict[str, Any]],
) -> None:
    payload = run(table, queries, result_type=result_type)

    assert sorted((row["big"] for row in payload["data"]), key=repr) == sorted(
        ["333466480986292423", "333466480986292423", 1000, None, None], key=repr
    )


def test_grouping_sets_levels_keep_nullable_big_integers_exact(
    app_context: None, table: SqlaTable
) -> None:
    """SQLite has no ``GROUPING SETS``, so this runs the per-level fallback."""
    query = {
        **QUERY,
        "columns": ["grp", "big"],
        "metrics": ["sum_val"],
        "grouping_sets": [["grp", "big"], ["grp"]],
    }

    data = run(table, [query])["data"]

    assert {row["big"] for row in data} == {"333466480986292423", 1000, None}


def test_exports_keep_nullable_big_integers_exact(
    app_context: None, table: SqlaTable
) -> None:
    query = {**QUERY, "columns": ["big", "grp"]}

    csv_data = run(table, [query], result_format=ChartDataResultFormat.CSV)["data"]
    xlsx_data = run(table, [query], result_format=ChartDataResultFormat.XLSX)["data"]
    arrow_data = run(table, [query], result_format=ChartDataResultFormat.ARROW)["data"]

    assert csv_data.decode("utf-8-sig") == (
        "Big Value,grp\n333466480986292423,a\n,a\n1000,b\n333466480986292423,b\n,c\n"
    )
    assert cells(xlsx_data) == [
        ["Big Value", "grp"],
        ["333466480986292423", "a"],
        [None, "a"],
        [1000, "b"],
        ["333466480986292423", "b"],
        [None, "c"],
    ]
    arrow = pa.ipc.open_stream(arrow_data).read_all()
    assert arrow.schema.field("big").type == pa.int64()
    assert arrow.column("big").to_pylist() == [LARGE, None, 1000, LARGE, None]


def test_metrics_stay_float(app_context: None, table: SqlaTable) -> None:
    by_group = rows(run(table, [{**QUERY, "columns": ["grp"], "metrics": ["max_big"]}]))

    assert by_group == [
        {"grp": "a", "max_big": 3.334664809862924e17},
        {"grp": "b", "max_big": 3.334664809862924e17},
        {"grp": "c", "max_big": None},
    ]


def test_post_processed_dimension_stays_float(
    app_context: None, table: SqlaTable
) -> None:
    cumulative = [
        {
            "operation": "cum",
            "options": {"operator": "sum", "columns": {"sum_val": "sum_val"}},
        }
    ]

    post_processing = rows(
        run(
            table,
            [
                {
                    **QUERY,
                    "columns": ["grp", "big"],
                    "metrics": ["sum_val"],
                    "post_processing": cumulative,
                }
            ],
        )
    )
    post_processed_context = rows(
        run(
            table,
            [{**QUERY, "columns": ["big"]}],
            result_type=ChartDataResultType.POST_PROCESSED,
        )
    )

    assert post_processing == [
        {"grp": "a", "big": 3.334664809862924e17, "sum_val": 30},
        {"grp": "a", "big": None, "sum_val": 20},
        {"grp": "b", "big": 1000.0, "sum_val": 60},
        {"grp": "b", "big": 3.334664809862924e17, "sum_val": 100},
        {"grp": "c", "big": None, "sum_val": 150},
    ]
    assert post_processed_context == [
        {"big": 1000.0},
        {"big": 3.334664809862924e17},
        {"big": 3.334664809862924e17},
        {"big": None},
        {"big": None},
    ]


def test_contribution_totals_unchanged(app_context: None, table: SqlaTable) -> None:
    contribution = [
        {
            "operation": "contribution",
            "options": {
                "orientation": "column",
                "columns": ["sum_val"],
                "rename_columns": ["%sum_val"],
            },
        }
    ]
    queries = [
        {
            **QUERY,
            "columns": ["big"],
            "metrics": ["sum_val"],
            "post_processing": contribution,
        },
        {**QUERY, "metrics": ["sum_val"]},
    ]

    payload = create_context(table, queries).get_payload()

    assert [rows(query) for query in payload["queries"]] == [
        [
            {"big": 1000.0, "sum_val": 30, "%sum_val": 0.2},
            {"big": 3.334664809862924e17, "sum_val": 50, "%sum_val": 1 / 3},
            {"big": None, "sum_val": 70, "%sum_val": 0.4666666666666667},
        ],
        [{"sum_val": 150}],
    ]


def test_time_offset_query_stays_float_and_null_keys_join(
    app_context: None, table: SqlaTable
) -> None:
    """The row whose dimension is NULL still finds its offset value."""
    query = {
        **QUERY,
        "columns": ["big"],
        "metrics": ["sum_val"],
        "time_offsets": ["3 days ago"],
        "granularity": "ts",
        "time_range": "2024-01-04 : 2024-01-06",
        "is_timeseries": True,
        "extras": {"time_grain_sqla": "P1D"},
    }

    assert rows(run(table, [query])) == [
        {
            "__timestamp": pd.Timestamp("2024-01-04"),
            "big": 3.334664809862924e17,
            "sum_val": 40,
            "sum_val__3 days ago": 10,
        },
        {
            "__timestamp": pd.Timestamp("2024-01-05"),
            "big": None,
            "sum_val": 50,
            "sum_val__3 days ago": 20,
        },
    ]


TABLE_RAW = {
    "viz_type": "table",
    "query_mode": "raw",
    "all_columns": ["big"],
    "metrics": [],
}
TABLE_AGGREGATE = {
    "viz_type": "table",
    "query_mode": "aggregate",
    "groupby": ["big"],
    "metrics": ["sum_val"],
}
PIVOT = {
    "viz_type": "pivot_table_v2",
    "groupbyRows": ["big"],
    "groupbyColumns": [],
    "metrics": ["sum_val"],
    "aggregateFunction": "Sum",
    "rowTotals": True,
    "colTotals": True,
    "transposePivot": False,
    "combineMetric": False,
}
RAW_QUERY = [{**QUERY, "columns": ["big"]}]
AGGREGATE_QUERY = [{**QUERY, "columns": ["big"], "metrics": ["sum_val"]}]


@pytest.mark.parametrize(
    "form_data, queries, result_format, expected",
    [
        pytest.param(
            {**TABLE_RAW, "column_config": {"big": {"d3NumberFormat": ",.2f"}}},
            RAW_QUERY,
            ChartDataResultFormat.JSON,
            {
                "Big Value": {
                    0: "333,466,480,986,292,416.00",
                    1: "",
                    2: "1,000.00",
                    3: "333,466,480,986,292,416.00",
                    4: "",
                }
            },
            id="raw-table-number-format",
        ),
        pytest.param(
            {
                **TABLE_RAW,
                "column_config": {
                    "big": {
                        "currencyFormat": {"symbol": "USD", "symbolPosition": "prefix"}
                    }
                },
            },
            RAW_QUERY,
            ChartDataResultFormat.JSON,
            {"Big Value": {0: "$ 333P", 1: "", 2: "$ 1k", 3: "$ 333P", 4: ""}},
            id="raw-table-currency",
        ),
        pytest.param(
            {**TABLE_AGGREGATE, "column_config": {"big": {"d3NumberFormat": ",.2f"}}},
            AGGREGATE_QUERY,
            ChartDataResultFormat.JSON,
            {
                "Big Value": {0: "", 1: "1,000.00", 2: "333,466,480,986,292,416.00"},
                "Sum Val": {0: "70", 1: "30", 2: "50"},
            },
            id="aggregate-table-number-format",
        ),
        pytest.param(
            PIVOT,
            AGGREGATE_QUERY,
            ChartDataResultFormat.JSON,
            {
                "Sum Val": {
                    "1000.0": "30",
                    "3.334664809862924e+17": "50",
                    "nan": "70",
                    "Total (Sum)": "150",
                },
                "Total (Sum)": {
                    "1000.0": "30",
                    "3.334664809862924e+17": "50",
                    "nan": "70",
                    "Total (Sum)": "150",
                },
            },
            id="pivot-table",
        ),
        pytest.param(
            TABLE_RAW,
            RAW_QUERY,
            ChartDataResultFormat.CSV,
            'Big Value\n3.334664809862924e+17\n""\n1000.0\n3.334664809862924e+17\n""\n',
            id="raw-table-csv",
        ),
        pytest.param(
            TABLE_RAW,
            RAW_QUERY,
            ChartDataResultFormat.XLSX,
            [
                ["Big Value"],
                ["3.334664809862924e+17"],
                [None],
                [1000],
                ["3.334664809862924e+17"],
            ],
            id="raw-table-xlsx",
        ),
        pytest.param(
            PIVOT,
            AGGREGATE_QUERY,
            ChartDataResultFormat.CSV,
            ",Sum Val,Total (Sum)\n1000.0,30,30\n3.334664809862924e+17,50,50\n"
            "nan,70,70\nTotal (Sum),150,150\n",
            id="pivot-table-csv",
        ),
        pytest.param(
            PIVOT,
            AGGREGATE_QUERY,
            ChartDataResultFormat.XLSX,
            [
                [None, "Sum Val", "Total (Sum)"],
                ["1000.0", 30, 30],
                ["3.334664809862924e+17", 50, 50],
                ["nan", 70, 70],
                ["Total (Sum)", 150, 150],
            ],
            id="pivot-table-xlsx",
        ),
    ],
)
def test_post_processed_table_formatting_matches_float_baseline(
    app_context: None,
    table: SqlaTable,
    form_data: dict[str, Any],
    queries: list[dict[str, Any]],
    result_format: ChartDataResultFormat,
    expected: Any,
) -> None:
    """Report (client-processed) Tables and Pivot Tables are formatted from float
    columns exactly as before: a large nullable dimension must not reach client
    processing as an unformattable object column."""
    assert client_processed(table, queries, form_data, result_format) == expected


@pytest.mark.parametrize(
    "context_type", [ChartDataResultType.FULL, ChartDataResultType.POST_PROCESSED]
)
def test_queries_inherit_the_context_result_type(
    app_context: None, table: SqlaTable, context_type: ChartDataResultType
) -> None:
    """Saved charts, report/POST bodies and async rebuilds all hand each query the
    context's result type, so the per-query value that gates exact integers (and
    is part of the cache key) always agrees with the context."""
    queries = [{**QUERY, "columns": ["big"]}]
    form_data = {"viz_type": "table", "result_type": context_type.value}
    contexts = {
        "saved-chart": schema_context(table, queries, context_type),
        "report-body": schema_context(table, queries, context_type, form_data),
    }
    contexts["async-rebuild"] = load_serialized_query(
        serialize_query(contexts["saved-chart"], 0)
    )

    for context in contexts.values():
        assert context.result_type == context_type
        assert context.queries[0].result_type == context_type
    keys = {
        name: context.queries[0].cache_key(datasource=table.uid)
        for name, context in contexts.items()
    }
    assert len(set(keys.values())) == 1


def test_full_and_post_processed_requests_use_different_cache_keys(
    app_context: None, table: SqlaTable
) -> None:
    queries = [{**QUERY, "columns": ["big"]}]

    full = schema_context(table, queries, ChartDataResultType.FULL)
    post_processed = schema_context(table, queries, ChartDataResultType.POST_PROCESSED)

    assert full.queries[0].cache_key(datasource=table.uid) != post_processed.queries[
        0
    ].cache_key(datasource=table.uid)


def test_mixed_result_type_query_shares_the_plain_cache_key(
    app_context: None, table: SqlaTable
) -> None:
    """Documented residual: a query that sets its own ``result_type`` keeps it
    under a differently typed context. Its cache key is the plain request's, so
    it is gated (and cached) as that plain request: exact and deterministic."""
    mixed = [{**QUERY, "columns": ["big"], "result_type": "full"}]
    plain = [{**QUERY, "columns": ["big"]}]

    mixed_context = schema_context(table, mixed, ChartDataResultType.POST_PROCESSED)
    plain_context = schema_context(table, plain, ChartDataResultType.FULL)

    assert mixed_context.result_type == ChartDataResultType.POST_PROCESSED
    assert mixed_context.queries[0].result_type == ChartDataResultType.FULL
    assert mixed_context.queries[0].cache_key(
        datasource=table.uid
    ) == plain_context.queries[0].cache_key(datasource=table.uid)
    assert mixed_context.get_payload()["queries"][0]["data"][0] == {
        "big": "333466480986292423"
    }


@pytest.fixture
def result_cache(app: SupersetApp, tmp_path: Path) -> Iterator[Cache]:
    cache = Cache(
        config={
            "CACHE_TYPE": "FileSystemCache",
            "CACHE_DIR": str(tmp_path / "result_cache"),
            "CACHE_DEFAULT_TIMEOUT": 300,
        }
    )
    cache.init_app(app)
    regions = {CacheRegion.DEFAULT: cache, CacheRegion.DATA: cache}
    with patch("superset.common.utils.query_cache_manager._cache", regions):
        yield cache


def test_warm_cache_and_force_cached_keep_each_requests_shape(
    app_context: None, table: SqlaTable, result_cache: Cache
) -> None:
    queries = [{**QUERY, "columns": ["big"]}]

    def df_payload(result_type: ChartDataResultType, force_cached: bool = False) -> Any:
        context = create_context(table, queries, result_type=result_type, force=False)
        return context.get_df_payload(context.queries[0], force_cached=force_cached)

    first = df_payload(ChartDataResultType.FULL)
    warm = df_payload(ChartDataResultType.FULL)
    force_cached = df_payload(ChartDataResultType.FULL, force_cached=True)

    assert [bool(p["is_cached"]) for p in (first, warm, force_cached)] == [
        False,
        True,
        True,
    ]
    for payload in (first, warm, force_cached):
        assert payload["df"]["big"].dtype == object
        assert payload["df"]["big"].tolist() == [LARGE, None, 1000, LARGE, None]

    # the POST_PROCESSED request has its own entry and keeps the float frame
    post_processed = df_payload(ChartDataResultType.POST_PROCESSED)
    post_processed_warm = df_payload(ChartDataResultType.POST_PROCESSED)
    assert not post_processed["is_cached"]
    assert post_processed_warm["is_cached"]
    for payload in (post_processed, post_processed_warm):
        assert payload["df"]["big"].dtype == "float64"

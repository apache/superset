# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

"""Bullet chart data, preview, and update paths through the shared MCP tools."""

import importlib
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from types import SimpleNamespace
from typing import Any
from unittest.mock import MagicMock, Mock, patch
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import pytest
import pytz
from dateutil import tz as dateutil_tz
from dateutil.zoneinfo import get_zonefile_instance
from fastmcp import Client

from superset.mcp_service.app import mcp
from superset.mcp_service.chart import query_result as query_result_module
from superset.mcp_service.chart.chart_helpers import build_query_dicts_from_form_data
from superset.mcp_service.chart.compile import CompileResult
from superset.mcp_service.chart.query_result import validate_query_result_envelope
from superset.mcp_service.chart.schemas import (
    BulletChartConfig,
    ChartError,
    GetChartDataRequest,
)
from superset.mcp_service.chart.tool.get_chart_data import _query_from_form_data
from superset.utils import json, json as utils_json
from superset.utils.core import GenericDataType


def query_result_data(
    result: Any, *, temporal_json_numbers: bool = False
) -> tuple[list[Any] | None, Any]:
    """Validate a result envelope and return every query's canonical rows."""
    if failure := validate_query_result_envelope(
        result, temporal_json_numbers=temporal_json_numbers
    ):
        return None, failure
    return [query["data"] for query in result["queries"]], None


update_chart_module = importlib.import_module(
    "superset.mcp_service.chart.tool.update_chart"
)


@pytest.fixture
def mock_auth():
    """Mock MCP auth so Client.call_tool() doesn't need a real admin user."""
    import importlib
    from contextlib import contextmanager
    from unittest.mock import patch

    _gcd_module = importlib.import_module(
        "superset.mcp_service.chart.tool.get_chart_data"
    )

    @contextmanager
    def _noop_log_context(*_args: Any, **_kwargs: Any) -> Any:
        yield lambda **_kw: None

    # Neutralize event_logger.log_context: the default DBEventLogger would
    # otherwise insert a log row referencing our mock user_id and fail a
    # FK constraint against the real users table. Patch via the module
    # object directly — the `tool` package's __init__.py re-exports the
    # get_chart_data function under the same name, which shadows the
    # submodule binding in the package namespace, so a dotted-string patch
    # target resolves to the function and mock.patch cannot find
    # event_logger on it.
    mock_event_logger = Mock()
    mock_event_logger.log_context.side_effect = _noop_log_context
    with (
        patch("superset.mcp_service.auth.get_user_from_request") as mock_get_user,
        patch.object(_gcd_module, "event_logger", mock_event_logger),
    ):
        user = Mock()
        user.id = 1
        user.username = "admin"
        mock_get_user.return_value = user
        yield mock_get_user


@pytest.mark.parametrize(
    "value",
    [
        timedelta(0),
        timedelta(microseconds=-1),
        timedelta(days=1, seconds=2, microseconds=3),
        timedelta.max,
        timedelta.min,
        pd.Timedelta(0),
        pd.Timedelta(-1, unit="ns"),
        pd.Timedelta(1, unit="ns"),
        pd.Timedelta("1 days 00:00:02.000003004"),
        pd.Timedelta.max,
        pd.Timedelta.min,
    ],
)
def test_bullet_duration_projection_matches_chart_data_json(value: object) -> None:
    expected = json.loads(json.dumps(value, default=json.json_int_dttm_ser))

    data, failure = query_result_data(
        {"queries": [{"data": [{"value": value}], "rowcount": 1}]},
        temporal_json_numbers=True,
    )

    assert failure is None
    assert data == [[{"value": expected}]]


@pytest.mark.parametrize(
    "value",
    [
        np.timedelta64("NaT"),
        np.timedelta64(0, "ns"),
        np.timedelta64(-1, "ns"),
        np.timedelta64(1, "D"),
        np.timedelta64(123456789, "ns"),
        np.timedelta64(np.iinfo("int64").max, "s"),
        np.timedelta64(np.iinfo("int64").min + 1, "us"),
    ],
)
def test_bullet_numpy_duration_matches_real_dataframe_producer(value: object) -> None:
    from superset.dataframe import df_to_records

    produced = df_to_records(
        pd.DataFrame({"value": pd.Series([value], dtype=object)}),
        convert_big_integers=False,
    )[0]["value"]
    expected = (
        None
        if produced is None
        else json.loads(json.dumps(produced, default=json.json_int_dttm_ser))
    )

    data, failure = query_result_data(
        {"queries": [{"data": [{"value": value}], "rowcount": 1}]},
        temporal_json_numbers=True,
    )

    assert failure is None
    assert data == [[{"value": expected}]]


@pytest.mark.parametrize(
    "value",
    [
        np.timedelta64(1, "Y"),
        np.timedelta64(1, "ps"),
        np.timedelta64(np.iinfo("int64").max, "D"),
    ],
)
def test_bullet_numpy_duration_fails_closed_for_ambiguous_or_overflowing_value(
    value: np.timedelta64,
) -> None:
    data, failure = query_result_data(
        {"queries": [{"data": [{"value": value}], "rowcount": 1}]},
        temporal_json_numbers=True,
    )

    assert data is None
    assert failure is not None
    assert "invalid NumPy duration" in failure.error


def test_bullet_duration_subclasses_are_rejected_without_hooks() -> None:
    class HostileTimedelta(timedelta):
        calls = 0

        def _call(self, *_args: object) -> Any:
            type(self).calls += 1
            raise AssertionError("hostile timedelta hook executed")

        __abs__ = _call
        __eq__ = _call
        __lt__ = _call
        __str__ = _call

    class HostilePandasTimedelta(pd.Timedelta):
        calls = 0

        def _call(self, *_args: object) -> Any:
            type(self).calls += 1
            raise AssertionError("hostile pandas timedelta hook executed")

        __abs__ = _call
        __eq__ = _call
        __lt__ = _call
        __str__ = _call

    for value in (HostileTimedelta(seconds=1), HostilePandasTimedelta("1s")):
        data, failure = query_result_data(
            {"queries": [{"data": [{"value": value}], "rowcount": 1}]},
            temporal_json_numbers=True,
        )
        assert data is None
        assert failure is not None
        assert "unsupported or subclassed value" in failure.error
    assert HostileTimedelta.calls == 0
    assert HostilePandasTimedelta.calls == 0


def test_bullet_timedelta_envelope_uses_chart_data_wire_boundary(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    expected = {
        "queries": [
            {
                "data": [{"duration": "1 day, 0:00:02.000003"}],
                "rowcount": 1,
            }
        ]
    }
    exact_bytes = len(
        json.dumps(expected, ensure_ascii=False, separators=(",", ":")).encode()
    )
    monkeypatch.setattr(
        query_result_module, "MAX_QUERY_RESULT_VALUE_BYTES", exact_bytes
    )
    result = {
        "queries": [
            {
                "data": [{"duration": timedelta(days=1, seconds=2, microseconds=3)}],
                "rowcount": 1,
            }
        ]
    }

    data, failure = query_result_data(result, temporal_json_numbers=True)

    assert failure is None
    assert data == [[{"duration": "1 day, 0:00:02.000003"}]]
    assert len(json.dumps(result, separators=(",", ":")).encode()) == exact_bytes

    monkeypatch.setattr(
        query_result_module, "MAX_QUERY_RESULT_VALUE_BYTES", exact_bytes - 1
    )
    data, failure = query_result_data(
        {
            "queries": [
                {
                    "data": [
                        {"duration": timedelta(days=1, seconds=2, microseconds=3)}
                    ],
                    "rowcount": 1,
                }
            ]
        },
        temporal_json_numbers=True,
    )

    assert data is None
    assert failure is not None
    assert "aggregate JSON bytes" in failure.error


def test_bullet_temporal_mode_projects_timestamp_and_numpy_before_generic() -> None:
    timestamp = pd.Timestamp("2024-01-02 03:04:05.123456789")
    numpy_timestamp = np.datetime64("1969-12-31T23:59:59.999999999")
    result = {
        "queries": [
            {
                "data": [
                    {
                        "timestamp": timestamp,
                        "numpy_timestamp": numpy_timestamp,
                        "date": date(2024, 1, 2),
                        "pandas_nat": pd.NaT,
                        "numpy_nat": np.datetime64("NaT"),
                    }
                ]
            }
        ]
    }

    data, failure = query_result_data(result, temporal_json_numbers=True)

    assert failure is None
    assert data == [
        [
            {
                "timestamp": 1704164645123.456,
                "numpy_timestamp": -0.0010000000000287557,
                "date": 1704153600000.0,
                "pandas_nat": None,
                "numpy_nat": None,
            }
        ]
    ]


def test_non_bullet_temporal_mode_retains_canonical_iso_projection() -> None:
    result = {
        "queries": [
            {
                "data": [
                    {
                        "timestamp": pd.Timestamp("2024-01-02 03:04:05.123456789"),
                        "numpy_timestamp": np.datetime64("2024-01-02", "D"),
                    }
                ]
            }
        ]
    }

    data, failure = query_result_data(result)

    assert failure is None
    assert data == [
        [
            {
                "timestamp": "2024-01-02T03:04:05.123456789",
                "numpy_timestamp": "2024-01-02T00:00:00",
            }
        ]
    ]


def _query_context_stub(form_data: dict[str, Any] | None = None) -> Any:
    """Return the minimal real-shaped context needed by Jinja form-data seeding."""
    return SimpleNamespace(form_data=form_data or {}, queries=[])


class _AsyncContext:
    async def report_progress(self, *args: Any, **kwargs: Any) -> None:
        pass


@pytest.fixture
def mcp_server():
    return mcp


@pytest.mark.asyncio
async def test_unsaved_bullet_get_data_uses_strict_render_model(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    module = importlib.import_module("superset.mcp_service.chart.tool.get_chart_data")
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )

    class _Command:
        def __init__(self, query_context: Any) -> None: ...
        def validate(self) -> None: ...
        def run(self) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": [{"Revenue": "not numeric"}],
                        "colnames": ["Revenue"],
                        "coltypes": [GenericDataType.NUMERIC],
                    }
                ]
            }

    monkeypatch.setattr(
        module,
        "build_query_context_from_form_data",
        lambda *_a, **_k: _query_context_stub(),
    )
    monkeypatch.setattr(command_module, "ChartDataCommand", _Command)

    response = await _query_from_form_data(
        {
            "datasource": "1__table",
            "viz_type": "bullet",
            "metric": "Revenue",
        },
        GetChartDataRequest(form_data_key="bullet"),
        _AsyncContext(),
    )

    assert isinstance(response, ChartError)
    assert response.error_type == "MalformedBulletOutput"


@pytest.mark.asyncio
async def test_saved_bullet_get_data_uses_strict_render_model(
    mcp_server: Any,
    mock_auth: Any,
) -> None:
    from unittest.mock import patch

    module = importlib.import_module("superset.mcp_service.chart.tool.get_chart_data")
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    chart = SimpleNamespace(
        id=10,
        slice_name="Strict Bullet",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        query_context='{"queries": []}',
        params='{"viz_type": "bullet", "metric": "Revenue"}',
    )

    class _Command:
        def __init__(self, query_context: Any) -> None: ...
        def validate(self) -> None: ...
        def run(self) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": [{"Revenue": "not numeric"}],
                        "colnames": ["Revenue"],
                        "coltypes": [GenericDataType.NUMERIC],
                    }
                ]
            }

    with (
        patch.object(module, "find_chart_by_identifier", return_value=chart),
        patch.object(
            module,
            "validate_chart_dataset",
            return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
        ),
        patch(
            "superset.charts.schemas.ChartDataQueryContextSchema.load",
            return_value=_query_context_stub(),
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
    ):
        async with Client(mcp_server) as client:
            result = await client.call_tool(
                "get_chart_data", {"request": {"identifier": 10}}
            )

    payload = json.loads(result.content[0].text)
    assert payload["error_type"] == "MalformedBulletOutput"


@pytest.mark.asyncio
async def test_saved_bullet_get_data_projects_dataframe_timestamps_to_epoch(
    mcp_server: Any,
    mock_auth: Any,
) -> None:
    from unittest.mock import patch

    from superset.dataframe import df_to_records

    module = importlib.import_module("superset.mcp_service.chart.tool.get_chart_data")
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    dateutil_dublin = dateutil_tz.gettz("Europe/Dublin")
    dateutil_new_york = dateutil_tz.gettz("America/New_York")
    assert dateutil_dublin is not None
    assert dateutil_new_york is not None
    dublin_fold = datetime(
        2024,
        10,
        27,
        1,
        30,
        0,
        123456,
        tzinfo=dateutil_dublin,
        fold=1,
    )
    new_york_gap = datetime(2024, 3, 10, 2, 30, 0, 123456, tzinfo=dateutil_new_york)
    form_data = {
        "viz_type": "bullet",
        "metric": "Revenue",
        "groupby": ["Category"],
    }
    chart = SimpleNamespace(
        id=20,
        slice_name="Timestamp Bullet",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        query_context='{"queries": []}',
        params=json.dumps(form_data),
    )
    source_values = [
        pd.Timestamp("2024-01-02 03:04:05.123456789"),
        pd.Timestamp("2024-01-02 08:34:05.123456789+05:30"),
        pd.Timestamp(
            datetime(
                2024,
                11,
                3,
                1,
                30,
                tzinfo=ZoneInfo("America/New_York"),
                fold=0,
            )
        ),
        pd.Timestamp(
            datetime(
                2024,
                11,
                3,
                1,
                30,
                tzinfo=ZoneInfo("America/New_York"),
                fold=1,
            )
        ),
        pd.Timestamp("1969-12-31 23:59:59.999999999"),
        pd.NaT,
        dublin_fold,
        new_york_gap,
        datetime(2040, 7, 1, 12, 0, 0, 123456, tzinfo=dateutil_new_york),
        pd.Timestamp(dublin_fold),
        pd.Timestamp(new_york_gap),
    ]
    rows = df_to_records(
        pd.DataFrame(
            {
                "Category": pd.Series(source_values, dtype=object),
                "Revenue": range(1, len(source_values) + 1),
            }
        ),
        convert_big_integers=False,
    )

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": rows,
                        "colnames": ["Category", "Revenue"],
                        "coltypes": [GenericDataType.TEMPORAL, GenericDataType.NUMERIC],
                        "rowcount": len(rows),
                    }
                ]
            }

    with (
        patch.object(module, "find_chart_by_identifier", return_value=chart),
        patch.object(
            module,
            "validate_chart_dataset",
            return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
        ),
        patch(
            "superset.charts.schemas.ChartDataQueryContextSchema.load",
            return_value=_query_context_stub(),
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
    ):
        async with Client(mcp_server) as client:
            result = await client.call_tool(
                "get_chart_data", {"request": {"identifier": 20}}
            )

    payload = json.loads(result.content[0].text)
    assert [row["Category"] for row in payload["data"]] == [
        1704164645123.456,
        1704164645123.456,
        1730611800000.0,
        1730615400000.0,
        -0.0010000000000287557,
        None,
        1729989000123.456,
        1710052200123.456,
        2224774800123.456,
        1729989000123.456,
        1710052200123.456,
    ]


@pytest.mark.asyncio
async def test_transitionless_dateutil_dataframe_reaches_fastmcp_data(
    mcp_server: Any,
    mock_auth: Any,
) -> None:
    from unittest.mock import patch

    from dateutil.zoneinfo import get_zonefile_instance

    from superset.commands.chart.data.get_data_command import (
        ChartDataCommand as ProducerChartDataCommand,
    )
    from superset.common.chart_data import ChartDataResultType
    from superset.dataframe import df_to_records
    from superset.utils.json import json_int_dttm_ser

    module = importlib.import_module("superset.mcp_service.chart.tool.get_chart_data")
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    names = [
        "UTC",
        "GMT",
        "Universal",
        "Zulu",
        "EST",
        "HST",
        "MST",
        "Etc/GMT+1",
        "Etc/GMT-2",
    ]
    values = []
    for getter in (dateutil_tz.gettz, get_zonefile_instance().get):
        for name in names:
            timezone_value = getter(name)
            assert timezone_value is not None
            values.append(
                datetime(2040, 7, 1, 12, 34, 56, 123456, tzinfo=timezone_value)
            )
    rows = df_to_records(
        pd.DataFrame(
            {
                "Category": pd.Series(values, dtype=object),
                "Revenue": range(1, len(values) + 1),
            }
        ),
        convert_big_integers=False,
    )

    class _ProducerContext:
        result_type = ChartDataResultType.FULL

        def get_payload(self, **_kwargs: Any) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": rows,
                        "colnames": ["Category", "Revenue"],
                        "coltypes": [GenericDataType.TEMPORAL, GenericDataType.NUMERIC],
                        "rowcount": len(rows),
                    }
                ]
            }

    producer_result = ProducerChartDataCommand(
        _ProducerContext()  # type: ignore[arg-type]
    ).run()
    form_data = {
        "viz_type": "bullet",
        "metric": "Revenue",
        "groupby": ["Category"],
    }
    chart = SimpleNamespace(
        id=22,
        slice_name="Transitionless timestamps",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        query_context='{"queries": []}',
        params=json.dumps(form_data),
    )

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return producer_result

    with (
        patch.object(module, "find_chart_by_identifier", return_value=chart),
        patch.object(
            module,
            "validate_chart_dataset",
            return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
        ),
        patch(
            "superset.charts.schemas.ChartDataQueryContextSchema.load",
            return_value=_query_context_stub(),
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
    ):
        async with Client(mcp_server) as client:
            result = await client.call_tool(
                "get_chart_data", {"request": {"identifier": 22}}
            )

    payload = json.loads(result.content[0].text)
    assert payload["row_count"] == len(values)
    assert payload["total_rows"] == len(values)
    assert [row["Category"] for row in payload["data"]] == [
        json_int_dttm_ser(value) for value in values
    ]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("format_", "excel_engine"),
    [
        ("json", None),
        ("csv", None),
        ("excel", "openpyxl"),
        ("excel", "xlsxwriter"),
    ],
)
async def test_saved_cached_bullet_duration_producer_reaches_fastmcp_exports(
    mcp_server: Any,
    mock_auth: Any,
    format_: str,
    excel_engine: str | None,
) -> None:
    import base64
    import csv
    import io
    from contextlib import ExitStack
    from unittest.mock import patch

    from openpyxl import load_workbook

    from superset.commands.chart.data.get_data_command import (
        ChartDataCommand as ProducerChartDataCommand,
    )
    from superset.common.chart_data import ChartDataResultType
    from superset.dataframe import df_to_records

    module = importlib.import_module("superset.mcp_service.chart.tool.get_chart_data")
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    source_values = [
        timedelta(0),
        timedelta(microseconds=-1),
        timedelta(days=1, seconds=2, microseconds=3),
        pd.Timedelta(-1, unit="ns"),
        pd.Timedelta("1 days 00:00:02.000003004"),
        np.timedelta64(123456789, "ns"),
        np.timedelta64("NaT"),
    ]
    rows = df_to_records(
        pd.DataFrame(
            {
                "Duration": pd.Series(source_values, dtype=object),
                "Revenue": range(1, len(source_values) + 1),
            }
        ),
        convert_big_integers=False,
    )
    assert type(rows[5]["Duration"]) is pd.Timedelta
    assert rows[6]["Duration"] is None
    expected = [
        None
        if row["Duration"] is None
        else json.loads(json.dumps(row["Duration"], default=json.json_int_dttm_ser))
        for row in rows
    ]

    class _ProducerContext:
        result_type = ChartDataResultType.FULL

        def get_payload(self, **_kwargs: Any) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": rows,
                        "colnames": ["Duration", "Revenue"],
                        "coltypes": [GenericDataType.STRING, GenericDataType.NUMERIC],
                        "rowcount": len(rows),
                        "is_cached": True,
                        "cache_key": "duration-cache",
                        "cached_dttm": "2026-09-03T00:00:00+00:00",
                    }
                ]
            }

    producer_result = ProducerChartDataCommand(
        _ProducerContext()  # type: ignore[arg-type]
    ).run()
    form_data = {
        "viz_type": "bullet",
        "metric": "Revenue",
        "groupby": ["Duration"],
    }
    chart = SimpleNamespace(
        id=23,
        slice_name="Duration Bullet",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        query_context='{"queries": []}',
        params=json.dumps(form_data),
    )

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return producer_result

    with ExitStack() as stack:
        stack.enter_context(
            patch.object(module, "find_chart_by_identifier", return_value=chart)
        )
        stack.enter_context(
            patch.object(
                module,
                "validate_chart_dataset",
                return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
            )
        )
        stack.enter_context(
            patch(
                "superset.charts.schemas.ChartDataQueryContextSchema.load",
                return_value=_query_context_stub(),
            )
        )
        stack.enter_context(patch.object(command_module, "ChartDataCommand", _Command))
        if excel_engine == "xlsxwriter":
            stack.enter_context(
                patch.object(
                    module, "_create_excel_with_openpyxl", side_effect=ImportError
                )
            )
        async with Client(mcp_server) as client:
            result = await client.call_tool(
                "get_chart_data",
                {"request": {"identifier": 23, "format": format_}},
            )

    payload = json.loads(result.content[0].text)
    assert payload["row_count"] == len(rows)
    assert payload["total_rows"] == len(rows)
    assert payload["cache_status"]["cache_hit"] is True
    if format_ == "json":
        assert [row["Duration"] for row in payload["data"]] == expected
    elif format_ == "csv":
        decoded = list(csv.DictReader(io.StringIO(payload["csv_data"])))
        assert [row["Duration"] or None for row in decoded] == expected
    else:
        workbook = load_workbook(io.BytesIO(base64.b64decode(payload["excel_data"])))
        assert [cell.value for cell in workbook.active[1]] == ["Duration", "Revenue"]
        assert [
            workbook.active.cell(row=index + 2, column=1).value
            for index in range(len(expected))
        ] == expected


@pytest.mark.asyncio
async def test_form_data_key_bullet_duration_producer_uses_chart_data_wire(
    mcp_server: Any,
    mock_auth: Any,
) -> None:
    from unittest.mock import patch

    from superset.commands.chart.data.get_data_command import (
        ChartDataCommand as ProducerChartDataCommand,
    )
    from superset.common.chart_data import ChartDataResultType
    from superset.dataframe import df_to_records

    module = importlib.import_module("superset.mcp_service.chart.tool.get_chart_data")
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    rows = df_to_records(
        pd.DataFrame(
            {
                "Duration": pd.Series(
                    [timedelta(days=1, microseconds=3), pd.Timedelta(-1, "ns")],
                    dtype=object,
                ),
                "Revenue": [1, 2],
            }
        ),
        convert_big_integers=False,
    )
    expected = [
        json.loads(json.dumps(row["Duration"], default=json.json_int_dttm_ser))
        for row in rows
    ]

    class _ProducerContext:
        result_type = ChartDataResultType.FULL

        def get_payload(self, **_kwargs: Any) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": rows,
                        "colnames": ["Duration", "Revenue"],
                        "coltypes": [GenericDataType.STRING, GenericDataType.NUMERIC],
                        "rowcount": len(rows),
                    }
                ]
            }

    producer_result = ProducerChartDataCommand(
        _ProducerContext()  # type: ignore[arg-type]
    ).run()
    form_data = {
        "datasource": "1__table",
        "viz_type": "bullet",
        "metric": "Revenue",
        "groupby": ["Duration"],
    }

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return producer_result

    with (
        patch.object(
            module, "get_cached_form_data", return_value=json.dumps(form_data)
        ),
        patch.object(
            module,
            "build_query_context_from_form_data",
            return_value=_query_context_stub(),
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
    ):
        async with Client(mcp_server) as client:
            result = await client.call_tool(
                "get_chart_data",
                {"request": {"form_data_key": "duration-bullet"}},
            )

    payload = json.loads(result.content[0].text)
    assert payload["chart_id"] == 0
    assert [row["Duration"] for row in payload["data"]] == expected


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("grouped", "format_", "identifier_alias", "excel_engine"),
    [
        (False, "json", "id", None),
        (True, "json", "chart_id", None),
        (False, "csv", "id", None),
        (True, "csv", "chart_id", None),
        (False, "excel", "id", "openpyxl"),
        (True, "excel", "id", "openpyxl"),
        (False, "excel", "chart_id", "xlsxwriter"),
        (True, "excel", "chart_id", "xlsxwriter"),
    ],
)
async def test_saved_empty_bullet_get_data_fastmcp_returns_normalized_rows(
    mcp_server: Any,
    mock_auth: Any,
    grouped: bool,
    format_: str,
    identifier_alias: str,
    excel_engine: str | None,
) -> None:
    import base64
    import io
    from contextlib import ExitStack
    from unittest.mock import patch

    from openpyxl import load_workbook

    module = importlib.import_module("superset.mcp_service.chart.tool.get_chart_data")
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    form_data: dict[str, Any] = {
        "viz_type": "bullet",
        "metric": "Revenue",
    }
    if grouped:
        form_data["groupby"] = ["Region"]
    raw_columns = ["Revenue", *(["Region"] if grouped else [])]
    chart = SimpleNamespace(
        id=21,
        slice_name="Empty Bullet",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        query_context='{"queries": []}',
        params=json.dumps(form_data),
    )

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return {
                "queries": 2
                * [
                    {
                        "data": [],
                        "colnames": raw_columns,
                        "coltypes": [GenericDataType.STRING for _ in raw_columns],
                        "rowcount": 0,
                    }
                ]
            }

    with ExitStack() as stack:
        stack.enter_context(
            patch.object(module, "find_chart_by_identifier", return_value=chart)
        )
        stack.enter_context(
            patch.object(
                module,
                "validate_chart_dataset",
                return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
            )
        )
        stack.enter_context(
            patch(
                "superset.charts.schemas.ChartDataQueryContextSchema.load",
                return_value=_query_context_stub(),
            )
        )
        stack.enter_context(patch.object(command_module, "ChartDataCommand", _Command))
        if excel_engine == "xlsxwriter":
            stack.enter_context(
                patch.object(
                    module, "_create_excel_with_openpyxl", side_effect=ImportError
                )
            )
        async with Client(mcp_server) as client:
            result = await client.call_tool(
                "get_chart_data",
                {
                    "request": {
                        identifier_alias: 21,
                        "format": format_,
                    }
                },
            )

    payload = json.loads(result.content[0].text)
    assert "error_type" not in payload
    if format_ == "json":
        assert payload["data"] == []
        assert payload["row_count"] == 0
        assert payload["total_rows"] == 0
        assert len(payload["query_results"]) == 2
        for query_result in payload["query_results"]:
            assert query_result["data"] == []
            assert query_result["row_count"] == 0
            assert query_result["total_rows"] == 0
    elif format_ == "csv":
        assert payload["format"] == "csv"
        assert payload["row_count"] == 0
        assert payload["total_rows"] == 0
        assert payload["csv_data"] == (
            "Revenue,Region\r\n" if grouped else "Revenue\r\n"
        )
    else:
        assert payload["format"] == "excel"
        assert payload["row_count"] == 0
        assert payload["total_rows"] == 0
        workbook = load_workbook(io.BytesIO(base64.b64decode(payload["excel_data"])))
        assert [cell.value for cell in workbook.active[1]] == raw_columns


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "grouped",
    [
        False,
        True,
    ],
)
async def test_form_data_key_empty_bullet_fastmcp_returns_normalized_rows(
    mcp_server: Any,
    mock_auth: Any,
    grouped: bool,
) -> None:
    from unittest.mock import patch

    module = importlib.import_module("superset.mcp_service.chart.tool.get_chart_data")
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    form_data: dict[str, Any] = {
        "datasource": "1__table",
        "viz_type": "bullet",
        "metric": "Revenue",
    }
    if grouped:
        form_data["groupby"] = ["Region"]
    raw_columns = ["Revenue", *(["Region"] if grouped else [])]

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return {
                "queries": 2
                * [
                    {
                        "data": [],
                        "colnames": raw_columns,
                        "coltypes": [GenericDataType.STRING for _ in raw_columns],
                        "rowcount": 0,
                    }
                ]
            }

    with (
        patch.object(
            module, "get_cached_form_data", return_value=json.dumps(form_data)
        ),
        patch.object(
            module,
            "build_query_context_from_form_data",
            return_value=_query_context_stub(),
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
    ):
        async with Client(mcp_server) as client:
            result = await client.call_tool(
                "get_chart_data",
                {
                    "request": {
                        "form_data_key": "empty-bullet",
                        "format": "json",
                    }
                },
            )

    payload = json.loads(result.content[0].text)
    assert "error_type" not in payload
    assert payload["chart_id"] == 0
    assert payload["chart_type"] == "bullet"
    assert payload["row_count"] == 0
    assert payload["total_rows"] == 0
    assert payload["data"] == []
    assert len(payload["query_results"]) == 2
    for query_result in payload["query_results"]:
        assert query_result["data"] == []
        assert query_result["row_count"] == 0
        assert query_result["total_rows"] == 0


@pytest.mark.asyncio
@pytest.mark.parametrize("number_format", ["DURATION", "MEMORY_BINARY"])
@pytest.mark.parametrize("format_", ["json", "csv", "excel"])
async def test_saved_bullet_native_formatter_does_not_block_data_or_export(
    mcp_server: Any,
    mock_auth: Any,
    number_format: str,
    format_: str,
) -> None:
    """Native presentation presets must not make valid raw rows unreadable."""
    import base64
    import io

    from openpyxl import load_workbook

    module = importlib.import_module("superset.mcp_service.chart.tool.get_chart_data")
    form_data = {
        "viz_type": "bullet",
        "metric": "Revenue",
        "groupby": ["Region"],
        "y_axis_format": number_format,
    }
    chart = SimpleNamespace(
        id=21,
        slice_name="Native Bullet",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        query_context='{"queries": []}',
        params=json.dumps(form_data),
    )
    rows = [{"Revenue": 1024, "Region": "EU"}]
    with (
        patch.object(module, "find_chart_by_identifier", return_value=chart),
        patch.object(
            module,
            "validate_chart_dataset",
            return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
        ),
        patch(
            "superset.charts.schemas.ChartDataQueryContextSchema.load",
            return_value=_query_context_stub(),
        ),
        patch(
            "superset.commands.chart.data.get_data_command.ChartDataCommand"
        ) as command,
    ):
        command.return_value.run.return_value = {
            "queries": [
                {
                    "data": rows,
                    "colnames": ["Revenue", "Region"],
                    "coltypes": [GenericDataType.NUMERIC, GenericDataType.STRING],
                    "rowcount": 1,
                }
            ]
        }
        async with Client(mcp_server) as client:
            result = await client.call_tool(
                "get_chart_data",
                {"request": {"identifier": 21, "format": format_}},
            )
    payload = json.loads(result.content[0].text)
    assert "error_type" not in payload, payload
    if format_ == "json":
        assert payload["data"] == rows
    elif format_ == "csv":
        assert payload["csv_data"] == "Revenue,Region\r\n1024,EU\r\n"
    else:
        workbook = load_workbook(io.BytesIO(base64.b64decode(payload["excel_data"])))
        assert list(workbook.active.values) == [("Revenue", "Region"), (1024, "EU")]


@pytest.mark.asyncio
@pytest.mark.parametrize("format_", ["ascii", "vega_lite"])
async def test_bullet_numeric_and_temporal_categories_reach_real_mcp_entrypoint(
    format_: str,
) -> None:
    from contextlib import nullcontext

    from superset.mcp_service.app import mcp

    preview_module = importlib.import_module(
        "superset.mcp_service.chart.tool.get_chart_preview"
    )
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    form_data = {
        "viz_type": "bullet",
        "metric": "Revenue",
        "groupby": ["Category"],
    }
    chart = SimpleNamespace(
        id=121,
        slice_name="Number boundaries",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        params=utils_json.dumps(form_data),
    )
    rows = [
        {"Category": 9007199254740993, "Revenue": 1},
        {"Category": Decimal("1.0000000000000001"), "Revenue": 2},
        {"Category": Decimal("1.7976931348623159e308"), "Revenue": 3},
        {"Category": date(2026, 9, 2), "Revenue": 4},
        {
            "Category": datetime(2026, 9, 2, 3, 4, 5, tzinfo=timezone.utc),
            "Revenue": 5,
        },
        {
            "Category": datetime(
                2023,
                11,
                5,
                1,
                30,
                tzinfo=ZoneInfo("America/New_York"),
                fold=0,
            ),
            "Revenue": 6,
        },
        {
            "Category": datetime(
                2023,
                11,
                5,
                1,
                30,
                tzinfo=ZoneInfo("America/New_York"),
                fold=1,
            ),
            "Revenue": 7,
        },
    ]

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": rows,
                        "colnames": ["Category", "Revenue"],
                        "coltypes": [GenericDataType.TEMPORAL, GenericDataType.NUMERIC],
                    }
                ]
            }

    query_context = SimpleNamespace(
        form_data={},
        queries=[SimpleNamespace(metrics=["Revenue"], columns=["Category"])],
    )
    user = MagicMock(id=1, username="admin", roles=[], groups=[])
    with (
        patch("superset.mcp_service.auth.get_user_from_request", return_value=user),
        patch("superset.mcp_service.auth.check_tool_permission", return_value=True),
        patch.object(preview_module, "find_chart_by_identifier", return_value=chart),
        patch.object(preview_module.db.session, "refresh", return_value=None),
        patch.object(
            preview_module,
            "validate_chart_dataset",
            return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
        ),
        patch.object(
            preview_module.event_logger,
            "log_context",
            side_effect=lambda **_kwargs: nullcontext(),
        ),
        patch.object(
            preview_module,
            "build_query_context_from_form_data",
            return_value=query_context,
        ),
        patch(
            "superset.charts.data.form_data.set_query_context_form_data",
            return_value=None,
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
        patch.object(
            preview_module, "get_superset_base_url", return_value="http://localhost"
        ),
    ):
        async with Client(mcp) as client:
            result = await client.call_tool(
                "get_chart_preview",
                {"request": {"id": 121, "format": format_}},
            )

    payload = utils_json.loads(result.content[0].text)
    if format_ == "ascii":
        content = payload["content"]["ascii_content"]
        assert "9007199254740992" in content
        assert "Infinity" in content
        assert "1788307200000" in content
        assert "1699162200000" in content
        assert "1699165800000" in content
    else:
        specification = payload["content"]["specification"]
        bar = next(
            layer for layer in specification["layer"] if layer["mark"]["type"] == "bar"
        )
        category_field = bar["encoding"]["tooltip"][0]["field"]
        assert [row[category_field] for row in specification["data"]["values"]] == [
            "9007199254740992",
            "1",
            "Infinity",
            "1788307200000",
            "1788318245000",
            "1699162200000",
            "1699165800000",
        ]
        assert [row["Category"] for row in specification["data"]["values"][3:]] == [
            1788307200000.0,
            1788318245000.0,
            1699162200000.0,
            1699165800000.0,
        ]
        assert bar["encoding"]["tooltip"][0]["field"] == category_field
        assert "transform" not in specification


@pytest.mark.asyncio
@pytest.mark.parametrize("format_", ["ascii", "vega_lite"])
async def test_bullet_timestamp_categories_from_dataframe_reach_fastmcp(
    format_: str,
) -> None:
    from contextlib import nullcontext

    from superset.dataframe import df_to_records
    from superset.mcp_service.app import mcp

    preview_module = importlib.import_module(
        "superset.mcp_service.chart.tool.get_chart_preview"
    )
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    zoneinfo_tz = ZoneInfo("America/New_York")
    pytz_tz = pytz.timezone("America/New_York")
    dateutil_dublin = dateutil_tz.gettz("Europe/Dublin")
    dateutil_new_york = dateutil_tz.gettz("America/New_York")
    assert dateutil_dublin is not None
    assert dateutil_new_york is not None
    dublin_fold = datetime(
        2024,
        10,
        27,
        1,
        30,
        0,
        123456,
        tzinfo=dateutil_dublin,
        fold=1,
    )
    new_york_gap = datetime(2024, 3, 10, 2, 30, 0, 123456, tzinfo=dateutil_new_york)
    source_values = [
        pd.Timestamp("2024-01-02 03:04:05.123456789"),
        pd.Timestamp("2024-01-02 08:34:05.123456789+05:30"),
        pd.Timestamp(datetime(2024, 11, 3, 1, 30, tzinfo=zoneinfo_tz, fold=0)),
        pd.Timestamp(datetime(2024, 11, 3, 1, 30, tzinfo=zoneinfo_tz, fold=1)),
        pd.Timestamp(pytz_tz.localize(datetime(2024, 11, 3, 1, 30), is_dst=True)),
        pd.Timestamp(pytz_tz.localize(datetime(2024, 11, 3, 1, 30), is_dst=False)),
        pd.Timestamp("1969-12-31 23:59:59.999999999"),
        date(2024, 1, 2),
        pd.NaT,
        dublin_fold,
        new_york_gap,
        datetime(2040, 7, 1, 12, 0, 0, 123456, tzinfo=dateutil_new_york),
        datetime(
            2024,
            3,
            10,
            2,
            30,
            0,
            123456,
            tzinfo=dateutil_tz.tzoffset("EDT", -4 * 3600),
        ),
        datetime(2024, 3, 10, 6, 30, 0, 123456, tzinfo=timezone.utc),
        pd.Timestamp(dublin_fold),
        pd.Timestamp(new_york_gap),
    ]
    rows = df_to_records(
        pd.DataFrame(
            {
                "Category": pd.Series(source_values, dtype=object),
                "Revenue": range(1, len(source_values) + 1),
            }
        ),
        convert_big_integers=False,
    )
    if format_ == "ascii":
        # ASCII intentionally displays at most ten categories. Keep every
        # dateutil named-zone edge in this public-format invocation.
        rows = rows[9:12]
    expected = [
        "1704164645123.456",
        "1704164645123.456",
        "1730611800000",
        "1730615400000",
        "1730611800000",
        "1730615400000",
        "-0.0010000000000287557",
        "1704153600000",
        "null",
        "1729989000123.456",
        "1710052200123.456",
        "2224774800123.456",
        "1710052200123.456",
        "1710052200123.456",
        "1729989000123.456",
        "1710052200123.456",
    ]
    form_data = {
        "viz_type": "bullet",
        "metric": "Revenue",
        "groupby": ["Category"],
    }
    chart = SimpleNamespace(
        id=122,
        slice_name="Timestamp categories",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        params=utils_json.dumps(form_data),
    )

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": rows,
                        "colnames": ["Category", "Revenue"],
                        "coltypes": [GenericDataType.TEMPORAL, GenericDataType.NUMERIC],
                    }
                ]
            }

    query_context = SimpleNamespace(
        form_data={},
        queries=[SimpleNamespace(metrics=["Revenue"], columns=["Category"])],
    )
    user = MagicMock(id=1, username="admin", roles=[], groups=[])
    with (
        patch("superset.mcp_service.auth.get_user_from_request", return_value=user),
        patch("superset.mcp_service.auth.check_tool_permission", return_value=True),
        patch.object(preview_module, "find_chart_by_identifier", return_value=chart),
        patch.object(preview_module.db.session, "refresh", return_value=None),
        patch.object(
            preview_module,
            "validate_chart_dataset",
            return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
        ),
        patch.object(
            preview_module.event_logger,
            "log_context",
            side_effect=lambda **_kwargs: nullcontext(),
        ),
        patch.object(
            preview_module,
            "build_query_context_from_form_data",
            return_value=query_context,
        ),
        patch(
            "superset.charts.data.form_data.set_query_context_form_data",
            return_value=None,
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
        patch.object(
            preview_module, "get_superset_base_url", return_value="http://localhost"
        ),
    ):
        async with Client(mcp) as client:
            result = await client.call_tool(
                "get_chart_preview",
                {"request": {"id": 122, "format": format_}},
            )

    payload = utils_json.loads(result.content[0].text)
    if format_ == "ascii":
        content = payload["content"]["ascii_content"]
        for category in set(expected[9:12]):
            assert category[:20] in content
    else:
        specification = payload["content"]["specification"]
        bar = next(
            layer for layer in specification["layer"] if layer["mark"]["type"] == "bar"
        )
        category_field = bar["encoding"]["tooltip"][0]["field"]
        assert [row[category_field] for row in specification["data"]["values"]] == (
            expected
        )
        assert bar["encoding"]["tooltip"][0]["field"] == category_field
        assert [row["Category"] for row in specification["data"]["values"]] == [
            1704164645123.456,
            1704164645123.456,
            1730611800000.0,
            1730615400000.0,
            1730611800000.0,
            1730615400000.0,
            -0.0010000000000287557,
            1704153600000.0,
            None,
            1729989000123.456,
            1710052200123.456,
            2224774800123.456,
            1710052200123.456,
            1710052200123.456,
            1729989000123.456,
            1710052200123.456,
        ]


@pytest.mark.asyncio
async def test_transitionless_dateutil_dataframe_reaches_fastmcp_preview() -> None:
    from contextlib import nullcontext

    from superset.commands.chart.data.get_data_command import (
        ChartDataCommand as ProducerChartDataCommand,
    )
    from superset.common.chart_data import ChartDataResultType
    from superset.dataframe import df_to_records
    from superset.mcp_service.app import mcp
    from superset.mcp_service.chart.preview_utils import _javascript_number_string
    from superset.utils.json import json_int_dttm_ser

    preview_module = importlib.import_module(
        "superset.mcp_service.chart.tool.get_chart_preview"
    )
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    names = [
        "UTC",
        "GMT",
        "Universal",
        "Zulu",
        "EST",
        "HST",
        "MST",
        "Etc/GMT+1",
        "Etc/GMT-2",
    ]
    values = []
    for getter in (dateutil_tz.gettz, get_zonefile_instance().get):
        for name in names:
            timezone_value = getter(name)
            assert timezone_value is not None
            values.append(
                datetime(2040, 7, 1, 12, 34, 56, 123456, tzinfo=timezone_value)
            )
    rows = df_to_records(
        pd.DataFrame(
            {
                "Category": pd.Series(values, dtype=object),
                "Revenue": range(1, len(values) + 1),
            }
        ),
        convert_big_integers=False,
    )

    class _ProducerContext:
        result_type = ChartDataResultType.FULL

        def get_payload(self, **_kwargs: Any) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": rows,
                        "colnames": ["Category", "Revenue"],
                        "coltypes": [GenericDataType.TEMPORAL, GenericDataType.NUMERIC],
                        "rowcount": len(rows),
                    }
                ]
            }

    producer_result = ProducerChartDataCommand(
        _ProducerContext()  # type: ignore[arg-type]
    ).run()
    expected_numbers = [json_int_dttm_ser(value) for value in values]
    expected_categories = [
        _javascript_number_string(float(value)) for value in expected_numbers
    ]
    form_data = {
        "viz_type": "bullet",
        "metric": "Revenue",
        "groupby": ["Category"],
    }
    chart = SimpleNamespace(
        id=123,
        slice_name="Transitionless timestamps",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        params=utils_json.dumps(form_data),
    )

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return producer_result

    query_context = SimpleNamespace(
        form_data={},
        queries=[SimpleNamespace(metrics=["Revenue"], columns=["Category"])],
    )
    user = MagicMock(id=1, username="admin", roles=[], groups=[])
    with (
        patch("superset.mcp_service.auth.get_user_from_request", return_value=user),
        patch("superset.mcp_service.auth.check_tool_permission", return_value=True),
        patch.object(preview_module, "find_chart_by_identifier", return_value=chart),
        patch.object(preview_module.db.session, "refresh", return_value=None),
        patch.object(
            preview_module,
            "validate_chart_dataset",
            return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
        ),
        patch.object(
            preview_module.event_logger,
            "log_context",
            side_effect=lambda **_kwargs: nullcontext(),
        ),
        patch.object(
            preview_module,
            "build_query_context_from_form_data",
            return_value=query_context,
        ),
        patch(
            "superset.charts.data.form_data.set_query_context_form_data",
            return_value=None,
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
        patch.object(
            preview_module, "get_superset_base_url", return_value="http://localhost"
        ),
    ):
        async with Client(mcp) as client:
            result = await client.call_tool(
                "get_chart_preview",
                {"request": {"id": 123, "format": "vega_lite"}},
            )

    payload = utils_json.loads(result.content[0].text)
    specification = payload["content"]["specification"]
    bar = next(
        layer for layer in specification["layer"] if layer["mark"]["type"] == "bar"
    )
    category_field = bar["encoding"]["tooltip"][0]["field"]
    assert [row["Category"] for row in specification["data"]["values"]] == (
        expected_numbers
    )
    assert [row[category_field] for row in specification["data"]["values"]] == (
        expected_categories
    )
    assert bar["encoding"]["tooltip"][0]["field"] == category_field


@pytest.mark.asyncio
@pytest.mark.parametrize("format_", ["ascii", "vega_lite"])
async def test_duration_dataframe_reaches_fastmcp_bullet_preview(
    format_: str,
) -> None:
    from contextlib import nullcontext

    import numpy as np

    from superset.commands.chart.data.get_data_command import (
        ChartDataCommand as ProducerChartDataCommand,
    )
    from superset.common.chart_data import ChartDataResultType
    from superset.dataframe import df_to_records
    from superset.mcp_service.app import mcp

    preview_module = importlib.import_module(
        "superset.mcp_service.chart.tool.get_chart_preview"
    )
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    source_values = [
        timedelta(0),
        timedelta(days=1, seconds=2, microseconds=3),
        pd.Timedelta(-1, unit="ns"),
        np.timedelta64(123456789, "ns"),
        np.timedelta64("NaT"),
    ]
    rows = df_to_records(
        pd.DataFrame(
            {
                "Duration": pd.Series(source_values, dtype=object),
                "Revenue": [50, 120, 350, 0, 200],
            }
        ),
        convert_big_integers=False,
    )
    assert type(rows[3]["Duration"]) is pd.Timedelta
    assert rows[4]["Duration"] is None
    expected = [
        "null"
        if row["Duration"] is None
        else utils_json.loads(
            utils_json.dumps(row["Duration"], default=utils_json.json_int_dttm_ser)
        )
        for row in rows
    ]

    class _ProducerContext:
        result_type = ChartDataResultType.FULL

        def get_payload(self, **_kwargs: Any) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": rows,
                        "colnames": ["Duration", "Revenue"],
                        "coltypes": [GenericDataType.STRING, GenericDataType.NUMERIC],
                        "rowcount": len(rows),
                    }
                ]
            }

    producer_result = ProducerChartDataCommand(
        _ProducerContext()  # type: ignore[arg-type]
    ).run()
    form_data = {
        "viz_type": "bullet",
        "metric": "Revenue",
        "groupby": ["Duration"],
        "ranges": "100,200,300",
        "range_labels": "Low,Mid,High",
    }
    chart = SimpleNamespace(
        id=124,
        slice_name="Duration categories",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        params=utils_json.dumps(form_data),
    )

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return producer_result

    query_context = SimpleNamespace(
        form_data={},
        queries=[SimpleNamespace(metrics=["Revenue"], columns=["Duration"])],
    )
    user = MagicMock(id=1, username="admin", roles=[], groups=[])
    with (
        patch("superset.mcp_service.auth.get_user_from_request", return_value=user),
        patch("superset.mcp_service.auth.check_tool_permission", return_value=True),
        patch.object(preview_module, "find_chart_by_identifier", return_value=chart),
        patch.object(preview_module.db.session, "refresh", return_value=None),
        patch.object(
            preview_module,
            "validate_chart_dataset",
            return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
        ),
        patch.object(
            preview_module.event_logger,
            "log_context",
            side_effect=lambda **_kwargs: nullcontext(),
        ),
        patch.object(
            preview_module,
            "build_query_context_from_form_data",
            return_value=query_context,
        ),
        patch(
            "superset.charts.data.form_data.set_query_context_form_data",
            return_value=None,
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
        patch.object(
            preview_module, "get_superset_base_url", return_value="http://localhost"
        ),
    ):
        async with Client(mcp) as client:
            result = await client.call_tool(
                "get_chart_preview",
                {"request": {"id": 124, "format": format_}},
            )

    payload = utils_json.loads(result.content[0].text)
    if format_ == "ascii":
        for category in expected:
            assert category[:20] in payload["content"]["ascii_content"]
    else:
        specification = payload["content"]["specification"]
        bar = next(
            layer for layer in specification["layer"] if layer["mark"]["type"] == "bar"
        )
        category_field = bar["encoding"]["tooltip"][0]["field"]
        range_tooltip = next(
            item for item in bar["encoding"]["tooltip"] if item.get("title") == "Range"
        )
        assert [row[category_field] for row in specification["data"]["values"]] == (
            expected
        )
        assert [
            row[range_tooltip["field"]] for row in specification["data"]["values"]
        ] == ["Low", "Mid", "> High", "Low", "Mid"]
        range_layers = [
            layer for layer in specification["layer"] if layer["mark"]["type"] == "rect"
        ]
        assert range_layers
        assert all(
            layer["encoding"]["y"] == bar["encoding"]["y"] for layer in range_layers
        )


@pytest.mark.asyncio
@pytest.mark.parametrize("format_", ["ascii", "vega_lite"])
@pytest.mark.parametrize("saved", [True, False])
@pytest.mark.parametrize("stray_tokens", [False, True])
async def test_bullet_short_labels_and_case_distinct_dimensions_reach_fastmcp(
    format_: str,
    saved: bool,
    stray_tokens: bool,
) -> None:
    from contextlib import nullcontext

    from superset.mcp_service.app import mcp

    preview_module = importlib.import_module(
        "superset.mcp_service.chart.tool.get_chart_preview"
    )
    command_module = importlib.import_module(
        "superset.commands.chart.data.get_data_command"
    )
    form_data = {
        "viz_type": "bullet",
        "datasource": "1__table",
        "metric": "Revenue",
        "groupby": ["Region", "region"],
        "ranges": "10,20,30",
        "range_labels": "Low",
        "markers": "12,22",
        "marker_labels": "Plan",
        "marker_lines": "13,23",
        "marker_line_labels": "Forecast",
        "show_labels": True,
        "y_axis_format": ".1f",
    }
    if stray_tokens:
        for control in ("ranges", "markers", "marker_lines"):
            form_data[control] = f"{form_data[control]},nope,NaN,"
    chart = SimpleNamespace(
        id=121,
        slice_name="Number boundaries",
        viz_type="bullet",
        datasource_id=1,
        datasource_type="table",
        params=utils_json.dumps(form_data),
    )
    rows = [{"Region": "North", "region": "South", "Revenue": 15}]

    class _Command:
        def __init__(self, _query_context: Any) -> None: ...

        def validate(self) -> None: ...

        def run(self) -> dict[str, Any]:
            return {
                "queries": [
                    {
                        "data": rows,
                        "colnames": ["Region", "region", "Revenue"],
                        "coltypes": [
                            GenericDataType.STRING,
                            GenericDataType.STRING,
                            GenericDataType.NUMERIC,
                        ],
                    }
                ]
            }

    query_context = SimpleNamespace(
        form_data={},
        queries=[SimpleNamespace(metrics=["Revenue"], columns=["Region", "region"])],
    )
    user = MagicMock(id=1, username="admin", roles=[], groups=[])
    with (
        patch("superset.mcp_service.auth.get_user_from_request", return_value=user),
        patch("superset.mcp_service.auth.check_tool_permission", return_value=True),
        patch(
            "superset.commands.explore.form_data.get.GetFormDataCommand.run",
            return_value=utils_json.dumps(form_data),
        ),
        patch.object(preview_module, "find_chart_by_identifier", return_value=chart),
        patch.object(preview_module.db.session, "refresh", return_value=None),
        patch.object(
            preview_module,
            "validate_chart_dataset",
            return_value=SimpleNamespace(is_valid=True, warnings=[], error=None),
        ),
        patch.object(
            preview_module.event_logger,
            "log_context",
            side_effect=lambda **_kwargs: nullcontext(),
        ),
        patch.object(
            preview_module,
            "build_query_context_from_form_data",
            return_value=query_context,
        ),
        patch(
            "superset.charts.data.form_data.set_query_context_form_data",
            return_value=None,
        ),
        patch.object(command_module, "ChartDataCommand", _Command),
        patch.object(
            preview_module, "get_superset_base_url", return_value="http://localhost"
        ),
    ):
        request_payload: dict[str, str | int] = {"format": format_}
        if saved:
            request_payload["id"] = 121
        else:
            request_payload["form_data_key"] = "short-labels"
        async with Client(mcp) as client:
            result = await client.call_tool(
                "get_chart_preview",
                {"request": request_payload},
            )

    payload = utils_json.loads(result.content[0].text)
    assert payload["chart_type"] == "bullet"
    if format_ == "ascii":
        assert "North, South" in payload["content"]["ascii_content"]
    else:
        spec = payload["content"]["specification"]
        bar = next(layer for layer in spec["layer"] if layer["mark"]["type"] == "bar")
        category = bar["encoding"]["tooltip"][0]["field"]
        assert spec["data"]["values"][0][category] == "North, South"
        labels = {
            layer["encoding"]["x"]["datum"]: layer["encoding"]["text"]["value"]
            for layer in spec["layer"]
            if layer["mark"]["type"] == "text"
        }
        assert labels == {
            10.0: "Low",
            12.0: "Plan",
            22.0: "22.0",
            13.0: "Forecast",
            23.0: "23.0",
        }


def test_configless_bullet_rebind_preserves_native_sql_and_having() -> None:
    native_filters = [
        {
            "clause": "WHERE",
            "expressionType": "SQL",
            "sqlExpression": "region <> 'unknown'",
        },
        {
            "clause": "HAVING",
            "expressionType": "SIMPLE",
            "subject": "SavedRevenue",
            "operator": "GREATER_THAN",
            "comparator": 10,
        },
    ]
    form_data = {
        "viz_type": "bullet",
        "metric": "SavedRevenue",
        "groupby": ["Region"],
        "adhoc_filters": native_filters,
        "datasource": "1041__table",
    }
    chart = SimpleNamespace(id=55, datasource_id=10)
    dataset = SimpleNamespace(id=1041)

    def validate(
        config: Any,
        compiled_form_data: dict[str, Any],
        target_dataset: Any,
        *,
        run_compile_check: bool,
    ) -> CompileResult:
        assert isinstance(config, BulletChartConfig)
        assert not config.filters
        assert compiled_form_data["adhoc_filters"] == native_filters
        assert target_dataset is dataset
        assert run_compile_check is True
        return CompileResult(success=True)

    with (
        patch("superset.mcp_service.auth.has_dataset_access", return_value=True),
        patch("superset.daos.dataset.DatasetDAO.find_by_id", return_value=dataset),
        patch.object(update_chart_module, "validate_and_compile", side_effect=validate),
    ):
        result = update_chart_module._validate_update_against_dataset(
            None, form_data, chart, dataset_id=1041
        )

    assert result is None
    assert form_data["adhoc_filters"] == native_filters


def test_build_query_dicts_preserves_native_orderby_for_single_query(monkeypatch):
    monkeypatch.setattr(
        "superset.mcp_service.chart.chart_helpers.resolve_datasource_engine",
        lambda datasource_id, datasource_type: "base",
    )
    orderby = [["SUM(revenue)", False]]
    query = build_query_dicts_from_form_data(
        {
            "viz_type": "bullet",
            "metric": "SUM(revenue)",
            "groupby": ["region"],
            "orderby": orderby,
        },
        1,
        "table",
    )[0]

    assert query["columns"] == ["region"]
    assert query["metrics"] == ["SUM(revenue)"]
    assert query["orderby"] == orderby
